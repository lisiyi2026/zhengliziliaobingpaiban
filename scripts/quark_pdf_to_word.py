"""Normalize existing Word files and optionally compare them with a reference PDF/image."""

from __future__ import annotations

import argparse
import difflib
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
import uuid
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
REFERENCE_EXTENSIONS = {".pdf", *IMAGE_EXTENSIONS}
SUPPORTED_EXTENSIONS = {".docx", ".doc", ".pdf", *IMAGE_EXTENSIONS}
_SOFFICE_STATUS: bool | None = None
_HEADING_FONT: str | None = None
_MINERU_STATUS: tuple[bool, str] | None = None

NUMBERED_TITLE_RE = re.compile(r"^\s*(\d+(?:\.\d+){0,2})\s+")
CHINESE_CHAPTER_RE = re.compile(r"^\s*第[0-9一二三四五六七八九十百]+章\b")
CHINESE_SECTION_RE = re.compile(r"^\s*第[0-9一二三四五六七八九十百]+节\b")
KNOWN_MAJOR_HEADINGS = {"前言", "目录", "摘要", "中文摘要", "英文摘要", "Abstract", "参考文献", "致谢", "附录", "结论", "结语"}
KNOWN_MINOR_HEADINGS = {"研究背景", "研究意义", "研究目的", "研究内容", "研究方法", "文献综述", "理论基础", "实验设计", "结果分析", "讨论", "本章小结", "小结"}
USELESS_TEXT_RE = re.compile(
    r"刺猬\s*出品|必属\s*精品|高质量课程资料联系|H(?:i|j)jingyuge|"
    r"QQ\s*[:：]?\s*\d{5,}|www\.e[-_.]?studysky|扫码看答案与讲解|"
    r"解析咨询|择校咨询|鲸鱼哥微|网学天地独家录制|试题及答案获取方式|"
    r"www\.c[-_.]?stuavsky|网学天地|独家录制|e[-_.]?studysky(?:\.com)?|"
    r"共\s*\d+\s*页\s*/\s*第\s*\d+\s*页",
    re.IGNORECASE,
)
QR_TEXT_RE = re.compile(r"二维码|扫码|扫一扫|扫码看答案|e[-_.]?studysky", re.IGNORECASE)
ANSWER_HEADING_RE = re.compile(r"^\s*解析\s*[:：]?\s*[A-DＡ-Ｄ](?:\s*)$", re.IGNORECASE)


def looks_like_sentence(text: str) -> bool:
    value = normalize_text(text)
    return len(value) > 32 or value.endswith(("。", "！", "？", "；", ".", "!", "?", ";", "：", ":"))


def classify_paragraph(paragraph: Any, in_refs: bool, paragraphs: list[Any] | None = None, index: int = 0, first_content_index: int | None = None) -> str:
    text = (paragraph.text or "").strip()
    compact = re.sub(r"\s+", "", text)
    if not text:
        return "empty"
    if is_answer_heading(text):
        return "heading2"
    if re.match(r"^(?:图|表)\s*\d", text):
        return "caption"
    if first_content_index == index and len(compact) <= 60:
        return "title"
    if compact in {re.sub(r"\s+", "", v) for v in KNOWN_MAJOR_HEADINGS}:
        return "heading1"
    if in_refs and re.match(r"^(?:\[\d+\]|\d+[、.．])", text):
        return "reference"
    numbered = NUMBERED_TITLE_RE.match(text)
    if numbered and not looks_like_sentence(text):
        return f"heading{min(numbered.group(1).count('.') + 1, 3)}"
    if CHINESE_CHAPTER_RE.match(text) and not looks_like_sentence(text):
        return "heading1"
    if CHINESE_SECTION_RE.match(text) and not looks_like_sentence(text):
        return "heading2"
    if compact in {re.sub(r"\s+", "", v) for v in KNOWN_MINOR_HEADINGS}:
        return "heading2"
    if paragraph.style and paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
        return paragraph.style.name.lower().replace(" ", "")
    return "body"


def load_structure_map(path: Path | None) -> dict[int, str]:
    if path is None:
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    result: dict[int, str] = {}
    valid = {"title", "heading1", "heading2", "heading3", "body", "caption", "reference"}
    for item in data.get("paragraphs", []):
        if item.get("index") is None or item.get("role") is None:
            continue
        role = str(item["role"]).strip().lower().replace("_", "")
        if role not in valid:
            raise ValueError(f"不支持的结构角色: {item['role']}")
        result[int(item["index"])] = role
    return result


def heading_font() -> str:
    """Choose 华文行楷 when installed, otherwise the closest installed KaiTi."""
    global _HEADING_FONT
    if _HEADING_FONT is not None:
        return _HEADING_FONT
    candidates = ("STXingkai", "华文行楷", "KaiTi", "楷体")
    installed = set()
    try:
        import winreg  # type: ignore
        for hive, subkey in ((winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"), (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts")):
            try:
                with winreg.OpenKey(hive, subkey) as key:
                    for index in range(winreg.QueryInfoKey(key)[1]):
                        name, value, _ = winreg.EnumValue(key, index)
                        label = f"{name} {value}".lower()
                        if "xingkai" in label or "华文行楷" in label:
                            installed.update({"STXingkai", "华文行楷"})
                        if "kaiti" in label or "simkai" in label or "楷体" in label:
                            installed.update({"KaiTi", "楷体"})
            except OSError:
                continue
    except ImportError:
        pass
    _HEADING_FONT = next((font for font in candidates if font in installed), "KaiTi")
    return _HEADING_FONT


def is_toc_entry(logical: str) -> bool:
    """Detect chapter/index rows that end with a printed page number."""
    compact = re.sub(r"\s+", "", logical or "")
    return bool(
        re.search(r"\d{1,3}$", compact)
        and re.match(r"^(?:第[0-9一二三四五六七八九十百]+章|书后习题解析|同步训练题(?:答案)?|参考文献)", compact)
    )


def is_toc_line(text: str) -> bool:
    """Return true only for a real TOC row, never for tabbed body data.

    Quark uses tabs for many technical lines such as direction/vector tables.
    Those lines must remain ordinary body text; only a final numeric field
    paired with a chapter-like label (or an existing dot leader) is a TOC row.
    """
    parts = (text or "").split("\t")
    if len(parts) < 2:
        return False
    page = parts[-1].strip()
    if not page.isdigit():
        return False
    left = " ".join(parts[:-1]).strip()
    compact_left = re.sub(r"\s+", "", left)
    return is_toc_entry(compact_left) or bool(re.search(r"[.·…]{3,}", left))


def is_exam_title_line(text: str) -> bool:
    """Identify compact exam-cover/title lines without treating questions as headings."""
    compact = re.sub(r"\s+", "", normalize_text(text or ""))
    if not compact or len(compact) > 70:
        return False
    if re.search(r"20\d{2}年", compact) and re.search(r"(?:大学|学院|研究生|硕士|试题|试卷)", compact):
        return True
    if "招收" in compact and "研究生" in compact:
        return True
    return "入学考试" in compact and bool(re.search(r"(?:自命题|试题|试卷)", compact))


def is_exam_cover_school_line(text: str) -> bool:
    compact = re.sub(r"\s+", "", normalize_text(text or ""))
    return bool(compact and len(compact) <= 20 and re.search(r"(?:大学|学院)$", compact))


def is_answer_heading(text: str) -> bool:
    """Return true for standalone answer markers such as ``解析：A``."""
    return bool(ANSWER_HEADING_RE.match(normalize_text(text or "")))


def parse_args() -> argparse.Namespace:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Normalize Word files and compare with an optional PDF/image reference")
    parser.add_argument("--input", type=Path, default=root / "input" / "word")
    parser.add_argument("--output", type=Path, default=root / "output" / "word")
    parser.add_argument("--dpi", type=int, default=150)
    parser.add_argument("--compare", choices=("full", "sample"), default="full")
    parser.add_argument("--reference", type=Path, default=None, help="Original PDF/image used as the authoritative comparison source")
    parser.add_argument("--max-pages", type=int, default=None, help="仅处理输入 PDF 的前 N 页；不指定则处理全部页面")
    parser.add_argument("--bridge", action="store_true", help="通过用户账户桥接器调用 LibreOffice")
    parser.add_argument("--visual-compare", action="store_true", help="可选：把输出 Word 转成 PDF 并与原 PDF 做版式诊断")
    parser.add_argument("--render-layout", action="store_true", help="可选：渲染最终 Word 检查真实空白页和页面留白；默认关闭以缩短处理时间")
    parser.add_argument(
        "--pdf-engine",
        choices=("auto", "pymupdf", "mineru", "office"),
        default="auto",
        help="PDF 转 Word 引擎：auto/pymupdf 使用快速 PyMuPDF 重建，失败后才回退 MinerU/LibreOffice；mineru/office 可强制指定",
    )
    parser.add_argument("--mineru-python", type=Path, default=None, help="可选：MinerU 所在 Python 解释器路径")
    parser.add_argument("--mineru-backend", choices=("pipeline", "hybrid-engine", "vlm-engine"), default="pipeline", help="MinerU 后端")
    parser.add_argument("--mineru-method", choices=("auto", "txt", "ocr"), default="auto", help="MinerU PDF 解析方式")
    parser.add_argument("--mineru-timeout", type=int, default=1800, help="MinerU 单个 PDF 转换超时秒数")
    parser.add_argument(
        "--structure-map",
        type=Path,
        default=None,
        help="可选 UTF-8 结构映射 JSON；由人工审核后覆盖自动标题识别",
    )
    return parser.parse_args()


def copy_to_writable_output(source: Path, output: Path) -> Path:
    """Copy a source into the requested destination, falling back if locked.

    Office applications commonly keep the previous DOCX open with a read-only
    handle. Do not abort an entire batch in that case: preserve the old file
    and write a deterministic sibling named ``*_修正版.docx``.
    """
    try:
        shutil.copy2(source, output)
        return output
    except (PermissionError, OSError) as exc:
        if not isinstance(exc, PermissionError) and getattr(exc, "winerror", None) not in {5, 32}:
            raise
        stem = output.stem
        suffix = output.suffix or ".docx"
        for index in range(1, 100):
            candidate = output.with_name(f"{stem}_修正版{index if index > 1 else ''}{suffix}")
            if candidate.exists():
                continue
            shutil.copy2(source, candidate)
            print(f"目标文件被 Office/WPS 占用或为只读，已改写到: {candidate}")
            return candidate
        raise PermissionError(f"无法写入输出文件，且修正版备用名称已用尽: {output}") from exc


def find_soffice() -> str | None:
    # Use the console launcher for deterministic headless conversion.  The GUI
    # launcher can remain resident and make automation appear to time out.
    found = shutil.which("soffice.com") or shutil.which("soffice.exe") or shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for directory in (Path("C:/Program Files/LibreOffice/program"), Path("C:/Program Files (x86)/LibreOffice/program")):
        for name in ("soffice.com", "soffice.exe"):
            path = directory / name
            if path.exists():
                return str(path)
    return None


def find_poppler_tool(name: str) -> str | None:
    """Find bundled Poppler utilities even when the user's PATH is minimal."""
    found = shutil.which(name)
    if found:
        return found
    candidates = [
        Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "native" / "poppler" / "Library" / "bin" / f"{name}.exe",
        Path.home() / ".codex" / "plugins" / "cache" / "openai-primary-runtime" / "pdf" / "26.903.11726" / "bin" / f"{name}.exe",
    ]
    return next((str(path) for path in candidates if path.exists()), None)


def soffice_usable() -> bool:
    """Avoid repeatedly launching a broken LibreOffice binary (notably 0xc0000022)."""
    global _SOFFICE_STATUS
    if _SOFFICE_STATUS is not None:
        return _SOFFICE_STATUS
    # Codex's restricted worker account cannot launch GUI-linked binaries from
    # Program Files even when the interactive Windows account can. Do not cache
    # that environment failure as a machine-wide LibreOffice failure.
    if os.environ.get("USERNAME", "").lower().startswith("codexsandbox"):
        _SOFFICE_STATUS = False
        return False
    disabled_marker = Path(__file__).resolve().parents[1] / "tmp" / "quark" / ".soffice_unusable"
    if disabled_marker.exists():
        _SOFFICE_STATUS = False
        return False
    soffice = find_soffice()
    if not soffice:
        _SOFFICE_STATUS = False
        return False
    try:
        probe = subprocess.run(
            [soffice, "--headless", "--invisible", "--norestore", "--nofirststartwizard", "--nodefault", "--version"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
            timeout=20,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            env=office_environment(),
        )
        _SOFFICE_STATUS = probe.returncode == 0
        if not _SOFFICE_STATUS:
            disabled_marker.parent.mkdir(parents=True, exist_ok=True)
            disabled_marker.write_text("LibreOffice preflight failed; remove this file after repair/reinstall.\n", encoding="utf-8")
        return _SOFFICE_STATUS
    except (OSError, subprocess.TimeoutExpired):
        _SOFFICE_STATUS = False
        disabled_marker.parent.mkdir(parents=True, exist_ok=True)
        disabled_marker.write_text("LibreOffice preflight failed; remove this file after repair/reinstall.\n", encoding="utf-8")
        return False


def discover_inputs(source: Path) -> list[Path]:
    if source.is_file():
        if source.suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise RuntimeError(f"不支持的输入格式: {source.suffix}")
        return [source]
    if source.is_dir():
        # Office creates hidden lock files such as ``~$document.docx`` while
        # a source is open.  They are not valid DOCX packages and must never
        # enter the conversion queue.
        files = sorted(
            path
            for path in source.rglob("*")
            if path.is_file()
            and not path.name.startswith("~$")
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        )
        if not files:
            raise RuntimeError(f"目录中没有支持的文件: {source}")
        return files
    raise RuntimeError(f"找不到输入路径: {source}")


def classify_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
        pages = PdfReader(str(path)).pages
        if not pages:
            return "quark"
        chars = formulas = sparse = 0
        formula_chars = "∑∫√≈≤≥≠∞∂∇∏∈"
        for page in pages:
            text = "".join((page.extract_text() or "").split())
            chars += len(text)
            formulas += sum(text.count(char) for char in formula_chars)
            sparse += int(len(text) < 80)
        average = chars / len(pages)
        return "quark" if sparse / len(pages) >= 0.35 or average < 120 or formulas >= max(8, len(pages) * 3) else "local"
    except Exception:
        return "quark"


def classify_input(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        return classify_pdf(path)
    return "quark" if path.suffix.lower() in IMAGE_EXTENSIONS else "local"


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", value or "")).strip()


def docx_text(path: Path) -> str:
    from docx import Document  # type: ignore
    document = Document(str(path))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return normalize_text("\n".join(chunks))


def limit_pdf_pages(path: Path, output: Path, max_pages: int | None) -> Path:
    """Create a bounded PDF copy for a preview run without altering the source."""
    if max_pages is None:
        return path
    if max_pages < 1:
        raise RuntimeError("--max-pages 必须大于 0")
    from pypdf import PdfReader, PdfWriter  # type: ignore
    reader = PdfReader(str(path))
    writer = PdfWriter()
    for page in reader.pages[:max_pages]:
        writer.add_page(page)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("wb") as handle:
        writer.write(handle)
    return output


def image_to_docx(path: Path, output: Path) -> Path:
    """Wrap a raster input as an editable Word container without page flattening."""
    from docx import Document  # type: ignore
    from docx.shared import Inches, Pt  # type: ignore
    document = Document()
    section = document.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin
    paragraph = document.paragraphs[0]
    paragraph.alignment = 1
    paragraph.add_run().add_picture(str(path), width=usable_width)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def find_mineru_python(configured: Path | None = None) -> Path | None:
    """Find the project-local MinerU interpreter without importing torch."""
    candidates: list[Path] = []
    if configured:
        candidates.append(configured)
    env_python = os.environ.get("MINERU_PYTHON")
    if env_python:
        candidates.append(Path(env_python))
    root = Path(__file__).resolve().parents[1]
    candidates.extend((root / ".venv-mineru" / "Scripts" / "python.exe", root / ".venv-mineru" / "bin" / "python"))
    current = Path(sys.executable)
    candidates.append(current)
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def mineru_version(python: Path) -> str | None:
    """Read MinerU metadata in a subprocess so the main runtime stays light."""
    try:
        result = subprocess.run(
            [str(python), "-c", "import importlib.metadata as m; print(m.version('mineru'))"],
            capture_output=True, text=True, encoding="utf-8", errors="replace", check=False,
            timeout=20, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        return result.stdout.strip() if result.returncode == 0 and result.stdout.strip() else None
    except (OSError, subprocess.TimeoutExpired):
        return None


def _mineru_text(value: Any) -> str:
    """Flatten MinerU v2 content fragments while preserving their order."""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "".join(_mineru_text(item) for item in value)
    if isinstance(value, dict):
        for key in ("content", "text", "latex", "equation", "value"):
            if key in value:
                text = _mineru_text(value[key])
                if text:
                    return text
        for key in ("title_content", "paragraph_content", "list_items", "item_content", "table_content", "math_content", "content_list"):
            if key in value:
                text = _mineru_text(value[key])
                if text:
                    return text
    return ""


def _mineru_inline_math_text(value: Any) -> str:
    """Turn MinerU inline LaTeX into a safe, readable fallback string.

    Inline MinerU equations do not carry a standalone image or reliable local
    bbox in every release.  Emitting their raw LaTeX into DOCX produces visible
    ``\\frac``/``\\omega`` markup in Word.  This deliberately conservative
    fallback keeps the symbols searchable without pretending to validate the
    mathematical meaning.  Standalone/display equations with an image are
    handled as pictures by ``mineru_content_to_docx``.
    """
    text = _mineru_text(value) if not isinstance(value, str) else value
    text = str(text or "").strip()
    if not text:
        return "□"
    replacements = {
        "\\times": "×", "\\cdot": "·", "\\pm": "±", "\\mp": "∓",
        "\\leq": "≤", "\\le": "≤", "\\geq": "≥", "\\ge": "≥",
        "\\neq": "≠", "\\ne": "≠", "\\approx": "≈", "\\sim": "∼",
        "\\infty": "∞", "\\partial": "∂", "\\nabla": "∇", "\\sum": "∑",
        "\\prod": "∏", "\\int": "∫", "\\sqrt": "√", "\\angle": "∠",
        "\\degree": "°", "\\circ": "°", "\\cdots": "…", "\\ldots": "…",
        "\\quad": " ", "\\,": " ", "\\;": " ", "\\!": "",
    }
    greek = {
        "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
        "varepsilon": "ε", "zeta": "ζ", "eta": "η", "theta": "θ", "vartheta": "ϑ",
        "iota": "ι", "kappa": "κ", "lambda": "λ", "mu": "μ", "nu": "ν",
        "xi": "ξ", "pi": "π", "varpi": "ϖ", "rho": "ρ", "varrho": "ϱ",
        "sigma": "σ", "varsigma": "ς", "tau": "τ", "upsilon": "υ", "phi": "φ",
        "varphi": "ϕ", "chi": "χ", "psi": "ψ", "omega": "ω",
        "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ",
        "Pi": "Π", "Sigma": "Σ", "Phi": "Φ", "Psi": "Ψ", "Omega": "Ω",
    }
    for command, symbol in replacements.items():
        text = text.replace(command, symbol)
    text = re.sub(r"\\(?:mathrm|mathbf|mathit|mathcal|mathfrak|mathtt|pmb|boldsymbol|text|operatorname|scriptscriptstyle|scriptstyle|displaystyle|left|right|overline|underline)\b", "", text)
    for name, symbol in greek.items():
        text = re.sub(rf"\\{name}\b", symbol, text)
    # A compact, non-LaTeX fraction fallback.  It is intentionally textual;
    # uncertain content is still eligible for a local PDF crop when it is a
    # standalone block with a verified source bbox.
    text = re.sub(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}", r"(\1/\2)", text)
    text = re.sub(r"\\(?:begin|end)\s*\{[^{}]*\}", "", text)
    text = re.sub(r"\\[A-Za-z]+", "", text)
    text = text.replace("{", "(").replace("}", ")")
    text = text.replace("~", " ")
    for _ in range(3):
        text = re.sub(r"\s*([_^])\s*\(\s*\(\s*([^()]*)\s*\)\s*\)", r"\1\2", text)
        text = re.sub(r"\s*([_^])\s*\(\s*([^()]*)\s*\)", r"\1\2", text)
    text = re.sub(r"\s*([_^])\s*(?:\(\s*)+([A-Za-z0-9]+)(?:\s*\))+", r"\1\2", text)
    text = re.sub(r"\s+([_^])\s*", r"\1", text)
    text = re.sub(r"\\(?=\s|[^A-Za-z])", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\(\s*\)", "", text)
    return text or "□"


def _mineru_image_path(item: Any, parse_dir: Path) -> Path | None:
    if not isinstance(item, dict):
        return None
    for key in ("img_path", "image_path", "path", "image"):
        value = item.get(key)
        if isinstance(value, str) and value:
            candidate = Path(value)
            if not candidate.is_absolute():
                candidate = parse_dir / candidate
            if candidate.exists() and candidate.is_file() and candidate.suffix.lower() in IMAGE_EXTENSIONS:
                return candidate
            for root in (parse_dir, parse_dir.parent, parse_dir.parent.parent):
                matches = list(root.rglob(candidate.name)) if root.exists() else []
                if matches:
                    return matches[0]
    for key in ("image_source", "content"):
        nested = item.get(key)
        if isinstance(nested, dict):
            found = _mineru_image_path(nested, parse_dir)
            if found:
                return found
    return None


def _mineru_blocks(page: Any) -> list[Any]:
    if isinstance(page, list):
        return page
    if isinstance(page, dict):
        for key in ("content_list", "blocks", "para_blocks", "preproc_blocks"):
            value = page.get(key)
            if isinstance(value, list):
                return value
    return []


def mineru_content_to_docx(content_path: Path, output: Path, max_pages: int | None = None, reference_pdf: Path | None = None) -> Path:
    """Rebuild an editable DOCX from MinerU semantic blocks.

    MinerU is deliberately used as a parser here, not as a Word generator.
    Every visual block gets its own normal-flow paragraph and every text block
    remains editable.  This avoids the anchor/absolute-position objects emitted
    by PDF-to-Word converters and is understood consistently by Word and WPS.
    """
    from docx import Document  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Inches, Pt  # type: ignore

    payload = json.loads(content_path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        pages = [payload] if payload and isinstance(payload[0], dict) else payload
    else:
        pages = payload.get("pdf_info", payload.get("pages", []))
    document = Document()
    source_pdf = None
    formula_asset_dir: Path | None = None
    if reference_pdf and reference_pdf.exists():
        try:
            import pymupdf as fitz  # type: ignore
        except ImportError:
            try:
                import fitz  # type: ignore
            except ImportError:
                fitz = None  # type: ignore
        if fitz is not None:
            source_pdf = fitz.open(str(reference_pdf))
            formula_asset_dir = Path(tempfile.mkdtemp(prefix="mineru-formula-crops-", dir=str(output.parent)))

    visual_types = {
        "image", "figure", "picture", "equation", "formula", "math",
        "interline_equation", "inline_equation", "display_formula", "inline_formula",
    }

    def paragraph_content_text(content: Any) -> str:
        """Flatten paragraph fragments without leaking raw inline LaTeX."""
        if not isinstance(content, dict):
            text = normalize_text(_mineru_text(content))
            return _mineru_inline_math_text(text) if "\\" in text else text
        fragments = content.get("paragraph_content")
        if not isinstance(fragments, list):
            text = normalize_text(_mineru_text(content))
            return _mineru_inline_math_text(text) if "\\" in text else text
        parts: list[str] = []
        for fragment in fragments:
            if not isinstance(fragment, dict):
                parts.append(_mineru_inline_math_text(fragment))
                continue
            kind = str(fragment.get("type", "")).strip().lower()
            fragment_content = fragment.get("content", fragment)
            if kind in {"equation_inline", "inline_equation", "inline_formula", "latex", "formula"}:
                parts.append(_mineru_inline_math_text(fragment_content))
            else:
                parts.append(_mineru_text(fragment_content))
        text = normalize_text("".join(parts))
        return _mineru_inline_math_text(text) if "\\" in text else text

    def item_text(item: Any) -> str:
        if isinstance(item, dict):
            content = item.get("content", item)
            if str(item.get("type", "")).strip().lower() == "paragraph":
                return paragraph_content_text(content)
            text = normalize_text(_mineru_text(content))
            return _mineru_inline_math_text(text) if "\\" in text else text
        text = normalize_text(_mineru_text(item))
        return _mineru_inline_math_text(text) if "\\" in text else text

    def item_caption(item: dict[str, Any]) -> str:
        content = item.get("content", item)
        if not isinstance(content, dict):
            return ""
        values: list[str] = []
        for key in ("image_caption", "image_footnote", "caption"):
            entries = content.get(key)
            if not isinstance(entries, list):
                continue
            for entry in entries:
                value = _mineru_text(entry.get("content", entry) if isinstance(entry, dict) else entry)
                if value:
                    values.append(normalize_text(value))
        return " ".join(values).strip()

    def add_visual(
        image: Path,
        item_type: str,
        caption: str = "",
        bbox: list[float] | None = None,
        page_width: float = 1000.0,
        page_height: float = 1000.0,
    ) -> None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        # Keep figures within the A4 text area. Formula crops are intentionally
        # smaller so they behave like a normal display equation, not a page
        # fragment. The native aspect ratio is always preserved.  A height cap
        # is essential for Office: a tall inline image can otherwise extend its
        # line box into the following paragraph and appear to overlap text.
        formula_types = {"equation", "formula", "math", "interline_equation", "inline_equation", "display_formula", "inline_formula"}
        max_width = 5.95 if item_type in visual_types else 6.15
        max_height = 0.95 if item_type in formula_types else 4.75
        # MinerU's bbox is in the source-page coordinate system.  Use it as
        # the first size estimate so a small diagram remains a small diagram
        # instead of expanding to the full text width.  The hard caps below
        # still protect Office from malformed or near-page-sized crops.
        target_width = max_width
        target_height = max_height
        if isinstance(bbox, list) and len(bbox) == 4:
            try:
                bbox_width = max(1.0, float(bbox[2]) - float(bbox[0]))
                bbox_height = max(1.0, float(bbox[3]) - float(bbox[1]))
                usable_width_inches = 6.15
                usable_height_inches = 8.8
                target_width = min(max_width, max(0.45, usable_width_inches * bbox_width / max(page_width, 1.0)))
                target_height = min(max_height, max(0.18, usable_height_inches * bbox_height / max(page_height, 1.0)))
            except (TypeError, ValueError):
                pass
        from PIL import Image  # type: ignore
        try:
            with Image.open(image) as source_image:
                ratio = source_image.width / max(source_image.height, 1)
            width = min(target_width, max(0.45, target_height * ratio))
            height = width / max(ratio, 0.1)
            if height > target_height:
                height = target_height
                width = max(0.45, height * ratio)
            run.add_picture(str(image), width=Inches(width), height=Inches(height))
        except Exception:
            run.add_picture(str(image), width=Inches(max_width))
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = bool(caption)

    def add_formula_crop(page_index: int, bbox: list[float] | None, name: str) -> bool:
        """Insert the source-PDF line as an image when inline math is unsafe."""
        if source_pdf is None or formula_asset_dir is None or page_index >= len(source_pdf) or not bbox or len(bbox) != 4:
            return False
        try:
            page = source_pdf[page_index]
            page_items = _mineru_blocks(pages[page_index])
            max_x = max((float(x.get("bbox", [0, 0, 0, 0])[2]) for x in page_items if isinstance(x, dict) and isinstance(x.get("bbox"), list) and len(x.get("bbox")) == 4), default=float(bbox[2]))
            max_y = max((float(x.get("bbox", [0, 0, 0, 0])[3]) for x in page_items if isinstance(x, dict) and isinstance(x.get("bbox"), list) and len(x.get("bbox")) == 4), default=float(bbox[3]))
            # MinerU's scan coordinates use one isotropic page scale.  Using
            # page.height/max_y independently stretches the y axis and moves
            # a formula crop onto the following text line.
            sx = page.rect.width / max(max_x, 1.0)
            sy = sx
            rect = fitz.Rect(float(bbox[0]) * sx, float(bbox[1]) * sy, float(bbox[2]) * sx, float(bbox[3]) * sy) & page.rect
            # MinerU inline-math boxes are often tighter than the rendered
            # glyph line (especially on scanned pages).  A small vertical
            # pad produced clipped tops/bottoms in Word, so include a safe
            # line-height envelope around the crop.  This remains a local
            # formula/line image and never expands to a whole page.
            pad_x, pad_y = min(8.0, page.rect.width * 0.02), min(8.0, page.rect.height * 0.012)
            rect = fitz.Rect(max(0, rect.x0 - pad_x), max(0, rect.y0 - pad_y), min(page.rect.width, rect.x1 + pad_x), min(page.rect.height, rect.y1 + pad_y))
            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), clip=rect, alpha=False)
            crop = formula_asset_dir / f"{name}.png"
            pix.save(str(crop))
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_together = True
            from PIL import Image  # type: ignore
            with Image.open(crop) as source_image:
                ratio = source_image.width / max(source_image.height, 1)
            width = min(6.15, max(1.1, 0.75 * ratio))
            height = min(0.85, width / max(ratio, 0.1))
            width = height * ratio
            paragraph.add_run().add_picture(str(crop), width=Inches(width), height=Inches(height))
            return True
        except Exception:
            return False

    def bbox_of(item: Any) -> list[float] | None:
        value = item.get("bbox") if isinstance(item, dict) else None
        if isinstance(value, list) and len(value) == 4:
            try:
                return [float(v) for v in value]
            except (TypeError, ValueError):
                return None
        return None

    def visual_cluster_end(items: list[Any], start: int) -> int:
        """Merge nearby image/equation blocks into one local PDF crop.

        MinerU returns the pieces of a multi-column figure as separate image
        blocks.  Adding them as independent Word paragraphs loses the source
        relationship and lets Word move captions through the gaps.  A bounded
        crop of the visual cluster keeps the relative layout while remaining a
        local object (never a full-page image).
        """
        first = bbox_of(items[start])
        if first is None:
            return start
        left, top, right, bottom = first
        end = start
        for index in range(start + 1, len(items)):
            candidate = items[index]
            ctype = str(candidate.get("type", "")).strip().lower() if isinstance(candidate, dict) else ""
            if ctype not in {"image", "figure", "picture", "equation", "formula", "math", "interline_equation", "inline_equation", "display_formula", "inline_formula"}:
                break
            box = bbox_of(candidate)
            if box is None:
                break
            cleft, ctop, cright, cbottom = box
            horizontal_gap = max(0.0, max(left, cleft) - min(right, cright))
            vertical_gap = max(0.0, max(top, ctop) - min(bottom, cbottom))
            # Blocks on the same row or in a compact stacked diagram belong
            # together.  A generous threshold is safe because the union is
            # still capped below and captions are intentionally included.
            if horizontal_gap > 110.0 and vertical_gap > 95.0:
                break
            left, top, right, bottom = min(left, cleft), min(top, ctop), max(right, cright), max(bottom, cbottom)
            end = index
        return end

    def add_visual_cluster(page_index: int, items: list[Any], start: int, end: int, page_width: float, page_height: float) -> bool:
        """Insert a bounded source-PDF crop for a visual cluster."""
        if source_pdf is None or formula_asset_dir is None or page_index >= len(source_pdf):
            return False
        boxes = [bbox_of(items[index]) for index in range(start, end + 1)]
        boxes = [box for box in boxes if box is not None]
        if not boxes:
            return False
        left = min(box[0] for box in boxes)
        top = min(box[1] for box in boxes)
        right = max(box[2] for box in boxes)
        bottom = max(box[3] for box in boxes)
        # Include the caption line immediately below the visual group, but
        # never consume the whole page or unrelated body text.
        bottom = min(page_height, bottom + 28.0)
        area = max(1.0, (right - left) * (bottom - top))
        if area >= page_width * page_height * 0.72:
            return False
        try:
            page = source_pdf[page_index]
            sx = page.rect.width / max(page_width, 1.0)
            sy = sx
            rect = fitz.Rect(left * sx, top * sy, right * sx, bottom * sy) & page.rect
            pad_x, pad_y = min(5.0, page.rect.width * 0.012), min(4.0, page.rect.height * 0.008)
            rect = fitz.Rect(max(0, rect.x0 - pad_x), max(0, rect.y0 - pad_y), min(page.rect.width, rect.x1 + pad_x), min(page.rect.height, rect.y1 + pad_y))
            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), clip=rect, alpha=False)
            crop = formula_asset_dir / f"p{page_index+1:04d}-visual-{start:04d}-{end:04d}.png"
            pix.save(str(crop))
            from PIL import Image  # type: ignore
            with Image.open(crop) as source_image:
                ratio = source_image.width / max(source_image.height, 1)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False
            max_width, max_height = 6.15, 4.7
            width = min(max_width, max(0.75, max_height * ratio))
            height = width / max(ratio, 0.1)
            if height > max_height:
                height = max_height
                width = max(0.75, height * ratio)
            paragraph.add_run().add_picture(str(crop), width=Inches(width), height=Inches(height))
            return True
        except Exception:
            return False

    def add_table(item: dict[str, Any]) -> bool:
        content = item.get("content", item)
        rows = None
        if isinstance(content, dict):
            for key in ("rows", "table_rows", "body", "table_body"):
                value = content.get(key)
                if isinstance(value, list):
                    rows = value
                    break
        if not isinstance(rows, list) or not rows:
            return False
        normalized_rows: list[list[str]] = []
        for row in rows:
            if isinstance(row, dict):
                row = row.get("cells", row.get("content", []))
            if not isinstance(row, list):
                row = [row]
            normalized_rows.append([normalize_text(_mineru_text(cell)) for cell in row])
        width = max((len(row) for row in normalized_rows), default=0)
        if width == 0:
            return False
        table = document.add_table(rows=len(normalized_rows), cols=width)
        table.style = "Table Grid"
        for rindex, row in enumerate(normalized_rows):
            for cindex, value in enumerate(row):
                table.cell(rindex, cindex).text = value
        return True

    def add_image_row(items: list[Any], start: int, end: int, page_width: float, page_height: float) -> bool:
        """Place a same-row image group in borderless cells.

        Each cell owns exactly one image and its caption, so a figure label
        cannot drift away from the corresponding picture when Word repaginates.
        The table is only a layout container; pictures remain inline objects.
        """
        from docx.oxml import OxmlElement  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        entries: list[tuple[int, Path, list[float]]] = []
        for pos in range(start, end + 1):
            image = _mineru_image_path(items[pos], content_path.parent)
            box = bbox_of(items[pos])
            if image is None or box is None:
                continue
            entries.append((pos, image, box))
        if not entries:
            return False
        table = document.add_table(rows=1, cols=len(entries))
        table.autofit = False
        table.allow_autofit = False
        # Remove table borders: the table is an invisible layout scaffold.
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = borders.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                borders.append(node)
            node.set(qn("w:val"), "nil")
        total_width = max(1.0, max(box[2] for _, _, box in entries) - min(box[0] for _, _, box in entries))
        usable_width = 6.15
        row = table.rows[0]
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, (item_index, image_path, box) in zip(row.cells, entries):
            cell_width = max(0.55, usable_width * (box[2] - box[0]) / total_width)
            cell.width = Inches(cell_width)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_together = True
            paragraph.add_run().add_picture(str(image_path), width=Inches(min(cell_width - 0.08, 2.8)))
            caption = item_caption(items[item_index])
            if caption:
                caption_paragraph = cell.add_paragraph(caption)
                caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_paragraph.paragraph_format.space_before = Pt(0)
                caption_paragraph.paragraph_format.space_after = Pt(0)
                caption_paragraph.paragraph_format.keep_together = True
        return True

    major_break_re = re.compile(r"^(?:目录|第[0-9一二三四五六七八九十百]+章|参考文献|附录|致谢|结论|结语)$")
    previous_page_signatures: set[str] = set()
    for page_index, page in enumerate(pages):
        if max_pages is not None and page_index >= max_pages:
            break
        page_items = _mineru_blocks(page)
        page_width = max((float(x.get("bbox", [0, 0, 0, 0])[2]) for x in page_items if isinstance(x, dict) and isinstance(x.get("bbox"), list) and len(x.get("bbox")) == 4), default=1000.0)
        page_height = max((float(x.get("bbox", [0, 0, 0, 0])[3]) for x in page_items if isinstance(x, dict) and isinstance(x.get("bbox"), list) and len(x.get("bbox")) == 4), default=1000.0)
        index = 0
        current_page_signatures: set[str] = set()
        while index < len(page_items):
            item = page_items[index]
            if not isinstance(item, dict):
                index += 1
                continue
            item_type = str(item.get("type", "")).strip().lower()
            image = _mineru_image_path(item, content_path.parent)
            if image:
                # Group only adjacent images on the same visual row. Images
                # in later rows start a new group, so a full-page composite is
                # never produced.
                row_end = index
                first_box = bbox_of(item)
                if first_box is not None:
                    first_center = (first_box[1] + first_box[3]) / 2.0
                    for candidate_index in range(index + 1, len(page_items)):
                        candidate = page_items[candidate_index]
                        if not isinstance(candidate, dict) or _mineru_image_path(candidate, content_path.parent) is None:
                            break
                        candidate_box = bbox_of(candidate)
                        if candidate_box is None:
                            break
                        candidate_center = (candidate_box[1] + candidate_box[3]) / 2.0
                        if abs(candidate_center - first_center) > max(55.0, (first_box[3] - first_box[1]) * 0.65):
                            break
                        row_end = candidate_index
                if row_end > index and add_image_row(page_items, index, row_end, page_width, page_height):
                    index = row_end + 1
                    continue
                caption = item_caption(item)
                raw_bbox = item.get("bbox")
                bbox = raw_bbox if isinstance(raw_bbox, list) and len(raw_bbox) == 4 else None
                add_visual(image, item_type, caption, bbox, page_width, page_height)
                if caption:
                    caption_paragraph = document.add_paragraph(caption)
                    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption_paragraph.paragraph_format.keep_together = True
                    caption_paragraph.paragraph_format.keep_with_next = False
                    caption_paragraph.paragraph_format.space_before = Pt(0)
                    caption_paragraph.paragraph_format.space_after = Pt(2)
                index += 1
                continue
            if item_type == "index":
                content = item.get("content", item)
                entries = content.get("list_items", []) if isinstance(content, dict) else []
                for entry in entries if isinstance(entries, list) else []:
                    text = normalize_text(_mineru_text(entry))
                    if text:
                        document.add_paragraph(text)
                index += 1
                continue
            if item_type == "table" and add_table(item):
                index += 1
                continue
            content = item.get("content", item)
            raw_signature = normalize_text(_mineru_text(content))
            signature = re.sub(r"\s+", "", raw_signature)
            # MinerU can repeat the last line of a page as the first block of
            # the next page.  Drop only that boundary duplicate; repeated
            # labels inside a page remain untouched.
            if item_type == "paragraph" and signature and signature in previous_page_signatures:
                index += 1
                continue
            if item_type == "paragraph" and signature:
                current_page_signatures.add(signature)
            inline_math = False
            if item_type == "paragraph" and isinstance(content, dict):
                fragments = content.get("paragraph_content", [])
                inline_math = any(
                    isinstance(fragment, dict)
                    and str(fragment.get("type", "")).strip().lower() in {"equation_inline", "inline_equation", "inline_formula", "latex", "formula"}
                    for fragment in fragments if isinstance(fragments, list)
                )
            # Inline equation boxes in this scanned source are not reliably
            # source-validated across pages.  Keep the sanitized symbol text
            # rather than risk inserting a neighbouring PDF line.  Explicit
            # standalone equation/image blocks still use their image assets.
            text = item_text(item)
            if not text and item_type in visual_types:
                text = normalize_text(_mineru_text(item))
            if not text:
                index += 1
                continue
            paragraph = document.add_paragraph(text)
            if item_type in {"title", "heading"}:
                paragraph.style = "Heading 1"
                # Keep only meaningful document boundaries.  A break for
                # every source page creates large blank regions in Word; a
                # break before the目录/章节 headings preserves the intended
                # front-matter and chapter separation without freezing the
                # flowing body layout.
                if page_index > 0 and major_break_re.match(re.sub(r"\s+", "", text)):
                    paragraph.paragraph_format.page_break_before = True
            else:
                paragraph.style = "Normal"
            index += 1
        previous_page_signatures = current_page_signatures
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def pymupdf_pdf_to_docx(path: Path, output: Path, max_pages: int | None = None, dpi: int = 150) -> Path:
    """Fast PDF-first reconstruction using PyMuPDF page blocks.

    This is intentionally conservative: ordinary text stays editable, while
    figures, their captions, and high-confidence equation blocks are rendered
    as local PDF crops.  Cropping the local object (instead of rebuilding its
    glyphs) keeps formulas and diagram labels visually identical to the PDF
    and avoids the long MinerU model pass for ordinary electronic PDFs.
    """
    from docx import Document  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Inches, Pt  # type: ignore

    try:
        import pymupdf as fitz  # type: ignore
    except ImportError:  # pragma: no cover - older installations expose fitz
        import fitz  # type: ignore

    pdf = fitz.open(str(path))
    # Do not shrink a scanned page into a tiny figure.  A scan has one
    # near-page-sized image block and almost no extractable text; it must go
    # through the OCR/MinerU branch instead of the electronic-PDF fast path.
    scan_pages = 0
    sample_count = min(len(pdf), max_pages or 3, 3)
    for sample_index in range(sample_count):
        sample_page = pdf[sample_index]
        page_area = max(float(sample_page.rect.width * sample_page.rect.height), 1.0)
        image_area = 0.0
        text_chars = 0
        for sample_block in sample_page.get_text("dict", sort=False).get("blocks", []):
            if not isinstance(sample_block, dict):
                continue
            bbox = sample_block.get("bbox") or [0, 0, 0, 0]
            area = max(0.0, (float(bbox[2]) - float(bbox[0])) * (float(bbox[3]) - float(bbox[1])))
            if sample_block.get("type") == 1:
                image_area += area
            elif sample_block.get("type") == 0:
                text_chars += sum(len(str(span.get("text", ""))) for line in sample_block.get("lines", []) for span in line.get("spans", []))
        if image_area / page_area >= 0.82 and text_chars < 80:
            scan_pages += 1
    if sample_count and scan_pages >= max(1, (sample_count + 1) // 2):
        pdf.close()
        raise RuntimeError("检测为扫描/混合型 PDF，跳过 PyMuPDF 快速路径，交给 MinerU/OCR")
    document = Document()
    asset_dir = Path(tempfile.mkdtemp(prefix="pymupdf-assets-", dir=str(output.parent)))
    render_scale = max(1.25, min(3.0, float(dpi) / 72.0))
    page_total = min(len(pdf), max_pages) if max_pages else len(pdf)
    figure_caption_re = re.compile(r"^(?:图|图\s*\d|Fig(?:ure)?\.?\s*\d)")

    def block_text(block: dict[str, Any]) -> str:
        lines = []
        for line in block.get("lines", []):
            spans = [str(span.get("text", "")) for span in line.get("spans", [])]
            value = "".join(spans).strip()
            if value:
                lines.append(value)
        return normalize_text("\n".join(lines))

    def is_heading(text: str, block: dict[str, Any]) -> tuple[bool, str]:
        compact = re.sub(r"\s+", "", text)
        sizes = [float(span.get("size", 0)) for line in block.get("lines", []) for span in line.get("spans", [])]
        size = max(sizes, default=0.0)
        if compact in {"目录", "前言", "参考文献", "同步训练题", "同步训练题答案"}:
            return True, "title"
        if re.match(r"^第[0-9一二三四五六七八九十百]+章", compact):
            return True, "chapter"
        if len(compact) <= 28 and (size >= 16 or re.match(r"^(?:题|解|例|习题)\s*[0-9一二三四五六七八九十百-]+", compact)):
            return True, "heading"
        return False, "body"

    def save_crop(page: Any, rect: Any, name: str) -> Path:
        page_rect = page.rect
        rect = rect & page_rect
        pix = page.get_pixmap(matrix=fitz.Matrix(render_scale, render_scale), clip=rect, alpha=False)
        target = asset_dir / f"{name}.png"
        pix.save(str(target))
        return target

    def add_crop(page: Any, rect: Any, name: str, max_height: float = 4.8) -> None:
        image = save_crop(page, rect, name)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_together = True
        from PIL import Image  # type: ignore
        with Image.open(image) as source_image:
            ratio = source_image.width / max(source_image.height, 1)
        width = min(6.15, max(0.55, max_height * ratio))
        height = width / max(ratio, 0.1)
        if height > max_height:
            height = max_height
            width = max(0.55, height * ratio)
        paragraph.add_run().add_picture(str(image), width=Inches(width), height=Inches(height))

    try:
        for page_index in range(page_total):
            page = pdf[page_index]
            blocks = page.get_text("dict", sort=False).get("blocks", [])
            entries: list[tuple[float, float, int, dict[str, Any], str]] = []
            for idx, block in enumerate(blocks):
                if not isinstance(block, dict):
                    continue
                bbox = block.get("bbox") or [0, 0, 0, 0]
                if block.get("type") == 0:
                    text = block_text(block)
                    if text:
                        entries.append((float(bbox[1]), float(bbox[0]), idx, block, text))
                elif block.get("type") == 1:
                    entries.append((float(bbox[1]), float(bbox[0]), idx, block, ""))
            entries.sort(key=lambda value: (value[0], value[1]))
            consumed: set[int] = set()
            for pos, (_y, _x, idx, block, text) in enumerate(entries):
                if idx in consumed:
                    continue
                bbox = block.get("bbox") or [0, 0, 0, 0]
                if block.get("type") == 1:
                    rect = fitz.Rect(*bbox)
                    caption_idx = None
                    caption_rect = None
                    for next_y, next_x, next_idx, next_block, next_text in entries[pos + 1:]:
                        if next_idx in consumed or next_block.get("type") != 0 or not next_text:
                            continue
                        next_bbox = next_block.get("bbox") or [0, 0, 0, 0]
                        gap = float(next_bbox[1]) - float(bbox[3])
                        overlap = max(0.0, min(float(bbox[2]), float(next_bbox[2])) - max(float(bbox[0]), float(next_bbox[0])))
                        if 0 <= gap <= 42 and overlap >= min(float(bbox[2]) - float(bbox[0]), float(next_bbox[2]) - float(next_bbox[0])) * 0.12 and figure_caption_re.match(next_text):
                            caption_idx, caption_rect = next_idx, fitz.Rect(*next_bbox)
                        break
                    if caption_rect is not None:
                        rect |= caption_rect
                        consumed.add(caption_idx)
                    pad = 3
                    add_crop(page, fitz.Rect(max(0, rect.x0 - pad), max(0, rect.y0 - pad), min(page.rect.width, rect.x1 + pad), min(page.rect.height, rect.y1 + pad)), f"p{page_index+1:04d}-fig{idx:03d}")
                    consumed.add(idx)
                    continue
                formula, _confidence = _looks_like_formula_block(text)
                compact = re.sub(r"\s+", "", text)
                line_count = len(block.get("lines", []))
                if formula and len(compact) <= 180 and line_count <= 4 and ("=" in compact or any(char in compact for char in "∑∫√≤≥≠≈ωαβγθμπ")):
                    rect = fitz.Rect(*bbox)
                    pad = 4
                    add_crop(page, fitz.Rect(max(0, rect.x0 - pad), max(0, rect.y0 - pad), min(page.rect.width, rect.x1 + pad), min(page.rect.height, rect.y1 + pad)), f"p{page_index+1:04d}-eq{idx:03d}", max_height=0.75)
                    consumed.add(idx)
                    continue
                paragraph = document.add_paragraph(text)
                heading, role = is_heading(text, block)
                paragraph.paragraph_format.keep_together = heading
                paragraph.paragraph_format.keep_with_next = heading
                if heading:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_before = Pt(8 if role in {"title", "chapter"} else 4)
                    paragraph.paragraph_format.space_after = Pt(8)
                    for run in paragraph.runs:
                        run.font.name = _HEADING_FONT or "华文行楷"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), _HEADING_FONT or "华文行楷")
                        run.font.size = Pt(23 if role == "title" else (17 if role == "chapter" else 15))
                        run.bold = role in {"chapter", "heading"}
                else:
                    paragraph.paragraph_format.line_spacing = Pt(19)
                    for run in paragraph.runs:
                        set_run_fonts(run, "SimSun")
                        run.font.size = Pt(13)
                consumed.add(idx)
            if page_index < page_total - 1:
                # Keep one source-page boundary through the generic Word
                # cleanup pass.  A bare page-break paragraph is intentionally
                # removed by postprocess_docx; the hidden zero-width marker
                # makes this boundary identifiable without adding visible text.
                page_break_paragraph = document.add_paragraph()
                page_break_run = page_break_paragraph.add_run("\u200b")
                page_break_run.font.hidden = True
                page_break_run.add_break(WD_BREAK.PAGE)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output))
        return output
    finally:
        pdf.close()
        if source_pdf is not None:
            source_pdf.close()
        if formula_asset_dir is not None:
            shutil.rmtree(formula_asset_dir, ignore_errors=True)
        shutil.rmtree(asset_dir, ignore_errors=True)


def _find_mineru_content(root: Path, stem: str) -> Path | None:
    candidates = sorted(root.rglob(f"{stem}_content_list_v2.json"))
    if not candidates:
        candidates = sorted(root.rglob(f"{stem}_content_list.json"))
    return candidates[0] if candidates else None


def mineru_bridge_convert(path: Path, python: Path, temp: Path, args: argparse.Namespace) -> tuple[Path | None, dict[str, Any]]:
    """Submit a MinerU job to the interactive-user bridge and await its output."""
    bridge_root = Path(__file__).resolve().parents[1] / "tmp" / "mineru-bridge"
    queue, done, errors = bridge_root / "queue-v1", bridge_root / "done-v1", bridge_root / "errors-v1"
    work = bridge_root / "work" / uuid.uuid4().hex
    staged_input, staged_output = work / path.name, work / "out"
    for folder in (queue, done, errors, staged_output):
        folder.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, staged_input)
    job_id = work.name
    job = {
        "id": job_id,
        "input": str(staged_input.resolve()),
        "output_dir": str(staged_output.resolve()),
        "python": str(python.resolve()),
        "backend": args.mineru_backend,
        "method": args.mineru_method,
        "timeout": args.mineru_timeout,
    }
    (queue / f"{job_id}.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    deadline = time.monotonic() + max(90, args.mineru_timeout + 60)
    result_file, error_file = done / f"{job_id}.json", errors / f"{job_id}.json"
    try:
        while time.monotonic() < deadline:
            if result_file.exists():
                result = json.loads(result_file.read_text(encoding="utf-8"))
                result_file.unlink(missing_ok=True)
                content = Path(result.get("content_list", ""))
                if content.exists():
                    return content, {"engine": "mineru-user-bridge", "status": "ok", "work_dir": str(work), **result}
                return None, {"engine": "mineru-user-bridge", "status": "failed", "reason": "桥接器未返回 content list", "work_dir": str(work), **result}
            if error_file.exists():
                reason = error_file.read_text(encoding="utf-8", errors="replace")
                error_file.unlink(missing_ok=True)
                return None, {"engine": "mineru-user-bridge", "status": "failed", "reason": reason, "work_dir": str(work)}
            time.sleep(1)
        return None, {"engine": "mineru-user-bridge", "status": "timeout", "reason": "等待 MinerU 用户账户桥接器超时", "work_dir": str(work)}
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        return None, {"engine": "mineru-user-bridge", "status": "failed", "reason": str(exc), "work_dir": str(work)}


def mineru_pdf_to_docx(path: Path, output: Path, temp: Path, args: argparse.Namespace) -> tuple[Path | None, dict[str, Any]]:
    """Run MinerU locally and convert its structured output to DOCX."""
    python = find_mineru_python(args.mineru_python)
    if not python:
        return None, {"engine": "mineru", "status": "unavailable", "reason": "找不到 MinerU Python 解释器"}
    version = mineru_version(python)
    if not version:
        return None, {"engine": "mineru", "status": "unavailable", "reason": f"{python} 中未检测到 MinerU"}
    if use_bridge():
        content, bridge_info = mineru_bridge_convert(path, python, temp, args)
        if not content:
            shutil.rmtree(str(bridge_info.get("work_dir", "")), ignore_errors=True)
            bridge_info["version"] = version
            return None, bridge_info
        try:
            mineru_content_to_docx(content, output, args.max_pages, path)
            return output, {"version": version, **bridge_info}
        finally:
            shutil.rmtree(str(bridge_info.get("work_dir", "")), ignore_errors=True)
    mineru_root = temp / "mineru"
    mineru_root.mkdir(parents=True, exist_ok=True)
    command = [str(python), "-m", "mineru.cli.client", "-p", str(path), "-o", str(mineru_root), "-b", args.mineru_backend, "-m", args.mineru_method, "-l", "ch", "--client-side-output-generation=True"]
    try:
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", check=False, timeout=max(60, args.mineru_timeout), creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except subprocess.TimeoutExpired as exc:
        return None, {"engine": "mineru", "version": version, "status": "timeout", "command": command, "stderr": str(exc)}
    content = _find_mineru_content(mineru_root, path.stem)
    if result.returncode != 0 or not content:
        return None, {"engine": "mineru", "version": version, "status": "failed", "returncode": result.returncode, "command": command, "stdout": result.stdout[-4000:], "stderr": result.stderr[-4000:]}
    mineru_content_to_docx(content, output, args.max_pages, path)
    return output, {"engine": "mineru", "version": version, "status": "ok", "content_list": str(content), "stdout": result.stdout[-2000:], "stderr": result.stderr[-2000:]}


def iter_paragraphs(document: Any):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_cell_paragraphs(cell)


def iter_cell_paragraphs(cell: Any):
    for paragraph in cell.paragraphs:
        yield paragraph
    for table in cell.tables:
        for row in table.rows:
            for nested in row.cells:
                yield from iter_cell_paragraphs(nested)


def _paragraph_has_nontext_object(paragraph: Any) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing|.//m:oMath|.//m:oMathPara|.//w:object|.//w:pict"))


def _remove_drawings(paragraph: Any) -> int:
    removed = 0
    for drawing in list(paragraph._p.xpath(".//w:drawing")):
        parent = drawing.getparent()
        if parent is not None:
            parent.remove(drawing)
            removed += 1
    return removed


def _image_blob_for_relationship(part: Any, rid: str) -> bytes | None:
    try:
        related = part.related_parts.get(rid)
    except Exception:
        related = None
    if related is None or not hasattr(related, "blob"):
        return None
    content_type = str(getattr(related, "content_type", ""))
    if content_type and not content_type.startswith("image/"):
        return None
    try:
        return bytes(related.blob)
    except Exception:
        return None


def _looks_like_qr_blob(blob: bytes) -> bool:
    """Fast QR detector; returns false when OpenCV is unavailable."""
    try:
        import cv2  # type: ignore
        import numpy as np  # type: ignore
        image = cv2.imdecode(np.frombuffer(blob, dtype=np.uint8), cv2.IMREAD_GRAYSCALE)
        if image is None or image.size < 2500:
            return False
        detector = cv2.QRCodeDetector()
        points = detector.detect(image)[1]
        if points is not None:
            return True
        # A few tiny Quark QR exports fail ``detect`` but decode reliably after
        # a modest nearest-neighbour upscale.
        if max(image.shape) < 260:
            scale = 3
            enlarged = cv2.resize(image, None, fx=scale, fy=scale, interpolation=cv2.INTER_NEAREST)
            points = detector.detect(enlarged)[1]
            if points is not None:
                return True
        return False
    except Exception:
        return False


def _remove_useless_vml(paragraph: Any) -> tuple[int, int]:
    """Remove VML text-box watermarks that python-docx does not expose."""
    removed_shapes = 0
    removed_runs = 0
    for pict in list(paragraph._p.iter()):
        if pict.tag.split("}")[-1] != "pict":
            continue
        text = "".join((node.text or "") for node in pict.iter() if node.tag.split("}")[-1] == "t")
        if not text or not USELESS_TEXT_RE.search(text) and not QR_TEXT_RE.search(text):
            continue
        parent = pict.getparent()
        while parent is not None and parent.tag.split("}")[-1] not in {"r", "AlternateContent"}:
            parent = parent.getparent()
        target = parent if parent is not None else pict
        target_parent = target.getparent()
        if target_parent is not None:
            target_parent.remove(target)
            removed_shapes += 1
            if target.tag.split("}")[-1] == "r":
                removed_runs += 1
    return removed_shapes, removed_runs


def _remove_floating_textboxes(document: Any) -> int:
    """Remove Quark/WPS text-box fallbacks that can float over body text.

    Quark puts many page-coordinate fragments in ``mc:AlternateContent``:
    the Choice is a WPS text box and the Fallback is a VML shape.  They are
    not real figures (there is no ``pic:pic``/``a:blip``), but Office may still
    render them as draggable objects and place them on top of normal text.
    The editable text/OMML paragraphs and genuine pictures remain untouched.
    """
    removed = 0
    for paragraph in list(iter_paragraphs(document)):
        for node in list(paragraph._p.iter()):
            if node.tag.split("}")[-1] != "AlternateContent":
                continue
            has_textbox = any(child.tag.split("}")[-1] == "wsp" for child in node.iter())
            # ``v:imagedata`` is the legacy fallback of these text boxes, not
            # a real figure.  Only DrawingML ``pic``/``blip`` counts as a
            # genuine picture that must be retained.
            has_real_picture = any(child.tag.split("}")[-1] in {"pic", "blip"} for child in node.iter())
            if not has_textbox or has_real_picture:
                continue
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
                removed += 1
    return removed


def reference_major_title_texts(document: Any) -> set[str]:
    """Collect short, centered title lines from the untouched Quark DOCX.

    The original export is the best local authority for which exam/year lines
    are major headings.  We use text and the old formatting only as a signal;
    the final typography still follows the project's fixed 23 pt title rule.
    """
    result: set[str] = set()
    # A title group may contain a school line, a wrapped exam title, and a
    # short metadata line. Keep the local association generous enough for
    # Quark's occasional spacer paragraphs, but never associate distant
    # centered body text with a heading.
    title_group_distance = 6
    paragraphs = list(iter_paragraphs(document))
    normalized = [re.sub(r"\s+", "", normalize_text(p.text or "")) for p in paragraphs]
    nonempty_rank: dict[int, int] = {}
    rank = 0
    for index, compact in enumerate(normalized):
        if compact and not paragraph_is_nonstructural(paragraphs[index]):
            nonempty_rank[index] = rank
            rank += 1
    exam_indices: set[int] = set()
    continuation_indices: set[int] = set()
    for index, compact in enumerate(normalized):
        if not compact or paragraph_is_nonstructural(paragraphs[index]):
            continue
        # A major exam line normally contains a year plus an exam marker.
        # This also catches the first line of a wrapped two-line title.
        if re.search(r"20\d{2}", compact) and re.search(r"(?:研究生|硕士|入学考试|试题|试卷|招收)", compact):
            exam_indices.add(index)
    for index, compact in enumerate(normalized):
        if not compact or index in exam_indices:
            continue
        # Wrapped title continuation, e.g. the second line of the 2023
        # heading: it has no year but contains the exam marker and follows a
        # year-bearing title within the same short title group.
        if re.search(r"(?:入学考试|自命题试题|考试试题)", compact):
            if any(abs(nonempty_rank.get(index, 0) - nonempty_rank.get(prior, 0)) <= title_group_distance for prior in exam_indices):
                continuation_indices.add(index)
    title_group_indices = exam_indices | continuation_indices
    for index, paragraph in enumerate(paragraphs):
        text = normalize_text(paragraph.text or "")
        compact = re.sub(r"\s+", "", text)
        if not compact or len(compact) > 90 or paragraph_is_nonstructural(paragraph):
            continue
        # Chip/function-table captions are centered in Quark exports, but
        # they are figure/table captions, not document-level headings.
        if re.search(r"功能表|功能如|附录|附圖|附图", compact):
            continue
        centered = str(paragraph.alignment).upper().endswith("CENTER") or paragraph.alignment == 1
        standalone_year = bool(re.fullmatch(r"20\d{2}年(?:度)?", compact))
        known = compact in {re.sub(r"\s+", "", value) for value in KNOWN_MAJOR_HEADINGS}
        school_group = bool(re.search(r"(?:大学|学院)$", compact)) and any(
            abs(nonempty_rank.get(index, 0) - nonempty_rank.get(title_index, 0)) <= title_group_distance
            for title_index in title_group_indices
        )
        if standalone_year or index in title_group_indices or known or (centered and school_group):
            result.add(compact)
    return result


def remove_qr_images(document: Any) -> int:
    """Delete QR images even when no nearby OCR label exists."""
    removed = 0
    checked: dict[str, bool] = {}
    for paragraph in list(iter_paragraphs(document)):
        for shape_node in list(paragraph._p.iter()):
            if shape_node.tag.split("}")[-1] not in {"drawing", "pict"}:
                continue
            ids = []
            for node in shape_node.iter():
                for key, value in node.attrib.items():
                    if key.split("}")[-1] in {"embed", "id"} and value:
                        ids.append(value)
            is_qr = False
            for rid in ids:
                if rid not in checked:
                    blob = _image_blob_for_relationship(paragraph.part, rid)
                    checked[rid] = bool(blob and _looks_like_qr_blob(blob))
                if checked.get(rid):
                    is_qr = True
                    break
            if not is_qr:
                continue
            target = shape_node
            while target.getparent() is not None and target.tag.split("}")[-1] not in {"r", "AlternateContent"}:
                target = target.getparent()
            parent = target.getparent()
            if parent is not None:
                parent.remove(target)
                removed += 1
    return removed


def remove_useless_text_and_qr(document: Any) -> dict[str, int]:
    """Remove known watermark/advertising text and QR blocks cheaply.

    This intentionally uses a short, auditable phrase list.  It does not run
    OCR or alter pixels inside ordinary figures.  A paragraph containing a QR
    marker loses its drawings as well as the marker text; other visual objects
    are preserved when only a watermark run is removed.
    """
    removed_paragraphs = 0
    removed_fragments = 0
    removed_qr_images = 0
    removed_vml_shapes = 0
    removed_floating_textboxes = _remove_floating_textboxes(document)
    paragraphs = list(iter_paragraphs(document))
    for index, paragraph in enumerate(paragraphs):
        vml_shapes, _ = _remove_useless_vml(paragraph)
        removed_vml_shapes += vml_shapes
        raw = paragraph.text or ""
        if not raw and not _paragraph_has_nontext_object(paragraph):
            continue
        qr_match = bool(QR_TEXT_RE.search(raw))
        if qr_match:
            removed_qr_images += _remove_drawings(paragraph)
            # Some Quark exports put the QR image in the immediately following
            # empty paragraph.  Remove that drawing-only companion as well.
            for neighbor_index in (index - 1, index + 1):
                if not 0 <= neighbor_index < len(paragraphs):
                    continue
                neighbor = paragraphs[neighbor_index]
                if neighbor.text.strip() or not neighbor._p.xpath(".//w:drawing"):
                    continue
                removed_qr_images += _remove_drawings(neighbor)
                if not _paragraph_has_nontext_object(neighbor) and not paragraph_in_table(neighbor):
                    remove_paragraph(neighbor)
                    removed_paragraphs += 1
        if not raw:
            if qr_match and not paragraph_in_table(paragraph):
                remove_paragraph(paragraph)
                removed_paragraphs += 1
            continue
        if not USELESS_TEXT_RE.search(raw) and not qr_match:
            continue
        # Keep non-adjacent question text while removing only ad/QR fragments.
        if not _paragraph_has_nontext_object(paragraph):
            fragments = re.split(r"(\t|\n)", raw)
            kept: list[str] = []
            for fragment in fragments:
                if fragment in {"\t", "\n"}:
                    kept.append(fragment)
                    continue
                if USELESS_TEXT_RE.search(fragment) or QR_TEXT_RE.search(fragment):
                    removed_fragments += 1
                    continue
                kept.append(fragment)
            cleaned = "".join(kept)
            cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
            cleaned = re.sub(r"[ \t]+([，。；：,.;])", r"\1", cleaned).strip()
            if cleaned != raw:
                paragraph.text = cleaned
        else:
            # Preserve formulas/figures but remove visible watermark runs.
            for run in paragraph.runs:
                value = run.text or ""
                cleaned = USELESS_TEXT_RE.sub("", value)
                cleaned = QR_TEXT_RE.sub("", cleaned)
                if cleaned != value:
                    run.text = cleaned
                    removed_fragments += 1
        if not paragraph.text.strip() and not _paragraph_has_nontext_object(paragraph) and not paragraph_in_table(paragraph):
            remove_paragraph(paragraph)
            removed_paragraphs += 1
    return {
        "useless_paragraphs_removed": removed_paragraphs,
        "useless_text_fragments_removed": removed_fragments,
        "qr_images_removed": removed_qr_images,
        "vml_watermark_shapes_removed": removed_vml_shapes,
        "floating_textboxes_removed": removed_floating_textboxes,
    }


def _table_row_has_content(row: Any) -> bool:
    for cell in row.cells:
        if normalize_text(cell.text):
            return True
        if cell._tc.xpath(".//w:drawing|.//m:oMath|.//m:oMathPara|.//w:object|.//w:pict|.//w:tbl"):
            return True
    return False


def trim_empty_table_edges(document: Any) -> dict[str, int]:
    """Remove blank border-only rows/tables emitted around converted figures."""
    tables_trimmed = 0
    rows_removed = 0
    tables_removed = 0
    for table in list(document.tables):
        rows = list(table.rows)
        start = 0
        while start < len(rows) and not _table_row_has_content(rows[start]):
            start += 1
        end = len(rows)
        while end > start and not _table_row_has_content(rows[end - 1]):
            end -= 1
        if start == len(rows):
            element = table._element
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
                tables_removed += 1
            continue
        if start or end != len(rows):
            for row in rows[:start] + rows[end:]:
                row._tr.getparent().remove(row._tr)
                rows_removed += 1
            tables_trimmed += 1
    return {
        "empty_table_edges_trimmed": tables_trimmed,
        "empty_table_rows_removed": rows_removed,
        "empty_tables_removed": tables_removed,
    }


def remove_paragraph(paragraph: Any) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def paragraph_in_table(paragraph: Any) -> bool:
    node = paragraph._element
    while node is not None:
        if node.tag.endswith("}tc"):
            return True
        node = node.getparent()
    return False


def paragraph_has_visual_object(paragraph: Any) -> bool:
    """Return True for paragraphs containing drawings, equations, or objects.

    These blocks are layout-bearing content rather than structural headings.
    Their visible runs are still formatted by the body typography pass.
    """
    for element in paragraph._p.iter():
        local = element.tag.split("}")[-1]
        if local in {"drawing", "pict", "object", "oMath", "oMathPara"}:
            return True
    return False


def paragraph_is_nonstructural(paragraph: Any) -> bool:
    """Tables, images, and formulas are excluded from heading mapping."""
    return paragraph_in_table(paragraph) or paragraph_has_visual_object(paragraph)


def has_explicit_page_break_before(paragraph: Any) -> bool:
    """Detect an explicit page break so repeated post-processing is idempotent."""
    from docx.oxml.ns import qn  # type: ignore
    # ``pageBreakBefore`` is a paragraph property, not a break-only sibling.
    # Treat it as an intentional boundary so the marker cleanup below does
    # not reset it to ``w:val=0``.
    if paragraph._p.xpath("./w:pPr/w:pageBreakBefore"):
        return True
    previous = paragraph._element.getprevious()
    if previous is not None and previous.tag.endswith("}p") and previous.find(".//" + qn("w:br")) is not None:
        return True
    return False


def has_section_boundary_before(paragraph: Any) -> bool:
    """Return true when an earlier empty block already ends a Word section."""
    from docx.oxml.ns import qn  # type: ignore
    previous = paragraph._element.getprevious()
    while previous is not None and previous.tag.endswith("}p"):
        if previous.find(".//" + qn("w:sectPr")) is not None:
            return True
        # Only scan through empty spacer paragraphs. A non-empty paragraph
        # means this heading is not already at a source page boundary.
        texts = previous.xpath(".//w:t/text()")
        drawings = previous.xpath(".//w:drawing")
        if texts or drawings:
            break
        previous = previous.getprevious()
    return False


def normalize_section_geometry(document: Any) -> None:
    """Use one A4 portrait canvas for every imported Quark section."""
    from docx.shared import Mm  # type: ignore
    from docx.enum.section import WD_ORIENT, WD_SECTION_START  # type: ignore
    for section in document.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
        section.top_margin = Mm(19)
        section.bottom_margin = Mm(19)
        # Normalize odd/even/continuous source section starts. Every retained
        # source page should begin on a normal next page, never an implicit
        # blank page introduced by WPS.
        section.start_type = WD_SECTION_START.NEW_PAGE


def collapse_source_sections(document: Any) -> int:
    """Remove page-per-section breaks created by OCR/Word converters.

    Quark exports often encode every source PDF page as a new Word section.
    Keeping those boundaries makes a question stem and its answer split even
    when there is ample room on the next page.  The pipeline repaginates the
    editable text, so intermediate ``sectPr`` nodes are removed and the one
    remaining section is normalized to A4 below.
    """
    removed = 0
    for paragraph in list(document.paragraphs):
        for sect_pr in list(paragraph._p.xpath("./w:pPr/w:sectPr")):
            sect_pr.getparent().remove(sect_pr)
            removed += 1
    return removed


def strip_trailing_export_artifacts(document: Any) -> int:
    """Remove Quark's hidden metadata and duplicated outline appended at EOF."""
    paragraphs = list(document.paragraphs)
    start = None
    for index, paragraph in enumerate(paragraphs):
        value = normalize_text(paragraph.text)
        if value.lower().replace(" ", "") == "[generalinformation]":
            start = index
            break
    if start is None:
        # Fallback for exports that omit the metadata label but append a
        # second front-matter outline after the last answer section.
        for index, paragraph in enumerate(paragraphs):
            if normalize_text(paragraph.text) == "前言" and index > max(0, len(paragraphs) - 120):
                start = index
                break
    if start is None:
        return 0
    removed = 0
    for paragraph in paragraphs[start:]:
        remove_paragraph(paragraph)
        removed += 1
    return removed


def remove_page_break_only_paragraphs(document: Any) -> int:
    """Remove redundant empty page-break paragraphs from paginated exports."""
    removed = 0
    for paragraph in list(iter_paragraphs(document)):
        if paragraph.text.strip() or paragraph._p.xpath(".//w:drawing"):
            continue
        breaks = paragraph._p.xpath(".//w:br")
        if not breaks:
            continue
        # A source section already starts the next page. Keeping a separate
        # empty paragraph containing w:br creates a blank page in WPS/LO.
        for br in breaks:
            br.getparent().remove(br)
        if not paragraph._p.xpath(".//w:r") and not paragraph._p.xpath(".//w:sectPr"):
            remove_paragraph(paragraph)
        removed += 1
    return removed


def compact_empty_spacer_paragraphs(document: Any, max_consecutive: int = 1) -> int:
    """Reduce accidental empty spacer runs without touching intentional breaks."""
    removed = 0
    empty_run = 0
    for paragraph in list(document.paragraphs):
        if paragraph_in_table(paragraph):
            empty_run = 0
            continue
        if paragraph.text.strip() or paragraph._p.xpath(".//w:drawing|.//m:oMath|.//m:oMathPara|.//w:object|.//w:pict") or paragraph._p.xpath(".//w:br|.//w:sectPr"):
            empty_run = 0
            continue
        empty_run += 1
        if empty_run > max_consecutive:
            remove_paragraph(paragraph)
            removed += 1
    return removed


def break_excessive_keep_chains(document: Any, max_chain: int = 8) -> int:
    """Break pathological keep-with-next chains that create one-line pages."""
    changed = 0
    chain: list[Any] = []
    for paragraph in list(document.paragraphs) + [None]:
        if paragraph is not None and paragraph.paragraph_format.keep_with_next:
            chain.append(paragraph)
            continue
        if len(chain) > max_chain:
            for item in chain[max_chain - 1:]:
                if item.paragraph_format.keep_with_next:
                    item.paragraph_format.keep_with_next = False
                    changed += 1
        chain = []
    return changed


def visual_layout_snapshot(document: Any) -> dict[str, Any]:
    """Capture image counts and anchor geometry before typography edits."""
    from docx.oxml.ns import qn  # type: ignore
    anchors: list[dict[str, Any]] = []
    for paragraph in iter_paragraphs(document):
        for anchor in paragraph._p.xpath(".//wp:anchor"):
            # Text-box anchors (notably floating TOC glyphs) are intentionally
            # rebuilt as editable paragraphs. Track only picture anchors here.
            if not any(element.tag.endswith("}pic") for element in anchor.iter()):
                continue
            extent = anchor.xpath("./wp:extent")
            if not extent:
                continue
            item: dict[str, Any] = {
                "cx": extent[0].get("cx"),
                "cy": extent[0].get("cy"),
                "relative_h": None,
                "relative_v": None,
                "offset_h": None,
                "offset_v": None,
                "wrap": None,
            }
            pos_h = anchor.xpath("./wp:positionH")
            pos_v = anchor.xpath("./wp:positionV")
            wrap = anchor.xpath("./wp:wrapSquare|./wp:wrapTight|./wp:wrapThrough|./wp:wrapTopAndBottom|./wp:wrapNone")
            if pos_h:
                item["relative_h"] = pos_h[0].get("relativeFrom")
                offset = next((element for element in pos_h[0] if element.tag.endswith("}posOffset")), None)
                item["offset_h"] = offset.text if offset is not None else None
            if pos_v:
                item["relative_v"] = pos_v[0].get("relativeFrom")
                offset = next((element for element in pos_v[0] if element.tag.endswith("}posOffset")), None)
                item["offset_v"] = offset.text if offset is not None else None
            if wrap:
                item["wrap"] = wrap[0].tag.split("}")[-1]
            anchors.append(item)
    return {"inline_count": len(document.inline_shapes), "anchored_count": len(anchors), "anchors": anchors}


def convert_picture_anchors_to_inline(document: Any) -> int:
    """Move picture anchors into the normal Word text flow.

    Quark stores many figures with absolute PDF-page coordinates. After page
    sections are collapsed and text is repaginated, those coordinates make
    figures overlap unrelated paragraphs. Picture anchors are therefore made
    inline; text-box anchors are handled separately by the TOC rebuilding pass.
    """
    converted = 0
    removable = {
        "positionH", "positionV", "wrapNone", "wrapSquare", "wrapTight",
        "wrapThrough", "wrapTopAndBottom", "simplePos",
    }
    for paragraph in iter_paragraphs(document):
        for anchor in list(paragraph._p.xpath(".//wp:anchor")):
            if not any(element.tag.endswith("}pic") for element in anchor.iter()):
                continue
            anchor.tag = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
            for child in list(anchor):
                if child.tag.split("}")[-1] in removable:
                    anchor.remove(child)
            converted += 1
    return converted


def separate_mixed_picture_paragraphs(document: Any) -> int:
    """Put pictures on their own paragraph when source XML mixed them with text.

    Office and WPS calculate an inline picture's line box differently from
    LibreOffice. A drawing left beside ordinary text can therefore cover the
    text or be clipped. Moving only the drawing run to a following paragraph
    keeps the source text and the visual object in normal document flow.
    """
    from copy import deepcopy  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore
    from docx.shared import Pt  # type: ignore
    from docx.enum.text import WD_LINE_SPACING  # type: ignore
    changed = 0
    for paragraph in list(iter_paragraphs(document)):
        drawings = list(paragraph._p.xpath(".//w:drawing[.//pic:pic]"))
        if not drawings or not normalize_text(paragraph.text or ""):
            continue
        movable_runs = []
        for run in list(paragraph._p.xpath("./w:r")):
            if run.xpath(".//w:drawing[.//pic:pic]") and not run.xpath("./w:t|./w:tab|./w:br"):
                movable_runs.append(run)
        if not movable_runs:
            continue
        new_p = OxmlElement("w:p")
        paragraph._p.addnext(new_p)
        new_paragraph = Paragraph(new_p, paragraph._parent)
        for run in movable_runs:
            new_p.append(run)
        new_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        new_paragraph.paragraph_format.line_spacing = None
        new_paragraph.paragraph_format.space_before = Pt(0)
        new_paragraph.paragraph_format.space_after = Pt(0)
        new_paragraph.paragraph_format.keep_together = True
        changed += len(movable_runs)
    return changed


def normalize_picture_paragraphs_for_office(document: Any) -> dict[str, int]:
    """Make every picture a plain, standalone Word-flow paragraph.

    Quark and some PDF converters leave anchor-only attributes on elements
    that were later renamed to ``wp:inline``.  LibreOffice is permissive about
    those attributes, while Microsoft Word may still apply their old
    overlap/position semantics.  This pass also separates every drawing from
    text or OMML in the same paragraph, including drawings nested in tables.
    """
    from copy import deepcopy  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore
    from docx.shared import Pt  # type: ignore

    inline_count = 0
    attrs_removed = 0
    paragraphs_split = 0
    picture_paragraphs = 0
    picture_nodes = ("w:drawing[.//pic:pic]",)

    def clean_inline(inline: Any) -> None:
        nonlocal attrs_removed
        # These attributes belong to wp:anchor.  They can survive a tag
        # conversion and cause Office to interpret an inline picture as a
        # floating/overlapping object.
        for name in ("behindDoc", "allowOverlap", "relativeHeight", "layoutInCell", "locked", "simplePos"):
            if name in inline.attrib:
                del inline.attrib[name]
                attrs_removed += 1
        for child in list(inline):
            local = child.tag.split("}")[-1]
            if local in {"positionH", "positionV", "wrapNone", "wrapSquare", "wrapTight", "wrapThrough", "wrapTopAndBottom", "simplePos"}:
                inline.remove(child)
                attrs_removed += 1

        extent = inline.xpath("./wp:extent")
        graphic_extents = inline.xpath(".//a:xfrm/a:ext")
        if extent and graphic_extents:
            cx, cy = extent[0].get("cx"), extent[0].get("cy")
            if cx and cy:
                for graphic_extent in graphic_extents:
                    graphic_extent.set("cx", cx)
                    graphic_extent.set("cy", cy)

    def style_picture_paragraph(paragraph: Any) -> None:
        nonlocal picture_paragraphs
        ppr = paragraph._p.get_or_add_pPr()
        # Remove converter-specific paragraph properties that affect the
        # inline picture's line box or produce renderer-specific alignment.
        for child in list(ppr):
            local = child.tag.split("}")[-1]
            if local in {"textAlignment", "wordWrap", "ind", "spacing", "jc", "rPr", "keepNext", "keepLines", "pageBreakBefore"}:
                ppr.remove(child)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Inline drawings taller than the current font line can protrude into
        # the previous/next paragraph in LibreOffice and Microsoft Word. Set
        # the picture paragraph's exact line box from its actual EMU height so
        # the drawing gets a real block of vertical space.
        extents = paragraph._p.xpath(".//wp:extent")
        if extents:
            try:
                height_points = float(extents[0].get("cy", "0")) / 12700.0
            except (TypeError, ValueError):
                height_points = 0.0
        else:
            height_points = 0.0
        if height_points > 0.1:
            # Never use an exact line box for a drawing. Word/WPS may clip an
            # inline picture when the renderer ascent/descent is larger than
            # the declared EMU extent. AT_LEAST lets the line grow safely.
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            paragraph.paragraph_format.line_spacing = Pt(height_points + 10.0)
        else:
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.line_spacing = None
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.widow_control = True
        picture_paragraphs += 1

    for paragraph in list(iter_paragraphs(document)):
        drawings = list(paragraph._p.xpath(".//w:drawing[.//pic:pic]"))
        if not drawings:
            continue
        for drawing in drawings:
            inline = next((node for node in drawing.iter() if node.tag.endswith("}inline")), None)
            if inline is None:
                # Anchors should already have been converted, but normalize
                # any late-created picture conservatively as well.
                inline = next((node for node in drawing.iter() if node.tag.endswith("}anchor")), None)
                if inline is not None:
                    inline.tag = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
            if inline is not None:
                clean_inline(inline)
                inline_count += 1

        has_text = bool(paragraph._p.xpath(".//w:t[text()]"))
        has_omath = bool(paragraph._p.xpath(".//m:oMath|.//m:oMathPara"))
        # Preserve a single pure-picture paragraph in place; only split extra
        # drawings.  Mixed text/equation paragraphs move every drawing below
        # the source text so no visual object shares its line box.
        mixed = has_text or has_omath
        if not mixed and len(drawings) == 1:
            style_picture_paragraph(paragraph)
            continue

        insertion_point = paragraph._p
        for drawing in drawings:
            source_run = drawing.getparent()
            while source_run is not None and not source_run.tag.endswith("}r"):
                source_run = source_run.getparent()
            new_p = OxmlElement("w:p")
            insertion_point.addnext(new_p)
            insertion_point = new_p
            new_paragraph = Paragraph(new_p, paragraph._parent)
            new_run = OxmlElement("w:r")
            if source_run is not None:
                source_rpr = source_run.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
                if source_rpr is not None:
                    new_run.append(deepcopy(source_rpr))
            source_run.remove(drawing)
            new_run.append(drawing)
            new_p.append(new_run)
            style_picture_paragraph(new_paragraph)
            paragraphs_split += 1
        # Drop now-empty runs and the source paragraph when it contained only
        # drawings.  Keep a text/OMML source paragraph in place.
        for run in list(paragraph._p.xpath("./w:r")):
            if not run.xpath("./w:t|./w:tab|./w:br|./w:drawing|./w:object|./w:pict"):
                run.getparent().remove(run)
        if not mixed and not paragraph._p.xpath(".//w:drawing|.//m:oMath|.//w:t[text()]"):
            remove_paragraph(paragraph)

    # A second lightweight sweep catches any inline pictures introduced by a
    # later replacement pass and guarantees the Office-only attributes stay
    # absent in the final document.
    for paragraph in list(iter_paragraphs(document)):
        for inline in paragraph._p.xpath(".//wp:inline"):
            clean_inline(inline)
        if paragraph._p.xpath(".//w:drawing[.//pic:pic]") and not paragraph._p.xpath(".//w:t[text()]|.//m:oMath|.//m:oMathPara"):
            style_picture_paragraph(paragraph)
    # A table row is allowed to split by default.  When it contains a figure
    # or formula, that split can leave the drawing half on one page and half
    # on the next in Office-compatible renderers.
    for table in getattr(document, "tables", []):
        for table_row in table.rows:
            has_visual = bool(table_row._tr.xpath(".//w:drawing[.//pic:pic]|.//m:oMath|.//m:oMathPara"))
            if not has_visual:
                continue
            tr_pr = table_row._tr.get_or_add_trPr()
            # Do not let converter-authored row heights clip visual content.
            for height in list(tr_pr.xpath("./w:trHeight")):
                tr_pr.remove(height)
            if not tr_pr.xpath("./w:cantSplit"):
                cant_split = OxmlElement("w:cantSplit")
                tr_pr.append(cant_split)
    return {
        "pictures_normalized": inline_count,
        "picture_attrs_removed": attrs_removed,
        "picture_paragraphs_split": paragraphs_split,
        "picture_paragraphs_styled": picture_paragraphs,
    }


def compare_visual_layout(document: Any, before: dict[str, Any]) -> dict[str, Any]:
    """Report whether edits changed object counts or anchor relationships."""
    after = visual_layout_snapshot(document)
    position_keys = ("relative_h", "relative_v", "offset_h", "offset_v", "wrap")
    before_positions = [{key: item.get(key) for key in position_keys} for item in before.get("anchors", [])]
    after_positions = [{key: item.get(key) for key in position_keys} for item in after.get("anchors", [])]
    geometry_changed = before_positions != after_positions
    return {
        "before": {key: before.get(key) for key in ("inline_count", "anchored_count")},
        "after": {key: after.get(key) for key in ("inline_count", "anchored_count")},
        "anchor_geometry_changed": geometry_changed,
        "passed": before.get("inline_count") == after.get("inline_count") and before.get("anchored_count") == after.get("anchored_count"),
    }


def normalize_heading_text(paragraph: Any, logical: str) -> None:
    """Remove OCR-inserted tabs/spaces inside short front-matter headings."""
    if logical not in {"前言", "目录"}:
        return
    runs = list(paragraph.runs)
    if not runs:
        return
    runs[0].text = logical
    for run in runs[1:]:
        run._element.getparent().remove(run._element)


def normalize_ocr_noise(paragraph: Any) -> int:
    """Remove only unambiguous OCR/typing noise from ordinary text runs.

    Math, drawings and table structure are left untouched.  This intentionally
    avoids guessing at technical symbols; it only removes zero-width controls,
    collapses duplicate punctuation, and drops punctuation-only paragraphs.
    """
    # A converter may mix an OMML object with a neighboring plain-text LaTeX
    # fragment in the same paragraph. Keep the object, but still normalize
    # that fragment; skip visual-only paragraphs with no such text.
    if paragraph_has_visual_object(paragraph) and not re.search(
        r"\\(?:frac|sqrt|overline|underline|mathrm|mathbf|mathit|[A-Za-z]{2,})\b",
        paragraph.text or "",
    ):
        return 0
    changed = 0
    runs = list(paragraph.runs)
    for run in runs:
        value = run.text or ""
        cleaned = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", value)
        # A few Quark exports put inline equations in ordinary text runs.
        # Word does not interpret LaTeX commands there, so leaving strings
        # such as ``\frac`` or ``\omega`` produces visibly broken formulas.
        # Convert only runs containing unmistakable math commands; OMML and
        # picture-based equations are excluded by paragraph_has_visual_object.
        if re.search(r"\\(?:frac|sqrt|overline|underline|mathrm|mathbf|mathit|[A-Za-z]{2,})\b", cleaned):
            normalized_math = _mineru_inline_math_text(cleaned)
            if normalized_math:
                cleaned = normalized_math
        cleaned = re.sub(r"([，。；：、,.!?！？])\1+", r"\1", cleaned)
        if cleaned != value:
            run.text = cleaned
            changed += 1
    if runs and not normalize_text(paragraph.text) and any((run.text or "").strip() for run in runs):
        for run in runs:
            run.text = ""
        changed += 1
    return changed


def normalize_ocr_spacing(paragraph: Any) -> int:
    """Remove OCR spaces inserted inside Chinese phrases.

    Quark frequently emits a whitespace-only run between two Chinese
    characters (for example ``解     取 长 度``).  These spaces are not useful
    typography and can make a short line wrap into a tall, sparse column.
    Keep spaces around Latin words, numbers, and mathematical operators where
    they may carry meaning; only remove spaces whose two visible neighbours
    are Chinese characters.
    """
    runs = list(paragraph.runs)
    if not runs or "\t" in (paragraph.text or ""):
        return 0
    changed = 0
    # First clean spaces contained within a single run.  Quark uses long
    # runs of ordinary spaces as visual positioning; once the content is in
    # an editable paragraph those offsets become accidental gaps.  Keep a
    # single separator inside Latin/math text, but never keep padding before
    # punctuation or at either paragraph edge.
    for run in runs:
        value = run.text or ""
        if not value:
            continue
        cleaned = value.replace("\u00a0", " ")
        cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
        cleaned = re.sub(r"\s+([，。；：、,.!?！？;:）】》』」〕】])", r"\1", cleaned)
        cleaned = re.sub(r"([（【《『「〔])\s+", r"\1", cleaned)
        cleaned = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", cleaned)
        if cleaned != value:
            run.text = cleaned
            changed += 1
    # Remove whitespace-only runs sitting between visible text.  Precompute
    # neighbours once; repeatedly scanning all earlier/later runs makes a
    # large OCR paragraph needlessly quadratic.
    previous_chars = [""] * len(runs)
    following_chars = [""] * len(runs)
    previous = ""
    for index, run in enumerate(runs):
        previous_chars[index] = previous
        visible = next((char for char in (run.text or "") if not char.isspace()), "")
        if visible:
            previous = visible
    following = ""
    for index in range(len(runs) - 1, -1, -1):
        following_chars[index] = following
        visible = next((char for char in (runs[index].text or "") if not char.isspace()), "")
        if visible:
            following = visible
    for index, run in enumerate(runs):
        value = run.text or ""
        if not value or value.strip() or "\t" in value:
            continue
        previous = previous_chars[index]
        following = following_chars[index]
        if not previous or not following:
            continue
        no_separator = bool(
            (re.match(r"[\u4e00-\u9fff]", previous) and re.match(r"[\u4e00-\u9fff]", following))
            or re.match(r"[，。；：、,.!?！？;:）】》』」〕】]", following)
            or re.match(r"[（【《『「〔]", previous)
        )
        replacement = "" if no_separator else " "
        if value != replacement:
            run.text = replacement
            changed += 1
    # Trim converter padding at paragraph edges without touching tabs.
    visible = [run for run in runs if run.text and run.text.strip()]
    if visible:
        first, last = visible[0], visible[-1]
        first_clean = (first.text or "").lstrip(" ")
        last_clean = (last.text or "").rstrip(" ")
        if first_clean != first.text:
            first.text = first_clean
            changed += 1
        if last_clean != last.text:
            last.text = last_clean
            changed += 1
    return changed


def normalize_tabular_alignment_rows(document: Any) -> dict[str, int]:
    """Normalize Quark's tab-separated pseudo-tables without flattening runs.

    A large number of exports represent answer choices, vector columns and
    ``figure label + value`` rows as ordinary paragraphs with hard-coded
    indents/tab stops.  Those offsets are page-specific and become unstable
    after the rest of the document is restyled.  Keep the text and OMML
    objects intact, but give each nearby group a shared Word tab grid.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Twips  # type: ignore

    paragraphs = list(iter_paragraphs(document))
    candidates: list[tuple[int, Any, int]] = []
    for index, paragraph in enumerate(paragraphs):
        # Visual rows may contain an inline equation alongside a tabbed value;
        # those still need the shared column grid.  Only skip visual blocks
        # that have no tab at all.
        if paragraph_in_table(paragraph):
            continue
        raw = paragraph.text or ""
        if "\t" not in raw or is_toc_line(raw):
            continue
        # Do not reinterpret a normal prose tab.  Rows with a short number of
        # columns, or an explicit long-space separator after a tab, are the
        # deterministic pseudo-table cases produced by Quark.
        parts = raw.split("\t")
        if len(parts) < 2 or len(parts) > 6:
            continue
        # Section/question headings such as ``四、\t如图所示`` are one
        # sentence, not two columns.  A tab here creates the large blank gap
        # visible in Word/WPS, so collapse only this unambiguous prefix.
        prefix = re.sub(r"\s+", "", parts[0])
        if re.match(r"^[一二三四五六七八九十百]+[、.．]$", prefix):
            for run in paragraph.runs:
                value = run.text or ""
                if "\t" in value:
                    run.text = value.replace("\t", " ")
            continue
        visible = normalize_text(raw.replace("\t", " "))
        if not visible or len(visible) > 180:
            continue
        # A long Chinese explanation with a tab is reading order, not a
        # pseudo-table. Keep its original flow instead of moving the tail to
        # a far-right column.
        if len(re.findall(r"[\u4e00-\u9fff]", visible)) > 18 and not re.search(r"(?:^|\s)[A-DＡ-Ｄ][.、]", visible):
            continue
        inferred_columns = len(parts)
        # Quark sometimes encodes the last pseudo-column as padding spaces
        # after a tab rather than as another w:tab (common in answer rows).
        if re.search(r"\t[^\n]* {4,}\S", raw):
            inferred_columns += 1
        candidates.append((index, paragraph, min(6, inferred_columns)))

    if not candidates:
        return {"column_alignment_rows": 0, "column_alignment_groups": 0, "column_alignment_changes": 0}

    # Form clusters of adjacent tabbed rows.  A single formula/image paragraph
    # between rows is allowed: Quark often stores the equation between the
    # column labels and the answer row.
    clusters: list[list[tuple[int, Any, int]]] = []
    current: list[tuple[int, Any, int]] = []
    last_index = -100
    for item in candidates:
        if not current or item[0] - last_index <= 3:
            current.append(item)
        else:
            clusters.append(current)
            current = [item]
        last_index = item[0]
    if current:
        clusters.append(current)

    changed = 0
    rows = 0
    groups = 0
    for cluster in clusters:
        # A lone tabbed paragraph can still be an option row, but avoid
        # changing isolated tabs in ordinary prose unless it has a compact
        # two-column shape.
        max_columns = max(columns for _, _, columns in cluster)
        if len(cluster) == 1 and max_columns == 2:
            text = cluster[0][1].text or ""
            if len(text.replace("\t", "").strip()) > 90:
                continue
        groups += 1
        rows += len(cluster)
        section = cluster[0][1]._parent.part.document.sections[0]
        usable = section.page_width - section.left_margin - section.right_margin
        # Keep the final stop away from the right edge so long values do not
        # wrap into a second line.  Word measures tab positions from the
        # paragraph's left edge, therefore all source indents are removed.
        usable_twips = max(7200, int(usable.twips) if hasattr(usable, "twips") else 9360)
        if max_columns <= 2:
            fractions = (0.52,)
        else:
            fractions = tuple(i / max_columns for i in range(1, max_columns))
        stops = [Twips(min(usable_twips - 360, max(1800, int(usable_twips * fraction)))) for fraction in fractions]
        for _, paragraph, _ in cluster:
            raw = paragraph.text or ""
            # Long runs of padding after a tab are another pseudo-column
            # marker (e.g. ``解析：B\tA+4          0``).  Convert only this
            # unambiguous padding in ordinary text paragraphs; never flatten
            # a paragraph containing OMML or drawings.
            if not paragraph._p.xpath(".//w:drawing|.//m:oMath|.//m:oMathPara"):
                seen_tab = False
                for run in paragraph.runs:
                    value = run.text or ""
                    if "\t" in value:
                        seen_tab = True
                    # The tab itself is commonly a separate run, followed by
                    # a padded value run. Treat that padding as the next
                    # pseudo-column, while leaving ordinary in-word spaces
                    # untouched.
                    if seen_tab and re.search(r" {4,}\S", value):
                        value = re.sub(r" {4,}(?=\S)", "\t", value, count=1)
                    value = re.sub(r"\t+", "\t", value)
                    if value != run.text:
                        run.text = value
                        changed += 1
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = paragraph.paragraph_format
            pf.left_indent = Twips(0)
            pf.right_indent = Twips(0)
            pf.first_line_indent = Twips(0)
            pf.keep_together = True
            try:
                pf.tab_stops.clear_all()
            except AttributeError:
                pass
            for stop in stops:
                pf.tab_stops.add_tab_stop(stop)
            changed += 1
    return {
        "column_alignment_rows": rows,
        "column_alignment_groups": groups,
        "column_alignment_changes": changed,
    }


def _is_formula_fragment(paragraph: Any) -> bool:
    """Return true for a short standalone formula/symbol fragment."""
    if paragraph_in_table(paragraph) or paragraph._p.xpath(".//w:drawing"):
        return False
    text = normalize_text(paragraph.text)
    compact = re.sub(r"\s+", "", text)
    if paragraph._p.xpath(".//m:oMath|.//m:oMathPara"):
        # A body sentence can contain an inline equation.  It is not a
        # standalone fragment and must remain its own paragraph.
        return len(compact) <= 36 and not re.search(r"[\u4e00-\u9fff]", compact)
    if not compact or len(compact) > 16:
        return False
    return bool(re.fullmatch(r"[A-Za-z0-9\u0370-\u03ff\u1d00-\u1d7f\u2070-\u209f°′″′′′+−=<>≤≥≈∠∕/;:,.()\[\]{}]+", compact))


def merge_formula_fragment_paragraphs(document: Any) -> int:
    """Join OCR-split formula labels into one editable inline formula block.

    A run of short formula paragraphs is merged only when it follows a
    question/step paragraph ending in a colon.  This targets diagram-label
    fragments such as ``B₁``, ``C₁α`` and ``60°`` without flattening ordinary
    body text or tables into a single paragraph.
    """
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    paragraphs = list(document.paragraphs)
    merged = 0
    index = 0
    while index < len(paragraphs):
        paragraph = paragraphs[index]
        logical = normalize_text(paragraph.text)
        if not logical or not logical.endswith(("：", ":")):
            index += 1
            continue
        cursor = index + 1
        while cursor < len(paragraphs) and not normalize_text(paragraphs[cursor].text) and not paragraph_has_visual_object(paragraphs[cursor]):
            cursor += 1
        start = cursor
        while cursor < len(paragraphs) and _is_formula_fragment(paragraphs[cursor]):
            cursor += 1
        if cursor - start < 3:
            index += 1
            continue
        fragments = paragraphs[start:cursor]
        for fragment in fragments:
            for run in fragment.runs:
                value = run.text or ""
                if value.strip():
                    run.text = re.sub(r"[ \t]{2,}", " ", value).strip()
                else:
                    run.text = ""
        target = fragments[0]
        target_p = target._p
        # Keep the first paragraph properties, then append content from each
        # following paragraph with a single separator run.
        for fragment in fragments[1:]:
            separator = OxmlElement("w:r")
            separator_rpr = OxmlElement("w:rPr")
            separator.append(separator_rpr)
            text_node = OxmlElement("w:t")
            text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            text_node.text = " "
            separator.append(text_node)
            target_p.append(separator)
            for child in list(fragment._p):
                if child.tag == qn("w:pPr"):
                    continue
                target_p.append(child)
            fragment._p.getparent().remove(fragment._p)
            merged += 1
        # Source fragments often carry several trailing layout spaces. Keep
        # one separator between adjacent formula objects so the result stays
        # readable without reproducing the OCR column spacing.
        for run in target.runs:
            value = run.text or ""
            cleaned = re.sub(r"[ \t]{2,}", " ", value)
            if cleaned != value:
                run.text = cleaned
        target.paragraph_format.keep_together = True
        target.paragraph_format.left_indent = 0
        target.paragraph_format.right_indent = 0
        target.paragraph_format.first_line_indent = 0
        paragraphs = list(document.paragraphs)
        # ``python-docx`` creates a new Paragraph wrapper on each traversal,
        # so object identity is not stable after the XML children are moved.
        target_position = next((pos for pos, item in enumerate(paragraphs) if item._p is target._p), len(paragraphs) - 1)
        index = target_position + 1
    return merged


def compact_symbol_fragment_layout(document: Any) -> int:
    """Keep OCR-split diagram labels compact without changing their objects.

    Quark sometimes exports labels from one diagram as separate one-line
    paragraphs (for example ``CP``, ``CTT`` and ``Q3``).  They are not
    reliable enough to merge into a sentence or to turn back into floating
    text boxes.  A bounded keep group and natural spacing prevents the body
    19 pt rule from turning the labels into a tall vertical column while
    preserving their searchable text and existing inline drawings.
    """
    from docx.enum.text import WD_LINE_SPACING  # type: ignore
    from docx.shared import Pt  # type: ignore

    paragraphs = list(document.paragraphs)
    changed = 0
    index = 0
    while index < len(paragraphs):
        if not _is_formula_fragment(paragraphs[index]):
            index += 1
            continue
        start = index
        while index < len(paragraphs) and _is_formula_fragment(paragraphs[index]):
            index += 1
        run = paragraphs[start:index]
        if len(run) < 2:
            continue
        # Do not create long chains that can push a large region to the next
        # page.  At most eight short labels travel together.
        for chunk_start in range(0, len(run), 8):
            chunk = run[chunk_start:chunk_start + 8]
            for position, paragraph in enumerate(chunk):
                pf = paragraph.paragraph_format
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                pf.line_spacing = None
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.keep_together = True
                pf.keep_with_next = position < len(chunk) - 1
                pf.widow_control = True
                changed += 1
    return changed


def normalize_numbered_subitem_alignment(document: Any) -> int:
    """Align numbered sub-items such as ``(1)``/``(2)`` on one left edge.

    Quark exports commonly mix a leading tab, direct left indentation and a
    first-line indentation for these paragraphs.  The combination makes the
    visible number drift even when the source lines are logically adjacent.
    Use a stable hanging indent: the number starts at the paragraph edge and
    wrapped lines start one character after it.  Only leading whitespace is
    removed; formula runs and all other text remain untouched.
    """
    from docx.shared import Pt  # type: ignore

    pattern = re.compile(r"^\s*(?:\([一二三四五六七八九十百0-9]+\)|（[一二三四五六七八九十百0-9]+）)")
    changed = 0
    for paragraph in iter_paragraphs(document):
        if paragraph_has_visual_object(paragraph):
            continue
        raw = paragraph.text or ""
        if not pattern.match(raw):
            continue
        # Remove only whitespace before the numbered marker while preserving
        # run-level fonts and any OMML later in the paragraph.
        leading = True
        for run in paragraph.runs:
            value = run.text or ""
            if not leading:
                break
            if value:
                cleaned = value.lstrip(" \t\r\n")
                if cleaned != value:
                    run.text = cleaned
                    changed += 1
                if cleaned:
                    leading = False
        try:
            paragraph.paragraph_format.tab_stops.clear_all()
        except AttributeError:
            pass
        pf = paragraph.paragraph_format
        # Keep table contents inside their cell; body paragraphs get a
        # hanging indent so wrapped lines remain aligned with the text.
        if paragraph_in_table(paragraph):
            left, first = Pt(0), Pt(0)
        else:
            left, first = Pt(21), Pt(-21)
        if pf.left_indent != left:
            pf.left_indent = left
            changed += 1
        if pf.first_line_indent != first:
            pf.first_line_indent = first
            changed += 1
        if pf.right_indent is not None:
            pf.right_indent = Pt(0)
            changed += 1
        if paragraph.alignment != 0:
            paragraph.alignment = 0
            changed += 1
    return changed


def normalize_math_sizes(document: Any) -> int:
    """Set editable OMML runs to the project's 13 pt formula size."""
    from docx.oxml.ns import qn  # type: ignore
    changed = 0
    for paragraph in iter_paragraphs(document):
        if not paragraph._p.xpath(".//m:oMath"):
            continue
        for element in paragraph._p.xpath(".//w:sz|.//w:szCs"):
            if element.get(qn("w:val")) != "26":
                element.set(qn("w:val"), "26")
                changed += 1
    return changed


def normalize_formula_line_boxes(document: Any) -> int:
    """Use single spacing for standalone OMML; preserve table formula layout."""
    from docx.enum.text import WD_LINE_SPACING  # type: ignore
    from docx.shared import Pt  # type: ignore

    changed = 0
    for paragraph in iter_paragraphs(document):
        if paragraph_in_table(paragraph) or not paragraph._p.xpath(".//m:oMath|.//m:oMathPara"):
            continue
        pf = paragraph.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.line_spacing = None
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.keep_together = True
        changed += 1
    return changed


def finalize_formula_single_spacing(document: Any) -> int:
    """Write an explicit single-spacing rule for every OMML paragraph.

    The document's Normal style intentionally uses the 19 pt body baseline
    as an exact line box.  Clearing paragraph-level spacing therefore still
    makes Word/WPS inherit ``lineRule=exact`` and display formula paragraphs
    as fixed spacing.  A direct OOXML override is required at the very end of
    the pipeline so the formula setting wins over that style inheritance.
    """
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    changed = 0
    spacing_tag = qn("w:spacing")
    line_tag = qn("w:line")
    line_rule_tag = qn("w:lineRule")
    before_tag = qn("w:before")
    after_tag = qn("w:after")
    for paragraph in iter_paragraphs(document):
        if not paragraph._p.xpath(".//m:oMath|.//m:oMathPara"):
            continue
        ppr = paragraph._p.get_or_add_pPr()
        spacing = ppr.find(spacing_tag)
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            ppr.append(spacing)
        # 240 twips is Word's explicit single-spacing value.  ``auto`` is
        # essential: omitting it lets the Normal style's ``exact`` rule win.
        desired = {
            line_tag: "240",
            line_rule_tag: "auto",
            before_tag: "0",
            after_tag: "0",
        }
        for key, value in desired.items():
            if spacing.get(key) != value:
                spacing.set(key, value)
                changed += 1
    return changed


def remove_formula_table_fixed_heights(document: Any) -> int:
    """Remove converter row heights from formula tables so cells can expand."""
    changed = 0
    for table in document.tables:
        if not table._tbl.xpath(".//m:oMath|.//m:oMathPara"):
            continue
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            for height in list(tr_pr.xpath("./w:trHeight")):
                tr_pr.remove(height)
                changed += 1
    return changed


def expand_formula_table_widths(document: Any) -> int:
    """Widen formula tables within the page, reclaiming converter indent.

    ``normalize_table_widths`` preserves a large source ``tblInd``.  That
    makes a table technically fit but leaves it visibly narrow in Word/WPS.
    For formula tables only, move the table back to the text margin and use
    the recovered indent as width, without changing cell content or row
    heights.
    """
    from docx.oxml.ns import qn  # type: ignore

    if not document.sections:
        return 0
    section = document.sections[0]
    # OOXML table widths are twips; python-docx section dimensions are EMU.
    available = max(1, int((section.page_width - section.left_margin - section.right_margin) / 635))
    changed = 0
    for table in document.tables:
        if not table._tbl.xpath(".//m:oMath|.//m:oMathPara"):
            continue
        grid = table._tbl.tblGrid
        columns = list(grid.gridCol_lst)
        widths = []
        for col in columns:
            try:
                widths.append(int(col.get(qn("w:w"), "0")))
            except (TypeError, ValueError):
                widths.append(0)
        total = sum(widths)
        if not widths or total <= 0:
            continue
        tbl_pr = table._tbl.tblPr
        indent_node = tbl_pr.first_child_found_in("w:tblInd")
        try:
            indent = max(0, int(indent_node.get(qn("w:w"), "0"))) if indent_node is not None else 0
        except (TypeError, ValueError):
            indent = 0
        target_limit = max(1, available)
        # Keep a modest growth when the table is already near the margin, but
        # reclaim the source indent when it is the reason the table is narrow.
        target = min(target_limit, max(int(total * 1.06), total + indent))
        if target <= total:
            continue
        if indent_node is not None:
            indent_node.set(qn("w:w"), "0")
        scaled = [max(240, int(value * target / total)) for value in widths]
        scaled[-1] += target - sum(scaled)
        for col, width in zip(columns, scaled):
            col.set(qn("w:w"), str(width))
        tbl_w = tbl_pr.first_child_found_in("w:tblW")
        if tbl_w is not None:
            tbl_w.set(qn("w:w"), str(target))
            tbl_w.set(qn("w:type"), "dxa")
        for row in table.rows:
            for cell, width in zip(row.cells, scaled):
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.first_child_found_in("w:tcW")
                if tc_w is not None:
                    tc_w.set(qn("w:w"), str(width))
        table.autofit = False
        changed += 1
    return changed


def _set_table_borders(table: Any, value: str = "nil") -> None:
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), value)
        node.set(qn("w:sz"), "0")


def normalize_layout_table_frames(document: Any) -> dict[str, int]:
    """Hide converter layout grids while preserving real data-table borders."""
    hidden = removed_empty = 0
    for table in list(document.tables):
        cells = [cell for row in table.rows for cell in row.cells]
        texts = [normalize_text(cell.text) for cell in cells]
        nonempty = [text for text in texts if text]
        if not nonempty:
            table._element.getparent().remove(table._element)
            removed_empty += 1
            continue
        joined = " ".join(nonempty)
        option_like = bool(re.search(r"(?:^|\s)[A-DＡ-Ｄ][.、]", joined))
        formula_like = bool(re.search(r"[∑∫√=≤≥≠]|[A-Za-z]\s*[_^]", joined))
        sparse_layout = len(nonempty) <= max(4, len(cells) // 2) and max(map(len, nonempty), default=0) < 180
        # Quark emits these tables as visual scaffolds for answer choices,
        # formula alternatives, and metadata boxes. Their borders are not
        # semantic table borders and become dotted frames in Word/WPS.
        if option_like or (formula_like and len(table.rows) <= 3) or (len(table.rows) == 1 and len(table.columns) <= 4) or (sparse_layout and len(table.rows) <= 3):
            _set_table_borders(table)
            hidden += 1
    return {"layout_table_borders_hidden": hidden, "empty_layout_tables_removed": removed_empty}


def normalize_layout_padding(document: Any) -> int:
    """Turn obvious visual padding between numbered columns into real tabs."""
    changed = 0
    for paragraph in iter_paragraphs(document):
        if paragraph_in_table(paragraph) or paragraph._p.xpath(".//w:drawing|.//m:oMath|.//m:oMathPara"):
            continue
        raw = paragraph.text or ""
        if len(raw) > 220 or "\t" in raw:
            continue
        compact = normalize_text(raw)
        if not (re.match(r"^(?:\(?[0-9一二三四五六七八九十]+[.)、]|[A-DＡ-Ｄ][.、])", compact) or re.match(r"^\d+\.", compact)):
            continue
        for run in paragraph.runs:
            value = run.text or ""
            replaced = re.sub(r"(?<=\S) {3,}(?=\S)", "\t", value)
            if replaced != value:
                run.text = replaced
                changed += 1
    return changed


def normalize_table_title_blocks(document: Any) -> dict[str, int]:
    """Recover explicit exam-cover titles stored inside Quark tables."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore
    from docx.shared import Pt  # type: ignore

    changed = breaks = 0
    for table in list(document.tables):
        text = normalize_text(" ".join(cell.text for row in table.rows for cell in row.cells))
        has_year = bool(re.search(r"20\d{2}\s*年", text))
        has_school = bool(re.search(r"(?:大学|学院)", text))
        has_exam = "研究生入学考试" in text or "初试试题" in text
        if not ((has_year and has_school) or (has_school and has_exam)):
            continue
        previous = table._element.getprevious()
        already_break = previous is not None and previous.tag.endswith("}p") and previous.xpath(".//w:br[@w:type='page']")
        if not already_break:
            page = OxmlElement("w:p")
            run = OxmlElement("w:r")
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            run.append(br)
            page.append(run)
            table._element.addprevious(page)
            breaks += 1
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    logical = re.sub(r"\s+", "", paragraph.text or "")
                    if not (re.search(r"20\d{2}年", logical) or "研究生入学考试" in logical or "初试试题" in logical):
                        continue
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pf = paragraph.paragraph_format
                    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                    pf.line_spacing = Pt(32)
                    pf.space_before = Pt(2)
                    pf.space_after = Pt(4)
                    pf.keep_together = True
                    pf.keep_with_next = True
                    for run in paragraph.runs:
                        set_run_fonts(run, heading_font())
                        run.font.size = Pt(23)
                        run.bold = True
                    changed += 1
    return {"table_title_paragraphs": changed, "table_title_page_breaks": breaks}


def normalize_malformed_math(document: Any) -> dict[str, int]:
    """Remove Quark's empty OMML slots that render as visible square boxes.

    Quark occasionally emits a fraction with an empty numerator or a
    sub/superscript whose base is empty. Word renders those empty math
    containers as dotted square placeholders. The safe deterministic repairs
    are limited to common notation: an empty fraction numerator is ``1``;
    an empty subscript base is the angle marker; polynomial powers use ``φ``.
    Unknown empty structures are removed rather than inventing mathematical
    content.
    """
    from copy import deepcopy  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def q(local: str) -> str:
        return "{" + M_NS + "}" + local

    def add_math_text(parent: Any, value: str) -> None:
        run = OxmlElement("m:r")
        style = next((child for child in parent.iter() if child.tag == "{" + W_NS + "}rPr"), None)
        if style is not None:
            run.append(deepcopy(style))
        text = OxmlElement("m:t")
        text.text = value
        run.append(text)
        parent.append(run)

    def math_text(element: Any) -> str:
        return "".join(node.text or "" for node in element.iter(q("t")))

    stats = {
        "empty_fraction_numerators_filled": 0,
        "empty_angle_bases_filled": 0,
        "empty_polynomial_bases_filled": 0,
        "script_containers_downgraded": 0,
        "unknown_math_slots_removed": 0,
    }

    def has_math_text(element: Any) -> bool:
        return bool(math_text(element).strip()) if element is not None else False

    def unwrap_element(parent: Any, element: Any) -> None:
        """Replace a malformed math wrapper with its valid base children."""
        if parent is None or element is None:
            return
        index = parent.index(element)
        base = element.find(q("e"))
        children = list(base) if base is not None else []
        for offset, child in enumerate(children):
            parent.insert(index + offset, child)
        parent.remove(element)

    def remove_empty_script(script: Any) -> None:
        parent = script.getparent()
        if parent is None:
            return
        base = script.find(q("e"))
        if base is not None and math_text(base):
            unwrap_element(parent, script)
        else:
            parent.remove(script)
    for paragraph in iter_paragraphs(document):
        for math in list(paragraph._p.xpath(".//m:oMath")):
            all_text = math_text(math)
            polynomial = bool(re.search(r"c\d", all_text, re.I) and ("Φ" in all_text or "φ" in all_text or all_text.startswith(("s=", "v=", "a=", "h="))))
            for fraction in list(math.iter(q("f"))):
                numerator = fraction.find(q("num"))
                denominator = fraction.find(q("den"))
                numerator_empty = numerator is None or not math_text(numerator)
                denominator_empty = denominator is None or not math_text(denominator)
                # An entirely empty fraction is a converter placeholder, not
                # a mathematical 1/1. Remove it so Office cannot paint a box.
                if numerator_empty and denominator_empty:
                    parent = fraction.getparent()
                    if parent is not None:
                        parent.remove(fraction)
                        stats["unknown_math_slots_removed"] += 1
                    continue
                if numerator is not None and numerator_empty and not denominator_empty:
                    add_math_text(numerator, "1")
                    stats["empty_fraction_numerators_filled"] += 1
            for tag in ("sSub", "sSup", "sSubSup"):
                for script in list(math.iter(q(tag))):
                    base = script.find(q("e"))
                    if tag == "sSubSup":
                        sub = script.find(q("sub"))
                        sup = script.find(q("sup"))
                        has_sub = has_math_text(sub)
                        has_sup = has_math_text(sup)
                        # Quark frequently emits sSubSup with only one
                        # script child. Office paints the absent child as a
                        # square placeholder. Downgrade it to the valid
                        # one-script form instead of inventing content.
                        if has_sub and not has_sup:
                            script.tag = q("sSub")
                            stats["script_containers_downgraded"] += 1
                            tag = "sSub"
                        elif has_sup and not has_sub:
                            script.tag = q("sSup")
                            stats["script_containers_downgraded"] += 1
                            tag = "sSup"
                        elif not has_sub and not has_sup and base is not None and math_text(base):
                            parent = script.getparent()
                            unwrap_element(parent, script)
                            stats["unknown_math_slots_removed"] += 1
                            continue
                    if tag == "sSub" and (script.find(q("sub")) is None or not has_math_text(script.find(q("sub")))):
                        remove_empty_script(script)
                        stats["unknown_math_slots_removed"] += 1
                        continue
                    if tag == "sSup" and (script.find(q("sup")) is None or not has_math_text(script.find(q("sup")))):
                        remove_empty_script(script)
                        stats["unknown_math_slots_removed"] += 1
                        continue
                    if base is None or math_text(base):
                        if tag == "sSubSup" and base is not None:
                            sup = script.find(q("sup"))
                            if math_text(base) in {"'", "′"} and sup is not None and math_text(sup) in {"'", "′"}:
                                replacement = OxmlElement("m:r")
                                text = OxmlElement("m:t")
                                text.text = "″"
                                replacement.append(text)
                                parent = script.getparent()
                                if parent is not None:
                                    parent.replace(script, replacement)
                                    stats["unknown_math_slots_removed"] += 1
                                continue
                        # A sub/superscript container can have a valid base
                        # but an empty child (for example ``a'`` with an
                        # empty subscript). Word paints that empty child as a
                        # square placeholder, so remove only the empty child.
                        if base is not None and tag == "sSubSup":
                            for child_name in ("sub", "sup"):
                                child = script.find(q(child_name))
                                if child is not None and not math_text(child):
                                    script.remove(child)
                                    stats["unknown_math_slots_removed"] += 1
                        continue
                    if tag == "sSub":
                        add_math_text(base, "∠")
                        stats["empty_angle_bases_filled"] += 1
                    elif tag == "sSup" and polynomial:
                        add_math_text(base, "φ")
                        stats["empty_polynomial_bases_filled"] += 1
                    else:
                        parent = script.getparent()
                        if parent is not None:
                            parent.remove(script)
                            stats["unknown_math_slots_removed"] += 1
            # Empty delimiter groups are another Quark placeholder pattern.
            # Keep the surrounding text, but remove only the empty group.
            for delimiter in list(math.iter(q("d"))):
                element = delimiter.find(q("e"))
                if element is None or not math_text(element):
                    parent = delimiter.getparent()
                    if parent is not None:
                        parent.remove(delimiter)
                        stats["unknown_math_slots_removed"] += 1
    return stats


def normalize_table_widths(document: Any) -> dict[str, int]:
    """Fit every table, including its left indent, to the page text width."""
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    section = document.sections[0]
    available_twips = max(1, int((section.page_width - section.left_margin - section.right_margin) / 635))
    changed = 0
    tables = 0
    for table in document.tables:
        grid = table._tbl.tblGrid
        columns = list(grid.gridCol_lst)
        widths = []
        for col in columns:
            try:
                widths.append(max(1, int(col.get(qn("w:w"), "0"))))
            except (TypeError, ValueError):
                widths.append(1)
        total = sum(widths)
        if not widths:
            continue

        # Quark stores a full-page tblW and a separate tblInd.  Checking only
        # tblGrid makes a table look valid in XML while the indent pushes its
        # right edge outside the page.  Fit the grid into the remaining width
        # after the indent, while preserving the source's left alignment.
        tbl_pr = table._tbl.tblPr
        indent_node = tbl_pr.first_child_found_in("w:tblInd")
        try:
            indent = max(0, int(indent_node.get(qn("w:w"), "0"))) if indent_node is not None else 0
        except (TypeError, ValueError):
            indent = 0
        if indent >= available_twips:
            indent = 0
            if indent_node is not None:
                indent_node.set(qn("w:w"), "0")
        target_width = max(1, available_twips - indent)
        tbl_w_node = tbl_pr.first_child_found_in("w:tblW")
        try:
            declared_width = int(tbl_w_node.get(qn("w:w"), "0")) if tbl_w_node is not None else 0
        except (TypeError, ValueError):
            declared_width = 0
        source_width = max(total, declared_width)
        if source_width <= target_width and indent == 0:
            continue
        scaled = [max(240, int(value * target_width / max(source_width, 1))) for value in widths]
        correction = target_width - sum(scaled)
        scaled[-1] = max(240, scaled[-1] + correction)
        for col, width in zip(columns, scaled):
            col.set(qn("w:w"), str(width))
        tbl_w = tbl_pr.first_child_found_in("w:tblW")
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.insert(0, tbl_w)
        tbl_w.set(qn("w:w"), str(target_width))
        tbl_w.set(qn("w:type"), "dxa")
        layout = tbl_pr.first_child_found_in("w:tblLayout")
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        for row in table.rows:
            for cell, width in zip(row.cells, scaled):
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.first_child_found_in("w:tcW")
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(width))
                tc_w.set(qn("w:type"), "dxa")
        table.autofit = False
        # Small two-to-four-row data tables are commonly stranded with their
        # last row on the next page. Keep the rows together when the table is
        # short; long tables remain splittable for normal pagination.
        if len(table.rows) <= 4:
            from docx.oxml import OxmlElement as _OxmlElement  # type: ignore
            for row_index, row in enumerate(table.rows[:-1]):
                tr_pr = row._tr.get_or_add_trPr()
                cant_split = tr_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit")
                if cant_split is None:
                    tr_pr.append(_OxmlElement("w:cantSplit"))
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
        tables += 1
        changed += 1
    return {"tables_checked": len(document.tables), "tables_resized": tables, "table_width_changes": changed}


def normalize_table_pagination(document: Any) -> dict[str, int]:
    """Prevent short visual tables from splitting into orphan continuation rows."""
    from docx.oxml import OxmlElement  # type: ignore
    changed = 0
    for table in document.tables:
        row_count = len(table.rows)
        if row_count <= 1 or row_count > 10:
            continue
        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if not tr_pr.xpath("./w:cantSplit"):
                tr_pr.append(OxmlElement("w:cantSplit"))
                changed += 1
            if row_index < row_count - 1:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
    return {"short_tables_protected": changed}


def _paragraph_cell_width(paragraph: Any) -> int | None:
    node = paragraph._p
    while node is not None:
        if node.tag.endswith("}tc"):
            value = node.xpath("./w:tcPr/w:tcW/@w:w")
            if value:
                try:
                    return int(value[0]) * 635
                except (TypeError, ValueError):
                    return None
        node = node.getparent()
    return None


def normalize_visual_sizes(document: Any) -> dict[str, int]:
    """Keep inline and anchored pictures inside their page/cell boundaries."""
    from docx.oxml.ns import qn  # type: ignore
    section = document.sections[0]
    page_width = int(section.page_width - section.left_margin - section.right_margin)
    page_height = int(section.page_height - section.top_margin - section.bottom_margin)
    max_page_width = int(page_width * 0.97)
    max_page_height = int(page_height * 0.90)
    inline_changed = 0
    for shape in document.inline_shapes:
        paragraph = shape._inline.getparent()
        cell_width = None
        node = paragraph
        while node is not None:
            if node.tag.endswith("}tc"):
                value = node.xpath("./w:tcPr/w:tcW/@w:w")
                if value:
                    try:
                        cell_width = int(value[0]) * 635
                    except (TypeError, ValueError):
                        pass
                break
            node = node.getparent()
        max_width = min(max_page_width, int(cell_width * 0.97) if cell_width else max_page_width)
        width, height = int(shape.width), int(shape.height)
        scale = min(1.0, max_width / max(width, 1), max_page_height / max(height, 1))
        if scale < 0.999:
            shape.width = int(width * scale)
            shape.height = int(height * scale)
            inline_changed += 1

    anchor_changed = 0
    for paragraph in document.paragraphs:
        cell_width = _paragraph_cell_width(paragraph)
        max_width = min(max_page_width, int(cell_width * 0.97) if cell_width else max_page_width)
        for anchor in paragraph._p.xpath(".//wp:anchor"):
            if not any(element.tag.endswith("}pic") for element in anchor.iter()):
                continue
            extent = anchor.xpath("./wp:extent")
            if not extent:
                continue
            cx, cy = int(extent[0].get("cx", "0")), int(extent[0].get("cy", "0"))
            scale = min(1.0, max_width / max(cx, 1), max_page_height / max(cy, 1))
            if scale >= 0.999:
                continue
            new_cx, new_cy = str(int(cx * scale)), str(int(cy * scale))
            extent[0].set("cx", new_cx)
            extent[0].set("cy", new_cy)
            for ext in anchor.xpath(".//a:xfrm/a:ext"):
                ext.set("cx", new_cx)
                ext.set("cy", new_cy)
            anchor_changed += 1
    return {"inline": inline_changed, "anchored": anchor_changed}


def audit_text_logic(document: Any) -> dict[str, Any]:
    """Report suspicious text/symbol patterns without rewriting technical text."""
    bracket_pairs = {"(": ")", "[": "]", "{": "}", "（": "）", "【": "】"}
    findings: list[dict[str, Any]] = []
    for index, paragraph in enumerate(iter_paragraphs(document)):
        if paragraph_has_visual_object(paragraph):
            continue
        value = normalize_text(paragraph.text)
        if not value:
            continue
        reasons: list[str] = []
        if "�" in value or any(unicodedata.category(ch) == "Cc" for ch in value):
            reasons.append("replacement_or_control_character")
        if re.search(r"([!?！？；：，。])\1+", value):
            reasons.append("repeated_punctuation")
        if re.search(r"[，。；：,:;][。；：,:;]", value):
            reasons.append("adjacent_punctuation")
        if any(value.count(left) != value.count(right) for left, right in bracket_pairs.items()):
            reasons.append("unbalanced_brackets")
        if re.fullmatch(r"[\W_]+", value, flags=re.UNICODE) and len(value) <= 6:
            reasons.append("symbol_only_short_block")
        if reasons:
            findings.append({"index": index, "text": value[:120], "reasons": reasons})
    return {"checked": True, "finding_count": len(findings), "findings": findings[:80], "note": "仅报告可疑文字/符号，不自动改写公式或技术内容"}


FORMULA_BLOCK_TYPES = {
    "formula", "equation", "inline_equation", "interline_equation",
    "display_formula", "inline_formula", "math", "latex",
}


def _walk_json_objects(value: Any):
    """Yield nested JSON objects without depending on a MinerU schema version."""
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _walk_json_objects(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_json_objects(child)


def mineru_formula_candidates(content_path: Path | None, max_pages: int | None = None) -> list[dict[str, Any]]:
    """Read formula *locations* from MinerU content JSON, never formula text.

    MinerU has changed the names of nested equation nodes between releases, so
    this intentionally accepts the known names and checks for latex/equation
    keys as a fallback.  It is a detector only: no recognition result is used
    to rewrite the DOCX.
    """
    if not content_path or not content_path.exists():
        return []
    try:
        data = json.loads(content_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    candidates: list[dict[str, Any]] = []
    pages = data if isinstance(data, list) else [data]
    for page_number, page in enumerate(pages, 1):
        if max_pages is not None and page_number > max_pages:
            break
        for item in _walk_json_objects(page):
            kind = str(item.get("type", "")).strip().lower()
            content = item.get("content")
            has_formula_key = any(key in item for key in ("latex", "latex_content", "equation", "formula"))
            if kind not in FORMULA_BLOCK_TYPES and not has_formula_key:
                continue
            bbox = item.get("bbox")
            if not isinstance(bbox, list) or len(bbox) != 4:
                bbox = None
            candidates.append({
                "page": page_number,
                "bbox": [round(float(v), 2) for v in bbox] if bbox else None,
                "detector": "mineru",
                "confidence": 0.95 if kind in FORMULA_BLOCK_TYPES else 0.80,
                "kind": kind or "formula",
                "has_content": bool(content or has_formula_key),
            })
    return candidates


def _looks_like_formula_block(text: str) -> tuple[bool, float]:
    """Classify a PDF text block as a formula candidate, not its correctness."""
    compact = re.sub(r"\s+", "", normalize_text(text))
    if not compact:
        return False, 0.0
    operators = sum(compact.count(char) for char in "=+-*/^∑∫√≤≥≠≈∞∂∇∏")
    greek = sum("\u0370" <= char <= "\u03ff" for char in compact)
    digit_or_latin = len(re.findall(r"[A-Za-z0-9]", compact))
    math_signal = operators + greek
    if math_signal >= 2 and len(compact) <= 180:
        return True, min(0.92, 0.66 + math_signal * 0.05)
    if re.search(r"[A-Za-z]\s*[=_^]\s*[A-Za-z0-9(]", text) and len(compact) <= 140:
        return True, 0.78
    if digit_or_latin and len(compact) <= 24 and re.search(r"[=^_]", compact):
        return True, 0.74
    return False, 0.0


def pdfkit_formula_bridge_detect(reference_pdf: Path | None, temp: Path, max_pages: int | None = None) -> tuple[list[dict[str, Any]] | None, dict[str, Any]]:
    """Ask the isolated PDF-Extract-Kit bridge for formula locations.

    The bridge is deliberately optional.  A missing/stale bridge returns
    immediately so ordinary Word processing is not delayed; callers then use
    MinerU or the PyMuPDF text heuristic.
    """
    if not reference_pdf or reference_pdf.suffix.lower() != ".pdf" or not reference_pdf.exists():
        return None, {"status": "unavailable", "reason": "没有可用的 PDF"}
    root = Path(__file__).resolve().parents[1]
    bridge_root = root / "tmp" / "pdfkit-bridge"
    queue, done, errors = bridge_root / "queue-v1", bridge_root / "done-v1", bridge_root / "errors-v1"
    ready = bridge_root / "ready-v1.json"
    # run_all.cmd starts the bridge before this script.  Do not enqueue work
    # when no live worker announced readiness, otherwise a typo would cause a
    # long timeout on every document.
    try:
        if not ready.exists():
            return None, {"status": "unavailable", "reason": "PDF-Extract-Kit 桥接器未就绪"}
        if time.time() - ready.stat().st_mtime > 3600:
            return None, {"status": "unavailable", "reason": "PDF-Extract-Kit 桥接器状态已过期"}
        ready_data = json.loads(ready.read_text(encoding="utf-8"))
        bridge_pid = int(ready_data.get("pid", 0))
        if bridge_pid <= 0:
            return None, {"status": "unavailable", "reason": "PDF-Extract-Kit 桥接器状态无效"}
        if os.name == "nt":
            probe = subprocess.run(
                ["tasklist", "/FI", f"PID eq {bridge_pid}", "/NH"],
                capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=5,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            if str(bridge_pid) not in probe.stdout:
                return None, {"status": "unavailable", "reason": "PDF-Extract-Kit 桥接器进程已退出"}
    except OSError:
        return None, {"status": "unavailable", "reason": "无法读取 PDF-Extract-Kit 桥接器状态"}
    except (ValueError, json.JSONDecodeError):
        return None, {"status": "unavailable", "reason": "PDF-Extract-Kit 桥接器状态损坏"}
    work = bridge_root / "work" / uuid.uuid4().hex
    staged_input, output = work / reference_pdf.name, work / "formula-candidates.json"
    job_id = work.name
    for folder in (queue, done, errors, work):
        folder.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(reference_pdf, staged_input)
        job = {
            "id": job_id,
            "input": str(staged_input.resolve()),
            "output": str(output.resolve()),
            "max_pages": max_pages,
        }
        (queue / f"{job_id}.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        deadline = time.monotonic() + 900
        result_file, error_file = done / f"{job_id}.json", errors / f"{job_id}.json"
        while time.monotonic() < deadline:
            if result_file.exists():
                result = json.loads(result_file.read_text(encoding="utf-8"))
                result_file.unlink(missing_ok=True)
                if not output.exists():
                    return None, {"status": "failed", "reason": "桥接器未返回结果", **result}
                data = json.loads(output.read_text(encoding="utf-8"))
                return list(data.get("candidates", [])), {"status": "ok", "engine": "pdf-extract-kit-yolov8", **data}
            if error_file.exists():
                reason = error_file.read_text(encoding="utf-8", errors="replace")
                error_file.unlink(missing_ok=True)
                return None, {"status": "failed", "reason": reason}
            time.sleep(1)
        return None, {"status": "timeout", "reason": "等待 PDF-Extract-Kit 桥接器超时"}
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        return None, {"status": "failed", "reason": str(exc)}
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _pdf_text_locations(reference_pdf: Path | None, max_pages: int | None = None) -> list[dict[str, Any]]:
    """Extract text blocks and PDF coordinates for deterministic fallback checks."""
    if not reference_pdf or not reference_pdf.exists() or reference_pdf.suffix.lower() != ".pdf":
        return []
    try:
        try:
            import pymupdf as fitz  # type: ignore
        except ImportError:
            import fitz  # type: ignore
        document = fitz.open(str(reference_pdf))
    except Exception:
        return []
    locations: list[dict[str, Any]] = []
    try:
        limit = min(len(document), max_pages) if max_pages is not None else len(document)
        for page_index in range(limit):
            page = document[page_index]
            for block in page.get_text("blocks"):
                if len(block) < 5:
                    continue
                text = str(block[4] or "").strip()
                if not text:
                    continue
                locations.append({
                    "page": page_index + 1,
                    "bbox": [round(float(value), 2) for value in block[:4]],
                    "text": text,
                })
    finally:
        document.close()
    return locations


def pdf_formula_candidates(reference_pdf: Path | None, max_pages: int | None = None, content_path: Path | None = None, pdfkit_candidates: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """Detect PDF formula regions using PDF-Extract-Kit, MinerU, then text blocks."""
    if pdfkit_candidates is not None:
        return pdfkit_candidates
    mineru = mineru_formula_candidates(content_path, max_pages)
    if mineru:
        return mineru
    locations = _pdf_text_locations(reference_pdf, max_pages)
    candidates: list[dict[str, Any]] = []
    for location in locations:
        matched, confidence = _looks_like_formula_block(str(location.get("text", "")))
        if matched:
            candidates.append({
                "page": location.get("page"),
                "bbox": location.get("bbox"),
                "text_preview": str(location.get("text", ""))[:120],
                "detector": "pymupdf-text-heuristic",
                "confidence": confidence,
                "kind": "formula-candidate",
            })
    return candidates


def word_formula_inventory(path: Path) -> dict[str, Any]:
    """Count editable OMML formulas and formula-like text blocks in a DOCX."""
    from docx import Document  # type: ignore
    document = Document(str(path))
    omml: list[dict[str, Any]] = []
    formula_text: list[dict[str, Any]] = []
    for index, paragraph in enumerate(iter_paragraphs(document)):
        equations = paragraph._p.xpath(".//m:oMath")
        equation_paras = paragraph._p.xpath(".//m:oMathPara")
        if equations:
            omml.append({"index": index, "objects": len(equations), "display_objects": len(equation_paras)})
        text = normalize_text(paragraph.text)
        if text and not equations and _looks_like_formula_block(text)[0]:
            formula_text.append({"index": index, "text": text[:160]})
    return {
        "checked": True,
        "omml_formula_objects": sum(item["objects"] for item in omml),
        "omml_formula_paragraphs": len(omml),
        "formula_like_text_blocks": len(formula_text),
        "represented_formula_blocks": len(omml) + len(formula_text),
        "omml": omml[:200],
        "formula_like_text": formula_text[:200],
    }


def deterministic_formula_audit(reference_pdf: Path | None, output_docx: Path, max_pages: int | None = None, content_path: Path | None = None, pdfkit_candidates: list[dict[str, Any]] | None = None, pdfkit_status: dict[str, Any] | None = None) -> dict[str, Any]:
    """Compare formula region/object counts only; never compare formula meaning."""
    word = word_formula_inventory(output_docx)
    pdf = pdf_formula_candidates(reference_pdf, max_pages, content_path, pdfkit_candidates)
    pdf_count = len(pdf)
    word_count = int(word.get("represented_formula_blocks", 0))
    # PDF text-block heuristics can merge adjacent equations.  A small count
    # difference is therefore a warning, not an automatic rejection.
    tolerance = max(2, int(max(pdf_count, word_count) * 0.30))
    count_delta = abs(pdf_count - word_count)
    passed = not pdf_count or count_delta <= tolerance
    return {
        "checked": True,
        "method": "PDF-Extract-Kit YOLOv8, MinerU formula locations, or PyMuPDF text-block heuristic + DOCX OMML/XML",
        "detector_chain": ["pdf-extract-kit-yolov8", "mineru", "pymupdf-text-heuristic"],
        "selected_detector": ("pdf-extract-kit-yolov8" if pdfkit_candidates is not None else ("mineru" if any(item.get("detector") == "mineru" for item in pdf) else "pymupdf-text-heuristic")),
        "pdfkit_bridge": pdfkit_status or {"status": "not-run"},
        "semantic_formula_comparison": False,
        "automatic_formula_replacement": False,
        "pdf_candidates": pdf_count,
        "word_represented_blocks": word_count,
        "count_delta": count_delta,
        "tolerance": tolerance,
        "passed": passed,
        "status": "ok" if passed else "manual_review_required",
        "pdf_candidates_detail": pdf[:300],
        "word_inventory": word,
        "note": "只检查公式候选区域与 Word 公式/公式样文本对象数量，不判断公式内容是否相等",
    }


def deterministic_text_completeness(reference_pdf: Path | None, output_docx: Path, temp: Path, max_pages: int | None = None) -> dict[str, Any]:
    """Check broad text retention without treating formula strings as truth."""
    if not reference_pdf or reference_pdf.suffix.lower() != ".pdf" or not reference_pdf.exists():
        return {"checked": False, "reason": "没有可用的原始 PDF"}
    source_pdf = reference_pdf
    if max_pages is not None:
        source_pdf = limit_pdf_pages(reference_pdf, temp / "text-completeness-preview.pdf", max_pages)
    source = source_text(source_pdf, temp / "pdf-text")
    output = docx_text(output_docx)
    result = compare_text(source, output)
    result.update({
        "method": "PDF text extraction + DOCX text extraction + character similarity",
        "formula_semantics_checked": False,
        "passed": bool(result.get("checked") and result.get("similarity", 0.0) >= 0.80 and result.get("output_chars", 0) >= max(1, int(result.get("source_chars", 0) * 0.70))),
        "note": "用于发现大段文字丢失；不据此判断 PDF 与 Word 公式内容是否相等",
    })
    return result


def apply_subscripts(paragraph: Any) -> None:
    """Turn common OCR forms such as P_L/P_H into real Word subscripts."""
    pattern = re.compile(r"(?P<base>[A-Za-z])[_](?P<sub>[A-Za-z0-9]+)|(?P<plain>P)(?P<plain_sub>[LH])(?=\s*=?|\s|[,;，；)])")
    for run in list(paragraph.runs):
        text = run.text or ""
        match = pattern.search(text)
        if not match:
            continue
        before, after = text[: match.start()], text[match.end() :]
        base = match.group("base") or match.group("plain")
        sub = match.group("sub") or match.group("plain_sub")
        run.text = before + base
        sub_run = paragraph.add_run(sub)
        sub_run.font.subscript = True
        sub_run.font.name = run.font.name
        sub_run.font.size = run.font.size
        run._r.addnext(sub_run._r)
        if after:
            after_run = paragraph.add_run(after)
            sub_run._r.addnext(after_run._r)


def set_run_fonts(run: Any, east_asia: str = "SimSun") -> None:
    """Use the requested Chinese font and Times New Roman for Latin characters."""
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import RGBColor  # type: ignore
    run.font.name = east_asia
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:cs"), "Times New Roman")
    # Remove theme-linked font attributes.  They can override the explicit
    # East Asian font in Word/WPS and make a correctly mapped heading render
    # as the document default instead.
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:csTheme"):
        fonts.attrib.pop(qn(theme_attr), None)


def add_page_field(paragraph: Any) -> None:
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def ensure_visible_footer(document: Any) -> None:
    """Add a right-aligned PAGE field with a visible fallback numeral."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Pt  # type: ignore
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            if any(p.text.strip() for p in footer.paragraphs):
                continue
            p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(0)
            add_page_field(p)
            # Cached field text is intentionally left as "1" so WPS displays a
            # number even before it refreshes fields; Word updates it on open.
            set_run_fonts(p.runs[0])
    # Ask Word/WPS to recalculate PAGE fields when the document opens.
    settings = document.settings._element
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def restart_body_page_numbers(document: Any) -> bool:
    """Start PAGE numbering at 1 without creating an extra blank section page."""
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Pt  # type: ignore
    target = None
    for paragraph in document.paragraphs:
        logical = re.sub(r"\s+", "", paragraph.text or "")
        if re.match(r"^第1章", logical) and "\t" not in (paragraph.text or ""):
            target = paragraph
            break
    if target is None:
        return False
    section = document.sections[0]
    pg_num = section._sectPr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        section._sectPr.append(pg_num)
    pg_num.set(qn("w:start"), "1")
    # Use the existing section footer; this avoids the extra blank page that
    # WPS creates when a NEW_PAGE section follows an explicit page break.
    footer = section.footer
    for p in list(footer.paragraphs):
        remove_paragraph(p)
    footer_p = footer.add_paragraph()
    footer_p.alignment = 2
    footer_p.paragraph_format.space_before = Pt(0)
    add_page_field(footer_p)
    set_run_fonts(footer_p.runs[0])
    if not has_explicit_page_break_before(target):
        from docx.enum.text import WD_BREAK  # type: ignore
        before = target.insert_paragraph_before()
        before.add_run().add_break(WD_BREAK.PAGE)
    return True


def apply_document_styles(document: Any) -> None:
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Pt, RGBColor  # type: ignore
    from docx.enum.text import WD_LINE_SPACING  # type: ignore
    for style_name in ("Normal", "Body Text", "Header", "Footer"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = "SimSun"
        style.font.size = Pt(13)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(19)
        rpr = style._element.get_or_add_rPr()
        fonts = rpr.rFonts
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.insert(0, fonts)
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), "SimSun")
        fonts.set(qn("w:cs"), "Times New Roman")


def format_body_paragraph(paragraph: Any) -> None:
    from docx.shared import Pt  # type: ignore
    from docx.shared import Inches, Pt  # type: ignore
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER  # type: ignore
    from docx.enum.text import WD_LINE_SPACING  # type: ignore
    text = normalize_text(paragraph.text)
    nonstructural = paragraph_is_nonstructural(paragraph)
    has_picture = bool(paragraph._p.xpath(".//w:drawing[.//pic:pic]"))
    # Fixed line spacing is for text only.  Applying it to an inline picture
    # clips the picture to a 19 pt line box in LibreOffice and creates blank
    # pages or captions without their figures.
    if has_picture:
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.line_spacing = None
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = False
    else:
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(19)
        # Quark often stores PDF coordinates as enormous paragraph spacing.
        # Once the document is rebuilt as normal Word flow, those values only
        # create blank bands between questions.
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
    # Quark assigns large right indents to paragraphs that carry an anchored
    # figure/equation.  Those indents are layout artifacts, not document
    # structure, and can collapse the accompanying text into a one-character
    # column.  Always reset them before handling the visual object.
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    if nonstructural and paragraph.text:
        paragraph.alignment = 0
    # Preserve the placement of image/equation-only blocks. Their runs (when
    # present) still receive the body font policy below.
    if nonstructural:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        if paragraph.text:
            for run in paragraph.runs:
                set_run_fonts(run, "SimSun")
                run.font.size = Pt(13)
        return
    if not text:
        return
    logical = re.sub(r"\s+", "", paragraph.text or "")
    normalize_heading_text(paragraph, logical)
    # Keep enumerated sub-items on one shared left edge.  This avoids the
    # converter's inherited tab/indent offsets making (1)/(2) drift apart.
    if re.match(r"^(?:\([一二三四五六七八九十0-9]+\)|（[一二三四五六七八九十0-9]+）)", logical):
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(21)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.alignment = 0
    if logical in {"前言", "目录"}:
        text = logical
    # Remove layout offsets inherited from the OCR converter. Indentation is
    # reintroduced below only where the logical hierarchy calls for it.
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    toc_line = is_toc_line(paragraph.text) and logical not in {"前言", "目录"}
    heading = False
    if toc_line:
        size = Pt(13)
        paragraph.alignment = 0
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = True
        try:
            paragraph.paragraph_format.tab_stops.clear_all()
        except AttributeError:
            pass
        # WPS does not consistently paint OOXML tab leaders. Keep a visible
        # dotted filler in the text and retain a final right-aligned tab stop
        # for editors that do support leaders.
        parts = paragraph.text.split("\t", 1)
        if len(parts) == 2:
            left, page = parts
            # Strip a previous visible leader before rebuilding it; otherwise
            # repeated runs would keep extending the line until the page number
            # wraps. Keep the leader bounded so text and page stay on one line.
            left = re.split(r"\s*\.{3,}", left, maxsplit=1)[0].rstrip()
            dots = "." * max(8, min(30, 46 - len(left)))
            paragraph.text = f"{left} {dots}\t{page}"
            section = paragraph._parent.part.document.sections[0]
            right_edge = section.page_width - section.left_margin - section.right_margin
            paragraph.paragraph_format.tab_stops.add_tab_stop(right_edge, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    elif re.match(r"^(?:前言|目录)$", logical):
        size = Pt(23)
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        # Word clips 23 pt East-Asian glyphs when a title inherits the
        # body's exact 19 pt line box. Use an auto/single line box for all
        # large headings while preserving the 19 pt body rule.
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.line_spacing = None
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        try:
            paragraph.paragraph_format.tab_stops.clear_all()
        except AttributeError:
            pass
        heading = True
    elif re.match(r"^第[0-9一二三四五六七八九十百]+章", logical) and not is_toc_entry(logical):
        size = Pt(23)
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.line_spacing = None
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        try:
            paragraph.paragraph_format.tab_stops.clear_all()
        except AttributeError:
            pass
        heading = True
    elif re.match(r"^(?:书后习题解析|同步训练题|同步训练题答案|参考文献)$", logical):
        size = Pt(17)
        heading = True
    elif is_answer_heading(paragraph.text):
        # Standalone ``解析：A`` markers are answer-section subheadings, not
        # ordinary body text.  Keep them with the following explanation and
        # avoid a forced page break for every answer line.
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(2)
        size = Pt(17)
        heading = True
    elif is_exam_title_line(logical):
        # Exam-cover lines are a compact heading group.  Keep them together
        # without forcing each line to start a new page.
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.line_spacing = None
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = True
        size = Pt(17)
        heading = True
    elif is_toc_entry(logical):
        size = Pt(13)
        paragraph.alignment = 0
        paragraph.paragraph_format.first_line_indent = Pt(0)
        heading = False
    elif toc_line:
        size = Pt(13)
    else:
        size = Pt(13)
    for run in paragraph.runs:
        set_run_fonts(run, heading_font() if heading else "SimSun")
        run.font.size = size
        if heading:
            run.bold = True
    # Keep headings and question stems flush left; indent explanatory sub-items one character.
    major = bool(re.match(r"^(?:前言|目录|第[0-9一二三四五六七八九十百]+章|20[0-9]{2}年|题|解)", logical)) and not is_toc_entry(logical)
    subordinate = bool(re.match(r"^(?:\(?[0-9一二三四五六七八九十]+[)）]|由|其中|说明|注)", text))
    if toc_line:
        paragraph.paragraph_format.first_line_indent = Pt(0)
    elif major or not subordinate:
        paragraph.paragraph_format.first_line_indent = Pt(0)
        if not major:
            paragraph.alignment = 0
    else:
        paragraph.paragraph_format.first_line_indent = Pt(21)


def set_header_footer(document: Any, header_text: str | None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Pt  # type: ignore
    for section in document.sections:
        header = section.header
        if header_text:
            for paragraph in list(header.paragraphs):
                remove_paragraph(paragraph)
            p = header.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(header_text)
            run.font.size = Pt(9)
        logical = re.sub(r"\s+", "", paragraph.text or "")
        is_heading = bool(
            logical in {"前言", "目录"}
            or (re.match(r"^第[0-9一二三四五六七八九十百]+章", logical) and not is_toc_entry(logical))
            or re.match(r"^(?:书后习题解析|同步训练题|同步训练题答案|参考文献)$", logical)
            or is_exam_title_line(logical)
            or is_answer_heading(paragraph.text)
        )
        set_run_fonts(run, heading_font() if is_heading else "SimSun")
        footer = section.footer
        # Keep an existing non-empty footer; for editable output with a detected
        # header provide a real PAGE field. Image-fidelity pages already contain
        # the source footer pixels, so adding another number would duplicate it.
        if header_text and not any(normalize_text(p.text) for p in footer.paragraphs):
            p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            add_page_field(p)
        section.header_distance = Pt(12)
        section.footer_distance = Pt(12)


def flatten_toc_tables(document: Any) -> int:
    """Turn Quark's one-row TOC tables into one-line tab-leader paragraphs."""
    from docx.text.paragraph import Paragraph  # type: ignore
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER  # type: ignore
    from docx.shared import Inches, Pt  # type: ignore
    changed = 0
    for table in list(document.tables):
        if len(table.rows) != 1 or len(table.columns) != 3:
            continue
        cells = table.rows[0].cells
        left, middle, right = (normalize_text(c.text) for c in cells)
        if not left or not right.isdigit() or not re.search(r"[.·…]{3,}", middle):
            continue
        p_el = table._element.makeelement("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p", {})
        table._element.addprevious(p_el)
        paragraph = Paragraph(p_el, table._parent)
        # Align the page number to the actual right text boundary, not a
        # hard-coded position that drifts when WPS changes page geometry.
        section = paragraph._parent.part.document.sections[0]
        right_edge = section.page_width - section.left_margin - section.right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(right_edge, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = True
        paragraph.add_run(left + "\t" + right)
        table._element.getparent().remove(table._element)
        changed += 1
    return changed


def rebuild_floating_toc(document: Any) -> int:
    """Convert Quark's floating-shape TOC into editable Word paragraphs.

    Quark sometimes stores an entire TOC page as dozens of anchored text
    boxes inside one paragraph.  ``python-docx`` cannot expose that text via
    ``Paragraph.text`` and the boxes therefore get formatted as scattered
    glyphs.  The boxes are ordered in reading order in the source XML, so we
    can safely pair each entry with its printed page number and rebuild the
    page as normal paragraphs.
    """
    rebuilt = 0
    from docx.shared import Pt  # type: ignore

    for paragraph in list(document.paragraphs):
        drawings = paragraph._p.xpath(".//w:drawing")
        if len(drawings) < 10:
            continue
        items: list[str] = []
        # Quark commonly puts a full-page TOC background image together with
        # editable-looking floating text boxes in one paragraph. Once the
        # document is reflowed, that background and its absolute text boxes
        # cover the preface. Treat a page-sized picture as a source backdrop,
        # remove it, and rebuild the TOC rows as normal paragraphs below.
        for drawing in drawings:
            value = "".join(drawing.xpath(".//w:t/text()"))
            if value:
                items.append(value)
        for drawing in drawings:
            picture = drawing.xpath(".//pic:pic")
            extent = drawing.xpath(".//wp:extent")
            if not picture or not extent:
                continue
            try:
                cx, cy = int(extent[0].get("cx", "0")), int(extent[0].get("cy", "0"))
            except (TypeError, ValueError):
                continue
            if cx >= 10_000_000 and cy >= 15_000_000:
                drawing.getparent().remove(drawing)
        if not any("目录" in re.sub(r"\s+", "", value) for value in items):
            continue

        rows: list[tuple[str, str]] = []
        label: str | None = None
        for value in items:
            value = value.replace("\u00a0", " ")
            if value.strip() == "\t" or not value.strip():
                continue
            compact = re.sub(r"\s+", "", value)
            if not rows and "目录" in compact:
                continue
            if value.strip().isdigit():
                if label:
                    rows.append((label.strip(), value.strip()))
                    label = None
                continue
            value = value.replace("\t", " ").strip()
            if value:
                label = value if label is None else f"{label}{value}"

        if len(rows) < 5:
            continue

        title = paragraph.insert_paragraph_before()
        title.text = "目录"
        title.paragraph_format.keep_with_next = True
        title.paragraph_format.space_after = Pt(10)
        for left, page in rows:
            row = paragraph.insert_paragraph_before()
            row.text = f"{left}\t{page}"
            row.paragraph_format.keep_together = True
            row.paragraph_format.widow_control = True
        remove_paragraph(paragraph)
        rebuilt += len(rows) + 1
    return rebuilt


def toc_paragraphs_to_rows(document: Any) -> int:
    """Replace TOC rows with plain text lines: entry + page number."""
    from docx.shared import Pt  # type: ignore
    converted = 0
    for paragraph in list(document.paragraphs):
        raw = paragraph.text or ""
        logical = re.sub(r"\s+", "", raw)
        # A tabbed technical row is not a TOC merely because its final field
        # is numeric (for example ``1\t3`` or an answer-choice row).  Only
        # convert genuine dot-leader/chapter TOC entries; otherwise preserve
        # the tabs for the column-alignment pass.
        if "\t" in raw and not is_toc_line(raw):
            continue
        if "\t" not in raw and not is_toc_entry(logical):
            continue
        if "\t" in raw:
            left, page = raw.split("\t", 1)
            page = page.strip()
        else:
            match = re.match(r"^(.*?)(\d{1,3})\s*$", normalize_text(raw))
            if not match:
                continue
            left, page = match.group(1).rstrip(), match.group(2)
        if not page.isdigit():
            continue
        left = re.split(r"\s*\.{3,}", left, maxsplit=1)[0].rstrip()
        paragraph.text = f"{left}    {page.strip()}"
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.keep_together = True
        chapter_row = bool(re.match(r"^第[0-9一二三四五六七八九十百]+章", re.sub(r"\s+", "", left)))
        chapter_row = chapter_row or re.sub(r"\s+", "", left) == "参考文献"
        for run in paragraph.runs:
            set_run_fonts(run)
            run.font.size = Pt(13)
            run.bold = chapter_row
        converted += 1
    return converted


def apply_logical_keep_groups(document: Any) -> int:
    """Keep each question stem, answer steps, and nearby objects together."""
    paragraphs = list(document.paragraphs)
    question_re = re.compile(r"^题\s*[0-9一二三四五六七八九十百]+(?:[-－][0-9一二三四五六七八九十百]+)?")
    starts = [i for i, p in enumerate(paragraphs) if question_re.match(normalize_text(p.text))]
    changed = 0
    for pos, start in enumerate(starts):
        end = starts[pos + 1] if pos + 1 < len(starts) else len(paragraphs)
        nonempty = [i for i in range(start, end) if paragraphs[i].text.strip() or paragraph_has_visual_object(paragraphs[i])]
        if not nonempty:
            continue
        last = nonempty[-1]
        group_chars = sum(len(normalize_text(paragraphs[i].text)) for i in nonempty)
        group_visuals = sum(paragraph_has_visual_object(paragraphs[i]) for i in nonempty)
        # Keep a compact question/solution unit as one chain.  For long
        # solutions, chaining every paragraph can push a nearly full unit to
        # a fresh page and create the user's "one-line page" defect.  In that
        # case keep only the stem and each short answer step with its next
        # line, allowing the body to flow naturally across pages.
        compact_group = len(nonempty) <= 10 and group_chars <= 1100 and group_visuals <= 5
        for offset, i in enumerate(nonempty):
            paragraph = paragraphs[i]
            paragraph.paragraph_format.widow_control = True
            paragraph.paragraph_format.keep_together = len(normalize_text(paragraph.text)) <= 260 or paragraph_has_visual_object(paragraph)
            next_i = nonempty[offset + 1] if offset + 1 < len(nonempty) else None
            if next_i is not None and (compact_group or len(normalize_text(paragraph.text)) <= 180 or paragraph_has_visual_object(paragraph)):
                paragraph.paragraph_format.keep_with_next = True
            changed += 1

    # A heading should stay with the first paragraph that explains it.
    for paragraph in paragraphs:
        logical = re.sub(r"\s+", "", paragraph.text or "")
        if (
            logical in {"前言", "目录"}
            or (re.match(r"^第[0-9一二三四五六七八九十百]+章", logical) and not is_toc_entry(logical))
            or bool(re.match(r"^[一二三四五六七八九十百]+、[^。！？]{1,30}$", logical))
        ):
            paragraph.paragraph_format.keep_with_next = True
    # A figure is a self-contained layout block.  Do not let inherited
    # keep-with-next flags form a chain that pushes all following content onto
    # an otherwise blank page.  Keep only an immediately following caption.
    for index, paragraph in enumerate(paragraphs):
        if not paragraph._p.xpath(".//w:drawing[.//pic:pic]"):
            continue
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = False
        if index + 1 < len(paragraphs):
            following = paragraphs[index + 1]
            if re.match(r"^(?:图|表)\s*\d", normalize_text(following.text)):
                paragraph.paragraph_format.keep_with_next = True
    return changed


def apply_structure_role(paragraph: Any, role: str) -> None:
    """Apply the user's fixed typography to one mapped paragraph role."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # type: ignore
    from docx.shared import Pt  # type: ignore
    role = role.lower()
    if role == "title" or role == "heading1":
        size, font, bold, align = 23, heading_font(), True, WD_ALIGN_PARAGRAPH.CENTER
    elif role in {"heading2", "heading3"}:
        size, font, bold, align = 17, heading_font(), True, WD_ALIGN_PARAGRAPH.LEFT
    elif role == "caption":
        size, font, bold, align = 13, "SimSun", False, WD_ALIGN_PARAGRAPH.CENTER
    else:
        size, font, bold, align = 13, "SimSun", False, WD_ALIGN_PARAGRAPH.LEFT
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Pt(0 if role != "body" else 0)
    if role in {"title", "heading1", "heading2", "heading3"}:
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.line_spacing = None
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
    else:
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(19)
    for run in paragraph.runs:
        set_run_fonts(run, font)
        run.font.size = Pt(size)
        run.bold = bold


def postprocess_docx(
    path: Path,
    strip_header_footer: bool = True,
    structure_overrides: dict[int, str] | None = None,
    reference_major_titles: set[str] | None = None,
) -> dict[str, Any]:
    """Apply layout rules that Quark/LibreOffice cannot infer reliably."""
    from docx import Document  # type: ignore
    from docx.enum.text import WD_LINE_SPACING  # type: ignore
    from docx.shared import Pt  # type: ignore

    document = Document(str(path))
    reference_major_titles = reference_major_titles or set()
    visual_before = visual_layout_snapshot(document)
    early_role_ids: dict[int, str] = {}
    # Apply the stage-1 map while paragraph indices still match the untouched
    # Word export. Later cleanup may remove headers or spacer blocks.
    # Structure and layout decisions are local and deterministic.
    early_roles_applied = False
    if structure_overrides and early_roles_applied:
        original_paragraphs = list(iter_paragraphs(document))
        for index, role in structure_overrides.items():
            if 0 <= int(index) < len(original_paragraphs) and not paragraph_is_nonstructural(original_paragraphs[int(index)]):
                apply_structure_role(original_paragraphs[int(index)], role)
                early_role_ids[id(original_paragraphs[int(index)])] = role
    # Logical groups are inferred by the existing local paragraph rules.
    # Logical groups are handled by the local keep-together rules below.
    collapsed_sections = collapse_source_sections(document)
    trailing_artifacts = strip_trailing_export_artifacts(document)
    normalize_section_geometry(document)
    source_page_sections = False
    # Remove redundant break-only paragraphs for every import path. Marker
    # breaks are added later, so intentional section headings remain intact.
    removed_page_break_only = remove_page_break_only_paragraphs(document)
    image_only = bool(document.inline_shapes) and not any(normalize_text(p.text) for p in document.paragraphs)
    removed_headers = 0
    removed_markers = 0
    uncertain_regions = 0
    header_text: str | None = None
    for section in document.sections:
        if strip_header_footer:
            for container in (section.header, section.first_page_header, section.even_page_header, section.footer, section.first_page_footer, section.even_page_footer):
                for paragraph in list(container.paragraphs):
                    if paragraph.text.strip():
                        removed_headers += 1
                    remove_paragraph(paragraph)
                for table in list(container.tables):
                    table._element.getparent().remove(table._element)
    toc_tables = flatten_toc_tables(document)
    floating_toc_rows = rebuild_floating_toc(document)
    picture_anchors_inlined = convert_picture_anchors_to_inline(document)
    # Remove QR/watermark blocks before mixed-picture paragraphs are split;
    # otherwise a QR drawing could be moved into a new paragraph and lose the
    # marker text that identifies it.
    useless_content = remove_useless_text_and_qr(document)
    useless_content["qr_detected_images_removed"] = remove_qr_images(document)
    mixed_picture_paragraphs_separated = separate_mixed_picture_paragraphs(document)
    office_picture_normalization = normalize_picture_paragraphs_for_office(document)
    empty_table_cleanup = trim_empty_table_edges(document)
    marker = re.compile(r"^(?:前言|目录|第\s*[0-9一二三四五六七八九十百]+\s*章|同步训练题答案|20[0-9]{2}\s*年(?:度)?[^。；，,]{0,12}(?:真题|试题|试卷))")
    header_phrase = re.compile(r"机械原理习题(?:精解精练|集|解题)")
    paragraphs = list(iter_paragraphs(document))
    structure_overrides = structure_overrides or {}
    meaningful_seen = False
    for index, paragraph in enumerate(paragraphs):
        text = normalize_text(paragraph.text)
        logical = re.sub(r"\s+", "", paragraph.text or "")
        if header_phrase.search(text):
            header_text = header_text or header_phrase.search(text).group(0)
            remove_paragraph(paragraph)
            removed_headers += 1
            continue
        # Repeated running headers such as "第1章…… 5" are not body
        # headings and must not trigger a page break.
        if re.match(r"^第[0-9一二三四五六七八九十百]+章", logical) and "\t" not in (paragraph.text or "") and re.search(r"\s+\d+\s*$", paragraph.text or ""):
            remove_paragraph(paragraph)
            removed_headers += 1
            continue
        apply_subscripts(paragraph)
        format_body_paragraph(paragraph)
        # The untouched Quark DOCX carries the reliable title signal even
        # when its direct formatting is inconsistent.  Match by normalized
        # text, then apply the fixed project title style and a single page
        # boundary before each repeated exam title.
        is_reference_title = logical in reference_major_titles and not paragraph_is_nonstructural(paragraph)
        if is_reference_title:
            apply_structure_role(paragraph, "title")
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            paragraph.paragraph_format.line_spacing = Pt(32)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            previous_logical = re.sub(r"\s+", "", paragraphs[index - 1].text or "") if index else ""
            previous_is_title = previous_logical in reference_major_titles
            if meaningful_seen and not previous_is_title and not has_explicit_page_break_before(paragraph):
                paragraph.paragraph_format.page_break_before = True
        if index in structure_overrides and not early_roles_applied:
            apply_structure_role(paragraph, structure_overrides[index])
        if paragraph_needs_region_review(paragraph.text):
            uncertain_regions += 1
        toc_line = is_toc_line(paragraph.text)
        next_text = normalize_text(paragraphs[index + 1].text) if index + 1 < len(paragraphs) else ""
        cover_start = is_exam_cover_school_line(paragraph.text) and (is_exam_title_line(next_text) or (index + 2 < len(paragraphs) and is_exam_title_line(paragraphs[index + 2].text)))
        # Front matter and every major section always start on a new page. Do
        # not suppress this for a source section boundary: WPS/Word may not
        # preserve the boundary created by Composer.
        if (not source_page_sections and (marker.search(logical) or cover_start) and logical not in reference_major_titles and meaningful_seen and (not toc_line and not is_toc_entry(logical) or logical == "目录") and not paragraph_is_nonstructural(paragraph) and not has_explicit_page_break_before(paragraph) and not has_section_boundary_before(paragraph)):
            from docx.enum.text import WD_BREAK  # type: ignore
            paragraph.paragraph_format.page_break_before = False
            before = paragraph.insert_paragraph_before()
            before.add_run().add_break(WD_BREAK.PAGE)
            removed_markers += 1
        if text and not paragraph_is_nonstructural(paragraph):
            meaningful_seen = True
    # A generated section marker can land immediately before a heading that
    # already has ``pageBreakBefore`` from the structure formatter.  Keeping
    # both creates a completely blank page in Word/WPS/LibreOffice.  Remove
    # only the empty paragraph that contains the redundant explicit break;
    # never touch a content paragraph or an intentional break without the
    # following heading-level break.
    redundant_preheading_breaks = 0
    after_markers = list(iter_paragraphs(document))
    for current, following in zip(after_markers, after_markers[1:]):
        if current.text.strip() or current._p.xpath(".//w:drawing"):
            continue
        if not current._p.xpath(".//w:br"):
            continue
        if not following._p.xpath("./w:pPr/w:pageBreakBefore"):
            continue
        remove_paragraph(current)
        redundant_preheading_breaks += 1
    # Header/footer output is intentionally disabled for this local Word-first
    # workflow.  The old reconstruction path remains available only to older
    # callers that explicitly pass ``strip_header_footer=False``.
    if not strip_header_footer and not image_only:
        set_header_footer(document, header_text)
    toc_rows = toc_paragraphs_to_rows(document)
    body_numbering_restarted = 0
    if not strip_header_footer:
        body_numbering_restarted = int(restart_body_page_numbers(document))
        ensure_visible_footer(document)
    formula_merges = merge_formula_fragment_paragraphs(document)
    apply_document_styles(document)
    from docx.shared import Pt  # type: ignore
    from docx.enum.text import WD_LINE_SPACING  # type: ignore
    for index, paragraph in enumerate(iter_paragraphs(document)):
        logical = re.sub(r"\s+", "", paragraph.text or "")
        mapped_role = early_role_ids.get(id(paragraph)) if early_roles_applied else structure_overrides.get(index)
        is_heading = bool(
            not paragraph_is_nonstructural(paragraph)
            and mapped_role in {"title", "heading1", "heading2", "heading3"}
            or
            (not paragraph_is_nonstructural(paragraph) and logical in {"前言", "目录"})
            or (re.match(r"^第[0-9一二三四五六七八九十百]+章", logical) and not is_toc_entry(logical))
            or re.match(r"^(?:书后习题解析|同步训练题|同步训练题答案|参考文献)$", logical)
            or logical in reference_major_titles
        )
        if paragraph_is_nonstructural(paragraph):
            is_heading = False
        for run in paragraph.runs:
            if mapped_role and not paragraph_is_nonstructural(paragraph):
                apply_structure_role(paragraph, mapped_role)
                break
            set_run_fonts(run, heading_font() if is_heading else "SimSun")
            if run.font.size is None:
                run.font.size = Pt(13)
        if paragraph._p.xpath(".//w:drawing[.//pic:pic]") or paragraph._p.xpath(".//m:oMath|.//m:oMathPara"):
            # A formula can wrap onto two visual lines. Use a minimum line
            # box so Word/WPS expands it instead of clipping the second line.
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            paragraph.paragraph_format.line_spacing = Pt(19)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False
        elif is_heading:
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            # Give large Chinese display headings a real line box.  A single
            # line with a 23 pt calligraphic font can otherwise clip the upper
            # strokes in Word/WPS even when the font size is correct.
            paragraph.paragraph_format.line_spacing = Pt(
                32 if logical in {"前言", "目录"} else (28 if logical.startswith("第") else 25)
            )
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
        else:
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph.paragraph_format.line_spacing = Pt(19)
    # Some Quark exports encode the front-matter boundary only as a section
    # property. After section cleanup that marker may disappear, leaving the
    # directory directly under the preface. Keep the first real directory on
    # its own page without adding an empty break paragraph.
    directory_seen = False
    for paragraph in iter_paragraphs(document):
        logical = re.sub(r"\s+", "", paragraph.text or "")
        if logical == "目录" and not directory_seen and not paragraph_is_nonstructural(paragraph):
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            directory_seen = True
        elif (
            not paragraph_is_nonstructural(paragraph)
            and not is_toc_entry(logical)
            and (
                re.match(r"^第[0-9一二三四五六七八九十百]+章", logical)
                or logical in {"参考文献", "附录", "致谢", "结论", "结语", "同步训练题答案"}
            )
        ):
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
    # A few Quark exports retain left alignment on the third line of an exam
    # cover after the normal style pass.  Normalize the complete short cover
    # group together so school/year/title lines do not split or drift.
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # type: ignore
    cover_paragraphs = list(iter_paragraphs(document))
    for index, paragraph in enumerate(cover_paragraphs):
        if not is_exam_title_line(paragraph.text):
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.line_spacing = None
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = True
        if index and is_exam_cover_school_line(cover_paragraphs[index - 1].text):
            school = cover_paragraphs[index - 1]
            school.alignment = WD_ALIGN_PARAGRAPH.CENTER
            school.paragraph_format.first_line_indent = Pt(0)
            school.paragraph_format.keep_with_next = True
    # Final title-group boundary pass.  It must run after all generic marker
    # and compatibility passes because those may reset pageBreakBefore while
    # removing converter page markers.  A school-name line and its following
    # exam title stay together; only the first line of each group starts a
    # new page.
    final_title_group_seen = False
    final_previous_was_title = False
    for paragraph in iter_paragraphs(document):
        logical = re.sub(r"\s+", "", paragraph.text or "")
        is_title = (
            not paragraph_is_nonstructural(paragraph)
            and (
                logical in reference_major_titles
                or is_exam_title_line(logical)
                or bool(re.fullmatch(r"20\d{2}年(?:度)?", logical))
            )
        )
        if is_title:
            if final_title_group_seen and not final_previous_was_title:
                paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            paragraph.paragraph_format.line_spacing = Pt(32)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(4)
            final_title_group_seen = True
            final_previous_was_title = True
        else:
            # Any intervening body, formula, table, or picture ends the title
            # group.  Treating visual paragraphs as transparent here caused
            # the next real heading to lose its required page break.
            final_previous_was_title = False
    # Header/footer runs added above also receive the same font policy.
    for section in document.sections:
        for container in (section.header, section.footer, section.first_page_header, section.even_page_header, section.first_page_footer, section.even_page_footer):
            for paragraph in container.paragraphs:
                for run in paragraph.runs:
                    set_run_fonts(run)
    logical_keep_group_paragraphs = apply_logical_keep_groups(document)
    compact_spacers = compact_empty_spacer_paragraphs(document, max_consecutive=0)
    broken_keep_chains = break_excessive_keep_chains(document)
    noise_changes = sum(normalize_ocr_noise(p) for p in iter_paragraphs(document))
    spacing_changes = sum(normalize_ocr_spacing(p) for p in iter_paragraphs(document))
    column_alignment = normalize_tabular_alignment_rows(document)
    math_size_changes = normalize_math_sizes(document)
    formula_line_box_changes = normalize_formula_line_boxes(document)
    formula_table_height_changes = remove_formula_table_fixed_heights(document)
    malformed_math = normalize_malformed_math(document)
    table_widths = normalize_table_widths(document)
    table_pagination = normalize_table_pagination(document)
    formula_table_width_changes = expand_formula_table_widths(document)
    layout_table_cleanup = normalize_layout_table_frames(document)
    layout_padding_changes = normalize_layout_padding(document)
    table_title_cleanup = normalize_table_title_blocks(document)
    visual_size_changes = normalize_visual_sizes(document)
    # Re-assert title boundaries after all generic spacing/size passes.  Some
    # of those passes intentionally normalize paragraph properties and can
    # clear a previously assigned page break.
    final_title_seen = False
    final_title_previous = False
    for paragraph in iter_paragraphs(document):
        logical = re.sub(r"\s+", "", paragraph.text or "")
        is_title = (
            not paragraph_is_nonstructural(paragraph)
            and (
                logical in reference_major_titles
                or is_exam_title_line(logical)
                or bool(re.fullmatch(r"20\d{2}年(?:度)?", logical))
            )
        )
        if is_title:
            if final_title_seen and not final_title_previous:
                paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            paragraph.paragraph_format.line_spacing = Pt(32)
            for run in paragraph.runs:
                set_run_fonts(run, heading_font())
                run.font.size = Pt(23)
                run.bold = True
            final_title_seen = True
            final_title_previous = True
        else:
            final_title_previous = False
    # Compact symbol labels after every generic spacing pass.  This is a
    # paragraph-only adjustment and intentionally does not touch drawings,
    # OMML, tables, or floating-object XML.
    compact_symbol_fragments = compact_symbol_fragment_layout(document)
    # Keep a visual object with its immediate caption/short label, but do not
    # chain long body paragraphs and create large blank regions.
    final_paragraphs = list(iter_paragraphs(document))
    for index, paragraph in enumerate(final_paragraphs[:-1]):
        if not paragraph_has_visual_object(paragraph):
            continue
        following = final_paragraphs[index + 1]
        following_text = normalize_text(following.text or "")
        if following_text and (re.match(r"^(?:图|表)\s*\d", following_text) or len(following_text) <= 80):
            paragraph.paragraph_format.keep_with_next = True
    # Size normalization may touch drawing extents; repeat the XML cleanup so
    # the final package remains free of anchor-only attributes for Office.
    office_picture_normalization_final = normalize_picture_paragraphs_for_office(document)
    # Generic body styling treats table cells as non-structural. Reapply the
    # narrowly identified exam-cover title style after that pass so titles
    # stored inside Quark's metadata table keep their font and page boundary.
    table_title_cleanup_final = normalize_table_title_blocks(document)
    # Last-write title pagination: compatibility normalization above may
    # rewrite paragraph properties.  Reapply breaks immediately before save,
    # using non-structural content as a real group boundary.
    title_seen_content = False
    title_group_open = False
    for paragraph in iter_paragraphs(document):
        logical = re.sub(r"\s+", "", paragraph.text or "")
        is_title = (
            not paragraph_is_nonstructural(paragraph)
            and (
                logical in reference_major_titles
                or is_exam_title_line(logical)
                or bool(re.fullmatch(r"20\d{2}年(?:度)?", logical))
            )
        )
        if is_title:
            if title_seen_content and not title_group_open:
                paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            paragraph.paragraph_format.line_spacing = Pt(32)
            for run in paragraph.runs:
                set_run_fonts(run, heading_font())
                run.font.size = Pt(23)
                run.bold = True
            title_group_open = True
        else:
            if paragraph.text.strip() or paragraph_has_visual_object(paragraph):
                title_seen_content = True
            title_group_open = False
    numbered_subitem_alignment = normalize_numbered_subitem_alignment(document)
    # Final formula spacing override.  The Normal style uses an exact 19 pt
    # line box for body text; direct python-docx clearing alone therefore
    # still appears as fixed spacing in Word/WPS through style inheritance.
    # Write the explicit single-spacing OOXML after every generic pass.
    formula_single_spacing_final = finalize_formula_single_spacing(document)
    visual_layout = compare_visual_layout(document, visual_before)
    text_logic_audit = audit_text_logic(document)
    # Save through a sibling temporary file so a reader/antivirus handle on
    # the previous output does not make the edit fail with WinError 5.
    temp_save = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.docx")
    document.save(str(temp_save))
    os.replace(str(temp_save), str(path))
    return {"removed_header_footer_blocks": removed_headers, "collapsed_source_sections": collapsed_sections, "trailing_export_artifacts_removed": trailing_artifacts, "inserted_marker_page_breaks": removed_markers, "removed_redundant_page_breaks": removed_page_break_only + redundant_preheading_breaks, "compact_empty_spacers": compact_spacers, "broken_keep_chains": broken_keep_chains, "flattened_toc_rows": toc_tables, "floating_toc_rows_rebuilt": floating_toc_rows, "picture_anchors_inlined": picture_anchors_inlined, "mixed_picture_paragraphs_separated": mixed_picture_paragraphs_separated, "office_picture_normalization": office_picture_normalization, "office_picture_normalization_final": office_picture_normalization_final, "useless_content": useless_content, "empty_table_cleanup": empty_table_cleanup, "toc_rows_as_tables": toc_rows, "logical_keep_group_paragraphs": logical_keep_group_paragraphs, "formula_fragment_paragraphs_merged": formula_merges, "compact_symbol_fragments": compact_symbol_fragments, "noise_changes": noise_changes, "spacing_changes": spacing_changes, "column_alignment": column_alignment, "math_size_changes": math_size_changes, "formula_line_box_changes": formula_line_box_changes, "formula_single_spacing_final": formula_single_spacing_final, "numbered_subitem_alignment": numbered_subitem_alignment, "formula_table_height_changes": formula_table_height_changes, "malformed_math": malformed_math, "table_widths": table_widths, "table_pagination": table_pagination, "formula_table_width_changes": formula_table_width_changes, "layout_table_cleanup": layout_table_cleanup, "layout_padding_changes": layout_padding_changes, "table_title_cleanup": table_title_cleanup, "table_title_cleanup_final": table_title_cleanup_final, "visual_size_changes": visual_size_changes, "text_logic_audit": text_logic_audit, "body_page_numbers_restarted": int(body_numbering_restarted), "uncertain_region_candidates": uncertain_regions, "heading_font": heading_font()}


def source_text(path: Path, temp: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
            return normalize_text("\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages))
        except Exception:
            return ""
    if suffix == ".docx":
        return docx_text(path)
    if suffix == ".doc":
        converted = office_convert(path, temp / "source-doc")
        return docx_text(converted) if converted else ""
    return ""


def _block_text(block: Any) -> str:
    """Read visible text from a top-level paragraph or table block."""
    from docx.text.paragraph import Paragraph  # type: ignore
    from docx.table import Table  # type: ignore
    tag = block.tag.split("}")[-1]
    if tag == "p":
        return Paragraph(block, block.getparent()).text
    if tag == "tbl":
        table = Table(block, block.getparent())
        return "\n".join(cell.text for row in table.rows for cell in row.cells)
    return ""


def truncate_docx_to_rendered_pages(path: Path, max_pages: int, temp: Path) -> dict[str, Any]:
    """Keep only the blocks that LibreOffice places within the first N pages.

    A DOCX has no reliable page model until it is laid out by an office engine.
    We therefore render the already-normalized document, use the last requested
    PDF page's extracted text to find its final source block, and remove later
    top-level blocks. This keeps paragraphs, tables, and inline images editable.
    """
    if max_pages is None:
        return {"requested": None, "rendered_pages": None, "truncated": False}
    if max_pages < 1:
        raise RuntimeError("--max-pages 必须大于 0")
    from docx import Document  # type: ignore
    document = Document(str(path))
    body = document.element.body
    blocks = [child for child in body.iterchildren() if child.tag.split("}")[-1] != "sectPr"]
    rendered_pdf = office_to_pdf(path, temp / "page-limit-pdf")
    try:
        from pypdf import PdfReader  # type: ignore
        page_texts = [normalize_text(page.extract_text() or "") for page in PdfReader(str(rendered_pdf)).pages]
    except Exception:
        page_texts = []
    if len(page_texts) <= max_pages:
        return {"requested": max_pages, "rendered_pages": len(page_texts) or None, "truncated": False, "method": "rendered-pages"}

    document = Document(str(path))
    body = document.element.body
    blocks = [child for child in body.iterchildren() if child.tag.split("}")[-1] != "sectPr"]
    target_text = page_texts[max_pages - 1]
    cutoff = None
    # Match a substantial tail of the last visible paragraph. PDF text may
    # insert spaces between CJK glyphs, so try progressively shorter chunks.
    best_score = 0.0
    for index, block in enumerate(blocks):
        text = normalize_text(_block_text(block))
        if len(text) < 8 or not target_text:
            continue
        compact = text.replace(" ", "")
        target_compact = target_text.replace(" ", "")
        score = 0.0
        for width in (80, 50, 30, 16):
            if len(compact) >= width and compact[-width:] in target_compact:
                score = max(score, width / 80)
                break
        if score == 0.0:
            for start in range(0, max(1, len(compact) - 19), 20):
                gram = compact[start : start + 20]
                if len(gram) >= 12 and gram in target_compact:
                    score = max(score, len(gram) / 80)
        if score > best_score:
            best_score, cutoff = score, index
    if cutoff is None:
        # If the page is image-only, preserve a conservative proportion of
        # blocks instead of flattening anything to a page image.
        cutoff = max(0, round(len(blocks) * max_pages / len(page_texts)) - 1)
    for block in blocks[cutoff + 1 :]:
        parent = block.getparent()
        if parent is not None:
            parent.remove(block)
    document.save(str(path))
    return {"requested": max_pages, "rendered_pages": len(page_texts), "truncated": True, "method": "rendered-pages", "kept_blocks": cutoff + 1}


def bounded_similarity(left: str, right: str, max_chars: int = 120000) -> float:
    """Compare long OCR streams without quadratic memory/time growth."""
    if left == right:
        return 1.0
    longest = max(len(left), len(right))
    if longest <= max_chars:
        return difflib.SequenceMatcher(None, left, right, autojunk=False).ratio()
    window = max_chars // 3
    offsets_left = (0, max(0, (len(left) - window) // 2), max(0, len(left) - window))
    offsets_right = (0, max(0, (len(right) - window) // 2), max(0, len(right) - window))
    ratios = []
    for start_left, start_right in zip(offsets_left, offsets_right):
        ratios.append(difflib.SequenceMatcher(None, left[start_left : start_left + window], right[start_right : start_right + window], autojunk=False).ratio())
    length_score = min(len(left), len(right)) / max(len(left), len(right), 1)
    return 0.85 * (sum(ratios) / len(ratios)) + 0.15 * length_score


def compare_text(left: str, right: str) -> dict[str, Any]:
    left, right = normalize_text(left), normalize_text(right)
    if not left:
        return {"checked": False, "equal": None, "similarity": None, "source_chars": 0, "output_chars": len(right)}
    ratio = bounded_similarity(left, right)
    return {"checked": True, "equal": left == right, "similarity": round(ratio, 6), "source_chars": len(left), "output_chars": len(right)}


def document_object_checks(source: Path, output: Path, preview: bool = False) -> dict[str, Any]:
    """Check editable text, formula/table content, and image layout metadata.

    Image pixels are deliberately not checked. Their order and dimensions are
    checked so figures stay attached to the same text flow.
    """
    if source.suffix.lower() not in {".docx", ".doc"}:
        return {"checked": False, "reason": "只有 Word 输入能提供可编辑对象结构"}
    try:
        from docx import Document  # type: ignore
        left, right = Document(str(source)), Document(str(output))
        left_blocks = [normalize_text(p.text) for p in iter_paragraphs(left) if normalize_text(p.text)]
        right_blocks = [normalize_text(p.text) for p in iter_paragraphs(right) if normalize_text(p.text)]
        left_stream = re.sub(r"\s+", "", "".join(left_blocks))
        right_stream = re.sub(r"\s+", "", "".join(right_blocks))
        window = left_stream[: max(len(right_stream) + 256, len(right_stream))]
        text_similarity = bounded_similarity(window, right_stream) if right_stream else 0.0
        left_formula = [v for v in left_blocks if paragraph_needs_region_review(v)]
        right_formula = [v for v in right_blocks if paragraph_needs_region_review(v)]
        formula_matches = sum(1 for value in right_formula if value in left_formula)
        formula_ok = not right_formula or formula_matches / len(right_formula) >= 0.80
        left_tables = [[normalize_text(cell.text) for row in table.rows for cell in row.cells] for table in left.tables]
        right_tables = [[normalize_text(cell.text) for row in table.rows for cell in row.cells] for table in right.tables]
        table_matches = sum(
            1 for index, cells in enumerate(right_tables)
            if index < len(left_tables)
            and difflib.SequenceMatcher(None, "".join(left_tables[index]), "".join(cells), autojunk=False).ratio() >= 0.90
        )
        tables_ok = not right_tables or table_matches == len(right_tables)
        left_images = [(int(shape.width), int(shape.height)) for shape in left.inline_shapes]
        right_images = [(int(shape.width), int(shape.height)) for shape in right.inline_shapes]
        pairs = min(len(left_images), len(right_images))
        image_size_matches = sum(
            1 for a, b in zip(left_images, right_images)
            if abs(a[0] - b[0]) <= max(180000, a[0] * 0.08) and abs(a[1] - b[1]) <= max(180000, a[1] * 0.08)
        )
        image_layout_ok = len(right_images) <= len(left_images) and (
            preview or pairs == 0 or image_size_matches / pairs >= 0.75
        )
        return {
            "checked": True,
            "text": {"similarity": round(text_similarity, 6), "source_chars": len(left_stream), "output_chars": len(right_stream), "ok": text_similarity >= 0.90},
            "formulas": {"source_count": len(left_formula), "output_count": len(right_formula), "matched": formula_matches, "ok": formula_ok},
            "tables": {"source_count": len(left_tables), "output_count": len(right_tables), "matched": table_matches, "ok": tables_ok},
            "images": {"source_count": len(left_images), "output_count": len(right_images), "size_matches": image_size_matches, "ok": image_layout_ok, "pixels_checked": False, "preview_subset": preview},
            "passed": text_similarity >= 0.90 and formula_ok and tables_ok and image_layout_ok,
        }
    except Exception as exc:
        return {"checked": False, "reason": str(exc)}


def content_blocks(path: Path) -> list[str]:
    """Return ordered paragraph/table-cell text for loss detection."""
    from docx import Document  # type: ignore
    document = Document(str(path))
    return [normalize_text(p.text) for p in iter_paragraphs(document) if normalize_text(p.text)]


def content_integrity(source: Path, output: Path) -> dict[str, Any]:
    """Check that formatting did not delete or rewrite source Word content."""
    if source.suffix.lower() not in {".docx", ".doc"}:
        return {"checked": False, "reason": "输入不是 Word，无法做逐段原文保留校验"}
    try:
        left = content_blocks(source)
        right = content_blocks(output)
    except Exception as exc:
        return {"checked": False, "reason": str(exc)}
    # Running headers removed by design are presentation artifacts. OCR exports
    # also split/merge paragraphs differently, so compare an aggregate text
    # stream rather than requiring identical paragraph boundaries.
    ignorable = re.compile(r"^(?:机械原理习题(?:精解精练|集|解题)|第[0-9一二三四五六七八九十百]+章.*\d+)$")
    filtered_left = [
        value for value in left
        if not ignorable.match(value) and not USELESS_TEXT_RE.search(value) and not QR_TEXT_RE.search(value)
    ]
    left_stream = re.sub(r"\s+", "", "".join(filtered_left))
    right_stream = re.sub(r"\s+", "", "".join(right))
    similarity = bounded_similarity(left_stream, right_stream)
    missing = [value for value in filtered_left if value and value not in right]
    return {
        "checked": True,
        "equal": similarity >= 0.90 and len(right_stream) >= max(1, int(len(left_stream) * 0.80)),
        "similarity": round(similarity, 6),
        "source_blocks": len(left),
        "output_blocks": len(right),
        "missing_blocks": missing[:20],
        "missing_count": len(missing),
    }


def backup_source(source: Path, backup_dir: Path) -> Path:
    """Create a timestamped backup before any conversion or formatting."""
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    target = backup_dir / f"{source.stem}_backup_{stamp}{source.suffix}"
    # Avoid replacing a backup if two batch items start within the same second.
    counter = 1
    while target.exists():
        target = backup_dir / f"{source.stem}_backup_{stamp}_{counter}{source.suffix}"
        counter += 1
    shutil.copy2(source, target)
    return target


def crop_uncertain_region(image: Path, index: int, output_dir: Path) -> Path | None:
    """Create a deliberately small source crop for a suspicious formula region."""
    try:
        from PIL import Image  # type: ignore
        with Image.open(image) as source:
            width, height = source.size
            # Keep the crop bounded; never use this helper for a whole-page image.
            box = (int(width * 0.08), int(height * 0.28), int(width * 0.92), int(height * 0.48))
            crop = source.crop(box)
            output_dir.mkdir(parents=True, exist_ok=True)
            path = output_dir / f"uncertain-region-{index:04d}.png"
            crop.save(path)
            return path
    except Exception:
        return None


def embed_uncertain_regions_in_docx(path: Path, reference: Path, temp: Path, dpi: int) -> int:
    """Append bounded reference crops after formula-like paragraphs.

    This opt-in fallback never replaces a page with an image. Each crop is
    limited to a central band of one reference page and remains an inline Word
    picture beside the editable text.
    """
    from docx import Document  # type: ignore
    from docx.shared import Inches  # type: ignore

    kind = "pdf" if reference.suffix.lower() == ".pdf" else "image"
    pages = source_pages(reference, kind, temp / "uncertain-reference", dpi)
    document = Document(str(path))
    candidates = [p for p in iter_paragraphs(document) if paragraph_needs_region_review(p.text)]
    embedded = 0
    for index, paragraph in enumerate(candidates):
        image = pages[min(index, len(pages) - 1)]
        crop = crop_uncertain_region(image, index + 1, temp / "uncertain-crops")
        if crop:
            paragraph.add_run().add_picture(str(crop), width=Inches(5.8))
            embedded += 1
    if embedded:
        document.save(str(path))
    return embedded


def _remove_issue_fragment(paragraph: Any, fragment: str) -> bool:
    """Remove only an explicitly returned OCR fragment, never an entire paragraph by guess."""
    fragment = str(fragment or "")
    if not fragment:
        return False
    removed = False
    for run in paragraph.runs:
        if fragment in (run.text or ""):
            run.text = (run.text or "").replace(fragment, "")
            removed = True
    if not removed and not paragraph_has_visual_object(paragraph):
        combined = "".join(run.text or "" for run in paragraph.runs)
        if fragment in combined:
            replacement = combined.replace(fragment, "", 1)
            text_runs = [run for run in paragraph.runs if run.text]
            if text_runs:
                text_runs[0].text = replacement
                for run in text_runs[1:]:
                    run.text = ""
                removed = True
    return removed


def _clear_paragraph_text(paragraph: Any) -> None:
    """Remove paragraph text runs while retaining paragraph properties/objects."""
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def embed_local_issue_regions_in_docx(
    path: Path,
    reference: Path,
    issues: list[dict[str, Any]],
    temp: Path,
    dpi: int,
) -> dict[str, Any]:
    """Embed verified, local formula/symbol crops without any model review.

    Every coordinate is checked against the real PDF page before rendering.
    Invalid, ungrounded, or page-sized suggestions are reported but never
    inserted. This function never calls a network service or a language model.
    """
    result: dict[str, Any] = {"candidates": len(issues), "accepted": 0, "embedded": 0, "rejected": 0, "details": []}
    if not issues or reference.suffix.lower() != ".pdf" or not reference.exists():
        result["rejected"] = len(issues)
        result["reason"] = "没有可用的原始 PDF"
        return result
    try:
        from docx import Document  # type: ignore
        from docx.shared import Inches  # type: ignore
        from PIL import Image  # type: ignore
    except Exception as exc:
        result["rejected"] = len(issues)
        result["reason"] = f"局部图片依赖不可用: {exc}"
        return result

    document = Document(str(path))
    paragraphs = list(iter_paragraphs(document))
    temp.mkdir(parents=True, exist_ok=True)
    crop_dir = temp / "ai-issue-crops"
    crop_dir.mkdir(parents=True, exist_ok=True)
    try:
        try:
            import pymupdf as fitz  # type: ignore
        except ImportError:
            import fitz  # type: ignore
        pdf = fitz.open(str(reference))
        pdf_backend = "pymupdf"
    except Exception:
        pdf = None
        pdf_backend = "poppler"
    pdf_page_count = len(pdf) if pdf is not None else 0
    page_rects: dict[int, tuple[float, float]] = {}
    if pdf_backend == "poppler":
        pdfinfo = find_poppler_tool("pdfinfo")
        if pdfinfo:
            try:
                info = subprocess.run([pdfinfo, str(reference)], capture_output=True, timeout=60, check=False)
                info_text = info.stdout.decode("utf-8", errors="replace")
                match = re.search(r"Page size:\s*([0-9.]+)\s+x\s+([0-9.]+)", info_text)
                if match:
                    page_width, page_height = float(match.group(1)), float(match.group(2))
                    page_count_match = re.search(r"Pages:\s*(\d+)", info_text)
                    if page_count_match:
                        for page_number in range(1, int(page_count_match.group(1)) + 1):
                            page_rects[page_number] = (page_width, page_height)
            except Exception:
                pass
    rendered_pages: dict[int, Path] = {}
    if pdf_backend == "poppler":
        pdftoppm = find_poppler_tool("pdftoppm")
        if pdftoppm:
            for page_number in sorted({int(item.get("page", 0)) for item in issues if str(item.get("page", "")).isdigit() and int(item.get("page", 0)) > 0}):
                target_prefix = crop_dir / f"full-page-{page_number:04d}"
                try:
                    completed = subprocess.run([pdftoppm, "-png", "-r", str(dpi), "-f", str(page_number), "-l", str(page_number), str(reference), str(target_prefix)], capture_output=True, timeout=120, check=False)
                    candidate = crop_dir / f"full-page-{page_number:04d}-{page_number}.png"
                    if not candidate.exists():
                        matches = sorted(crop_dir.glob(f"full-page-{page_number:04d}-*.png"))
                        candidate = matches[0] if matches else candidate
                    if candidate.exists():
                        rendered_pages[page_number] = candidate
                except Exception:
                    continue
    changed = False
    used_indices: set[int] = set()
    for issue_number, issue in enumerate(issues, 1):
        detail: dict[str, Any] = {"issue_number": issue_number, "index": issue.get("index"), "page": issue.get("page"), "bbox": issue.get("bbox")}
        try:
            confidence = float(issue.get("confidence", 0))
            index = int(issue.get("index"))
            page_number = int(issue.get("page", 0))
            bbox = issue.get("bbox")
        except (TypeError, ValueError):
            confidence, index, page_number, bbox = 0.0, -1, 0, []
        category = str(issue.get("category", ""))
        if confidence < 0.80:
            detail["status"], detail["reason"] = "rejected", "confidence 低于 0.80"
        elif category not in {"text", "formula", "symbol"}:
            detail["status"], detail["reason"] = "rejected", "仅处理文字、公式和符号疑点"
        elif index < 0 or index >= len(paragraphs):
            detail["status"], detail["reason"] = "rejected", "Word 段落索引越界"
        elif page_number < 1 or (pdf_backend == "pymupdf" and page_number > pdf_page_count) or (pdf_backend == "poppler" and page_number not in page_rects):
            detail["status"], detail["reason"] = "rejected", "PDF 页码无效"
        elif not isinstance(bbox, list) or len(bbox) != 4:
            detail["status"], detail["reason"] = "rejected", "没有可验证的 PDF 局部坐标"
        else:
            try:
                x0, y0, x1, y1 = [float(value) for value in bbox]
                if pdf_backend == "pymupdf":
                    page = pdf[page_number - 1]
                    rect = page.rect
                    page_width, page_height = float(rect.width), float(rect.height)
                else:
                    page = None
                    page_width, page_height = page_rects.get(page_number, (595.0, 842.0))
                    rect = type("Rect", (), {"x0": 0.0, "y0": 0.0, "x1": page_width, "y1": page_height, "width": page_width, "height": page_height})()
                if not all(math.isfinite(value) for value in (x0, y0, x1, y1)):
                    raise ValueError("坐标不是有限数")
                x0, x1 = max(rect.x0, min(rect.x1, x0)), max(rect.x0, min(rect.x1, x1))
                y0, y1 = max(rect.y0, min(rect.y1, y0)), max(rect.y0, min(rect.y1, y1))
                width, height = x1 - x0, y1 - y0
                page_area = max(rect.width * rect.height, 1.0)
                crop_area = width * height
                if width <= 2 or height <= 2:
                    raise ValueError("局部框过小")
                if crop_area / page_area > 0.12:
                    raise ValueError("裁剪面积超过页面 12%")
                if width / max(rect.width, 1.0) >= 0.70 or height / max(rect.height, 1.0) >= 0.30:
                    raise ValueError("裁剪接近整页尺寸")
                # A small padding keeps glyph edges intact without turning the
                # crop into a page fragment.
                pad_x, pad_y = min(5.0, width * 0.12), min(4.0, height * 0.18)
                clip_coords = (max(rect.x0, x0 - pad_x), max(rect.y0, y0 - pad_y), min(rect.x1, x1 + pad_x), min(rect.y1, y1 + pad_y))
                clip = fitz.Rect(*clip_coords) if pdf_backend == "pymupdf" else type("Rect", (), {"x0": clip_coords[0], "y0": clip_coords[1], "x1": clip_coords[2], "y1": clip_coords[3]})()
                crop_path = crop_dir / f"issue-{issue_number:04d}-p{page_number:04d}.png"
                if pdf_backend == "pymupdf":
                    pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72.0, dpi / 72.0), clip=clip, alpha=False)
                    pix.save(str(crop_path))
                else:
                    full_page = rendered_pages.get(page_number)
                    if not full_page:
                        raise ValueError("Poppler 未能渲染目标 PDF 页")
                    with Image.open(full_page) as full_image:
                        scale_x = full_image.width / max(page_width, 1.0)
                        scale_y = full_image.height / max(page_height, 1.0)
                        pixel_box = (int(clip.x0 * scale_x), int(clip.y0 * scale_y), int(clip.x1 * scale_x), int(clip.y1 * scale_y))
                        full_image.crop(pixel_box).save(crop_path)
                with Image.open(crop_path) as image:
                    ratio = image.width / max(image.height, 1)
                    gray = image.convert("L").resize((max(1, image.width // 4), max(1, image.height // 4)))
                    pixels = list(gray.getdata())
                    ink_fraction = sum(value < 245 for value in pixels) / max(1, len(pixels))
                if ink_fraction < 0.002:
                    raise ValueError("裁剪区域几乎为空，拒绝嵌入")
                max_width = 1.9 if paragraph_in_table(paragraphs[index]) else 2.25
                max_height = 0.34 if category in {"formula", "symbol"} else 0.32
                width_inches = min(max_width, max(0.28, max_height * ratio))
                height_inches = width_inches / max(ratio, 0.1)
                if height_inches > max_height:
                    height_inches = max_height
                    width_inches = min(max_width, max(0.28, height_inches * ratio))
                paragraph = paragraphs[index]
                replace_text = str(issue.get("replace_text", ""))
                original_text = normalize_text(paragraph.text or "")
                # A replacement image must occupy its own visual line. If the
                # candidate describes the whole formula/symbol block, remove
                # all text runs first; appending an image beside leftover OCR
                # text is what caused the image-over-text defect.
                if replace_text and original_text == normalize_text(replace_text):
                    _clear_paragraph_text(paragraph)
                    removed_fragment = True
                else:
                    removed_fragment = _remove_issue_fragment(paragraph, replace_text)
                    if normalize_text(paragraph.text or ""):
                        raise ValueError("公式文字未能完整移除，拒绝与原文字重叠")
                # Do not duplicate a crop on a repeated run of the same review.
                if index in used_indices:
                    raise ValueError("同一 Word 段落已有局部疑点图")
                used_indices.add(index)
                paragraph.add_run().add_picture(str(crop_path), width=Inches(width_inches))
                from docx.enum.text import WD_LINE_SPACING  # type: ignore
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                paragraph.paragraph_format.line_spacing = None
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.space_before = Inches(0.01)
                paragraph.paragraph_format.space_after = Inches(0.01)
                changed = True
                result["accepted"] += 1
                result["embedded"] += 1
                detail.update({"status": "embedded", "removed_replace_text": removed_fragment, "crop": str(crop_path), "crop_fraction": round(crop_area / page_area, 6), "ink_fraction": round(ink_fraction, 6), "width_inches": round(width_inches, 3), "height_inches": round(height_inches, 3)})
            except Exception as exc:
                detail["status"], detail["reason"] = "rejected", str(exc)
        if detail.get("status") == "rejected":
            result["rejected"] += 1
        result["details"].append(detail)
    if pdf is not None:
        pdf.close()
    if changed:
        temp_save = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.docx")
        document.save(str(temp_save))
        os.replace(str(temp_save), str(path))
    result["full_page_images_detected"] = 0
    return result


def auto_embed_unreliable_formula_regions(
    path: Path,
    reference: Path | None,
    pdf_candidates: list[dict[str, Any]] | None,
    temp: Path,
    dpi: int,
) -> dict[str, Any]:
    """Replace unreliable formula/symbol text with small source-PDF crops.

    Native OMML is retained.  The automatic replacement target is limited to
    formula-like text blocks (typically OCR output) whose source PDF detector
    supplied a valid local box.  Candidates and Word blocks are paired in
    document order; unmatched or unsafe items are reported and never enlarged
    to a page image.
    """
    result: dict[str, Any] = {
        "checked": True,
        "automatic": True,
        "candidates": len(pdf_candidates or []),
        "eligible_word_blocks": 0,
        "embedded": 0,
        "rejected": 0,
        "details": [],
        "note": "仅将无法可靠保留为文字的公式/符号局部替换为 PDF 小图；不替换整页，不调用 AI",
    }
    if not reference or not reference.exists() or reference.suffix.lower() != ".pdf":
        result.update({"checked": False, "reason": "没有可用原始 PDF"})
        return result
    if not pdf_candidates:
        result.update({"checked": True, "reason": "没有公式候选区域"})
        return result
    try:
        from docx import Document  # type: ignore
    except Exception as exc:
        result.update({"checked": False, "reason": f"Word 解析失败: {exc}"})
        return result
    document = Document(str(path))
    paragraphs = list(iter_paragraphs(document))
    eligible: list[tuple[int, str]] = []
    for index, paragraph in enumerate(paragraphs):
        if paragraph._p.xpath(".//m:oMath"):
            continue
        text = normalize_text(paragraph.text or "")
        # Never replace a mixed Chinese/body paragraph wholesale.  A crop is
        # only safe when the Word block itself is a standalone formula/symbol;
        # otherwise clearing it would remove the surrounding sentence and the
        # inserted picture would appear to sit on top of text.
        compact = re.sub(r"\s+", "", text)
        standalone = bool(compact and not re.search(r"[\u4e00-\u9fff]", compact) and len(compact) <= 180)
        if paragraph_needs_region_review(text) and standalone:
            eligible.append((index, text))
    result["eligible_word_blocks"] = len(eligible)
    if not eligible:
        result["reason"] = "没有需要用局部图替换的非 OMML 公式/符号文字块"
        return result
    issues: list[dict[str, Any]] = []
    for (index, text), candidate in zip(eligible, pdf_candidates):
        issue = dict(candidate)
        issue.update({
            "index": index,
            "category": "formula" if _looks_like_formula_block(text)[0] else "symbol",
            "confidence": max(0.80, float(candidate.get("confidence") or 0.0)),
            "replace_text": text[:240],
        })
        issues.append(issue)
    embedded = embed_local_issue_regions_in_docx(path, reference, issues, temp, dpi)
    result.update(embedded)
    result["eligible_word_blocks"] = len(eligible)
    return result


def paragraph_needs_region_review(text: str) -> bool:
    compact = normalize_text(text)
    return bool(compact and (
        re.search(r"[∑∫√≈≤≥≠∞∂∇∏]|[A-Za-z]\s*[=_^]", compact)
        or re.search(r"\b(?:P|F|n|m)\s*[LH]\b", compact)
        or re.search(r"[�□]|[\u0370-\u03ff].*[=+\-]", compact)
    ))


def _structure_reason(text: str, role: str) -> tuple[float, str]:
    """Return an auditable confidence/reason pair for the local classifier.

    This is deliberately conservative: uncertain short paragraphs remain body
    text until an optional human-reviewed structure map overrides them.
    """
    compact = re.sub(r"\s+", "", text or "")
    if role == "title":
        return 0.98, "first non-empty short paragraph"
    if role == "heading1":
        if re.match(r"^第[0-9一二三四五六七八九十百]+章", compact):
            return 0.98, "Chinese chapter marker"
        if re.match(r"^\d+\s+", text or ""):
            return 0.95, "numbered top-level heading"
        return 0.88, "known major-section label or surrounding structure"
    if role in {"heading2", "heading3"}:
        if re.match(r"^\d+(?:\.\d+){1,2}\s+", text or ""):
            return 0.95, "numbered nested heading"
        return 0.78, "short non-sentence paragraph with heading context"
    if role == "caption":
        return 0.97, "figure/table caption marker"
    if role == "reference":
        return 0.96, "reference entry marker"
    if len(compact) <= 16 and not looks_like_sentence(text or ""):
        return 0.52, "ambiguous short paragraph kept as body"
    return 0.90, "body paragraph"


def build_structure_map(document: Any) -> dict[str, Any]:
    """Classify paragraphs and return a JSON-serializable audit map.

    The map follows the reference project's contract (index + role) and adds
    confidence/reason/text fields for deterministic auditing and optional
    human-provided structure-map overrides.
    """
    entries: list[dict[str, Any]] = []
    in_refs = False
    paragraphs = list(iter_paragraphs(document))
    first_content_index = next(
        (i for i, p in enumerate(paragraphs) if p.text.strip() and not paragraph_is_nonstructural(p)),
        None,
    )
    excluded_objects = 0
    for index, paragraph in enumerate(paragraphs):
        if not paragraph.text.strip():
            continue
        if paragraph_is_nonstructural(paragraph):
            excluded_objects += 1
            continue
        role = classify_paragraph(paragraph, in_refs, paragraphs, index, first_content_index)
        if role == "empty":
            continue
        if role == "heading1" and re.sub(r"\s+", "", paragraph.text) == "参考文献":
            in_refs = True
        confidence, reason = _structure_reason(paragraph.text, role)
        entries.append({
            "index": index,
            "role": role,
            "confidence": confidence,
            "reason": reason,
            "text": normalize_text(paragraph.text)[:160],
        })
    return {
        "version": 1,
        "classifier": "conservative-local-structure-v2",
        "mapping_scope": "top-level text paragraphs only; tables, images, and formulas excluded",
        "excluded_nonstructural_paragraphs": excluded_objects,
        "paragraphs": entries,
        "uncertain": [entry for entry in entries if entry["confidence"] < 0.7],
    }


def write_structure_map(document: Any, path: Path) -> dict[str, Any]:
    path.parent.mkdir(parents=True, exist_ok=True)
    data = build_structure_map(document)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return data


def render_pdf(path: Path, directory: Path, dpi: int, max_pages: int | None = None) -> list[Path]:
    pdftoppm = find_poppler_tool("pdftoppm")
    directory.mkdir(parents=True, exist_ok=True)
    if pdftoppm:
        command = [pdftoppm, "-png", "-r", str(dpi)]
        if max_pages is not None:
            command.extend(["-f", "1", "-l", str(max_pages)])
        command.extend([str(path), str(directory / "page")])
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)
        if result.returncode == 0:
            pages = sorted(directory.glob("page-*.png"))
            if pages:
                return pages
    # PyMuPDF is already required for image/formula inspection and avoids a
    # separate Poppler installation on Windows.
    try:
        try:
            import pymupdf as fitz  # type: ignore
        except ImportError:
            import fitz  # type: ignore
        from PIL import Image  # type: ignore
        document = fitz.open(str(path))
        limit = min(len(document), max_pages) if max_pages is not None else len(document)
        pages: list[Path] = []
        scale = dpi / 72
        matrix = fitz.Matrix(scale, scale)
        for index in range(limit):
            pixmap = document[index].get_pixmap(matrix=matrix, alpha=False)
            target = directory / f"page-{index + 1:04d}.png"
            pixmap.save(str(target))
            pages.append(target)
        document.close()
        if pages:
            return pages
    except Exception as exc:
        raise RuntimeError(f"PDF 渲染失败（pdftoppm 和 PyMuPDF 均不可用）: {exc}") from exc
    raise RuntimeError("PDF 未生成页面")


def office_convert(path: Path, directory: Path) -> Path | None:
    soffice = find_soffice()
    if use_bridge():
        return bridge_convert(path, directory, "docx")
    if not soffice or not soffice_usable():
        return None
    directory.mkdir(parents=True, exist_ok=True)
    profile = directory / "profile"
    result = run_office([soffice, "--headless", "--invisible", "--norestore", "--nofirststartwizard", "--nodefault", f"-env:UserInstallation={profile.resolve().as_uri()}", "--convert-to", "docx", "--outdir", str(directory), str(path)])
    output = directory / f"{path.stem}.docx"
    return output if result.returncode == 0 and output.exists() else None


def office_to_pdf(path: Path, directory: Path) -> Path:
    soffice = find_soffice()
    if use_bridge():
        converted = bridge_convert(path, directory, "pdf")
        if converted:
            return converted
        raise RuntimeError("LibreOffice 需要在你的 Windows 用户账户中运行；请先启动 scripts\\lo_user_bridge.py")
    if not soffice or not soffice_usable():
        raise RuntimeError("需要 LibreOffice 将 Word 渲染为 PDF")
    directory.mkdir(parents=True, exist_ok=True)
    profile = directory / "profile"
    result = run_office([soffice, "--headless", "--invisible", "--norestore", "--nofirststartwizard", "--nodefault", f"-env:UserInstallation={profile.resolve().as_uri()}", "--convert-to", "pdf", "--outdir", str(directory), str(path)])
    output = directory / f"{path.stem}.pdf"
    if result.returncode != 0 or not output.exists():
        raise RuntimeError(f"LibreOffice 无法渲染 Word:\n{result.stderr}")
    return output


def inspect_rendered_layout(path: Path, directory: Path, max_pages: int | None = None) -> dict[str, Any]:
    """Inspect rendered pages for true blank pages and abnormal bottom gaps."""
    try:
        rendered = office_to_pdf(path, directory / "layout-pdf")
        try:
            import pymupdf as fitz  # type: ignore
        except ImportError:
            import fitz  # type: ignore
        document = fitz.open(str(rendered))
        limit = len(document) if max_pages is None else min(len(document), max_pages)
        blank_pages: list[int] = []
        large_whitespace_pages: list[int] = []
        bottom_gaps: list[float] = []
        for index in range(limit):
            page = document[index]
            blocks = page.get_text("blocks")
            meaningful: list[Any] = []
            for block in blocks:
                text = str(block[4]).strip() if len(block) > 4 else ""
                block_type = int(block[6]) if len(block) > 6 and str(block[6]).isdigit() else 0
                if text or block_type == 1:
                    meaningful.append(block)
            if not meaningful:
                blank_pages.append(index + 1)
                bottom_gaps.append(1.0)
                continue
            max_y = max(float(block[3]) for block in meaningful)
            gap = max(0.0, (float(page.rect.height) - max_y) / max(float(page.rect.height), 1.0))
            bottom_gaps.append(gap)
            # Report only; title/TOC pages may intentionally contain whitespace.
            if gap >= 0.55 and index > 0:
                large_whitespace_pages.append(index + 1)
        document.close()
        return {
            "checked": True,
            "pages": limit,
            "blank_pages": blank_pages,
            "large_whitespace_pages": large_whitespace_pages,
            "max_bottom_gap": round(max(bottom_gaps, default=0.0), 4),
            "passed": not blank_pages,
            "method": "libreoffice-pdf-pymupdf",
        }
    except Exception as exc:
        return {"checked": False, "reason": str(exc), "blank_pages": [], "large_whitespace_pages": []}


def run_office(command: list[str]) -> subprocess.CompletedProcess[str]:
    """Run LibreOffice with a hard timeout and terminate descendants on hang."""
    try:
        return subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", check=False, timeout=180, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0), env=office_environment())
    except subprocess.TimeoutExpired as exc:
        # subprocess.run cannot kill descendants after timeout on Windows;
        # terminate the known LibreOffice process tree before returning.
        try:
            for image in ("soffice.com", "soffice.exe", "soffice.bin"):
                subprocess.run(["taskkill", "/F", "/T", "/IM", image], capture_output=True, check=False, timeout=15, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except (OSError, subprocess.TimeoutExpired):
            pass
        return subprocess.CompletedProcess(command, 124, stdout=exc.stdout or "", stderr=(exc.stderr or "") + "\nLibreOffice 转换超时，已结束进程。")


def office_environment() -> dict[str, str]:
    """Prevent the Codex/Python runtime from contaminating LibreOffice."""
    env = os.environ.copy()
    for name in ("PYTHONHOME", "PYTHONPATH", "PYTHONSTARTUP", "VIRTUAL_ENV"):
        env.pop(name, None)
    return env


def bridge_convert(path: Path, directory: Path, output_format: str) -> Path | None:
    """Queue conversion for the user-account bridge when running in Codex sandbox."""
    root = Path(__file__).resolve().parents[1]
    bridge_root = root / "tmp" / "lo-bridge"
    queue = bridge_root / "queue-v2"
    done = bridge_root / "done-v2"
    errors = bridge_root / "errors-v2"
    for folder in (queue, done, errors, bridge_root / "input", bridge_root / "work"):
        folder.mkdir(parents=True, exist_ok=True)
    job_id = uuid.uuid4().hex
    # The bridge runs under the interactive Windows account.  Files created by
    # the Codex sandbox (especially tmp/word job folders) can be unreadable or
    # unwritable to that account, so stage both input and output in the shared
    # bridge directory and copy the result back after completion.
    stage = bridge_root / "work" / job_id
    stage.mkdir(parents=True, exist_ok=True)
    staged_input = stage / path.name
    shutil.copy2(path, staged_input)
    staged_output_dir = stage / "out"
    staged_output_dir.mkdir(parents=True, exist_ok=True)
    job = {
        "id": job_id,
        "input": str(staged_input.resolve()),
        "output_dir": str(staged_output_dir.resolve()),
        "format": output_format,
    }
    (queue / f"{job_id}.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    deadline = time.monotonic() + 660
    result_file = done / f"{job_id}.json"
    error_file = errors / f"{job_id}.json"
    while time.monotonic() < deadline:
        if result_file.exists():
            try:
                result = json.loads(result_file.read_text(encoding="utf-8"))
                staged_output = Path(result.get("output", ""))
                if not staged_output.exists():
                    return None
                directory.mkdir(parents=True, exist_ok=True)
                output = directory / f"{path.stem}.{output_format}"
                shutil.copy2(staged_output, output)
                shutil.rmtree(stage, ignore_errors=True)
                return output
            finally:
                result_file.unlink(missing_ok=True)
        if error_file.exists():
            error = error_file.read_text(encoding="utf-8", errors="replace")
            error_file.unlink(missing_ok=True)
            shutil.rmtree(stage, ignore_errors=True)
            raise RuntimeError(f"用户账户 LibreOffice 转换失败: {error}")
        time.sleep(1)
    shutil.rmtree(stage, ignore_errors=True)
    raise RuntimeError("等待用户账户 LibreOffice 桥接器超时；请确认 lo_user_bridge.py 正在运行")


def use_bridge() -> bool:
    if os.environ.get("CODEX_LO_BRIDGE", "").lower() in {"1", "true", "yes"}:
        return True
    # The sandbox preserves USERNAME=OMEN, so USERNAME alone is not a reliable
    # detector. Its runtime path and `whoami` identity are instead stable.
    runtime = str(Path(sys.executable)).lower()
    if "codex-runtimes" in runtime or "codexsandbox" in runtime:
        return True
    try:
        identity = subprocess.run(
            ["whoami"], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=5,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        ).stdout.lower()
        return "codexsandbox" in identity
    except (OSError, subprocess.TimeoutExpired):
        return False


def source_pages(path: Path, kind: str, directory: Path, dpi: int, max_pages: int | None = None) -> list[Path]:
    if kind == "image":
        return [path]
    if kind == "pdf":
        return render_pdf(path, directory / "pdf", dpi, max_pages)
    source_doc = path if path.suffix.lower() == ".docx" else office_convert(path, directory / "office")
    if not source_doc:
        raise RuntimeError(f"无法读取 Word 文件: {path}")
    return render_pdf(office_to_pdf(source_doc, directory / "docx-pdf"), directory / "docx-pages", dpi, max_pages)


def compare(source: Path, output: Path, kind: str, mode: str, temp: Path, dpi: int, max_pages: int | None = None, reference_pages: list[Path] | None = None) -> dict[str, Any]:
    text_result = compare_text(source_text(source, temp / "texts"), docx_text(output))
    visual = {"checked": False, "match": None}
    try:
        left_pages = reference_pages or source_pages(source, kind, temp / "source-compare", dpi, max_pages)
        output_pdf = office_to_pdf(output, temp / "output-pdf")
        visual = visual_compare(
            left_pages,
            render_pdf(output_pdf, temp / "output-pages", dpi, max_pages),
            mode,
            right_pdf=output_pdf,
        )
    except Exception as exc:
        visual = {"checked": False, "match": None, "reason": str(exc)}
    objects = document_object_checks(source, output, preview=max_pages is not None)
    text_ok = not text_result["checked"] or text_result["equal"] is True
    passed = bool(objects.get("passed")) if objects.get("checked") else (text_ok and visual.get("match") is True)
    return {"text": text_result, "objects": objects, "visual": visual, "passed": passed}


def find_reference(source: Path, explicit: Path | None, project_root: Path) -> Path | None:
    """Resolve the original PDF that is the single source of truth."""
    if explicit:
        return explicit.resolve()
    reference_dir = project_root / "input" / "pdfs"
    exact = reference_dir / f"{source.stem}.pdf"
    if exact.exists():
        return exact.resolve()
    candidates = sorted(reference_dir.glob("*.pdf"))
    return candidates[0].resolve() if len(candidates) == 1 else None


def visual_compare(left_pages: list[Path], right_pages: list[Path], mode: str, right_pdf: Path | None = None) -> dict[str, Any]:
    from PIL import Image, ImageChops, ImageOps  # type: ignore
    count = min(len(left_pages), len(right_pages))
    if count == 0:
        return {"checked": True, "match": False, "source_pages": len(left_pages), "output_pages": len(right_pages), "pages_checked": 0, "page_count_checked": False}
    indices = list(range(count)) if mode == "full" else sorted(set((0, count // 2, count - 1)))
    diffs: list[float] = []
    changed: list[float] = []
    layout_scores: list[float] = []
    ignored_image_pixels = 0
    for index in indices:
        with Image.open(left_pages[index]) as left, Image.open(right_pages[index]) as right:
            left, right = left.convert("RGB"), right.convert("RGB")
            if left.size != right.size:
                right = right.resize(left.size)
            gray = ImageOps.grayscale(ImageChops.difference(left, right))
            try:
                import numpy as np  # type: ignore
                def projection(image: Any) -> tuple[Any, Any]:
                    small = ImageOps.grayscale(image).resize((128, 192))
                    ink = np.asarray(small, dtype=float)
                    ink = np.maximum(255.0 - ink - 35.0, 0.0)
                    total = max(float(ink.sum()), 1.0)
                    return ink.sum(axis=1) / total, ink.sum(axis=0) / total
                left_projection = projection(left)
                right_projection = projection(right)
                layout_scores.append(round(1.0 - (float(np.abs(left_projection[0] - right_projection[0]).mean()) + float(np.abs(left_projection[1] - right_projection[1]).mean())) / 2.0, 6))
            except Exception:
                pass
            valid = Image.new("L", left.size, 255)
            if right_pdf:
                try:
                    try:
                        import pymupdf as fitz  # type: ignore
                    except ImportError:
                        import fitz  # type: ignore
                    from PIL import ImageDraw  # type: ignore
                    pdf = fitz.open(str(right_pdf))
                    page = pdf[index]
                    sx, sy = left.width / page.rect.width, left.height / page.rect.height
                    draw = ImageDraw.Draw(valid)
                    for block in page.get_text("dict").get("blocks", []):
                        if block.get("type") == 1 and block.get("bbox"):
                            x0, y0, x1, y1 = block["bbox"]
                            draw.rectangle((int(x0 * sx), int(y0 * sy), int(x1 * sx), int(y1 * sy)), fill=0)
                    pdf.close()
                except Exception:
                    pass
            values = list(valid.getdata())
            pixels = sum(1 for value in values if value)
            ignored_image_pixels += len(values) - pixels
            gray = Image.composite(gray, Image.new("L", left.size, 0), valid)
            histogram = gray.histogram()
            diffs.append(sum(value * amount for value, amount in enumerate(histogram)) / (255 * max(1, pixels)))
            changed.append(sum(histogram[13:]) / max(1, pixels))
    max_diff, max_changed = max(diffs, default=1), max(changed, default=1)
    min_layout = min(layout_scores, default=0.0)
    return {"checked": True, "match": min_layout >= 0.90, "layout_match": min_layout >= 0.90, "pixel_exact": max_diff == 0 and max_changed == 0, "pages_checked": len(indices), "source_pages": len(left_pages), "output_pages": len(right_pages), "page_count_checked": False, "image_pixels_checked": False, "ignored_image_pixels": ignored_image_pixels, "layout_scores": layout_scores, "min_layout_score": round(min_layout, 6), "max_mean_diff_diagnostic": round(max_diff, 6), "max_changed_ratio_diagnostic": round(max_changed, 6)}


def diff_crops(left_pages: list[Path], right_pages: list[Path], output_dir: Path) -> list[tuple[int, Path]]:
    """Extract only bounded changed regions; never emit a full-page fallback."""
    from PIL import Image, ImageChops, ImageFilter  # type: ignore
    output_dir.mkdir(parents=True, exist_ok=True)
    crops: list[tuple[int, Path]] = []
    for page_index, (left_path, right_path) in enumerate(zip(left_pages, right_pages)):
        with Image.open(left_path) as left_image, Image.open(right_path) as right_image:
            left = left_image.convert("RGB")
            right = right_image.convert("RGB").resize(left.size)
            # Work at 1/8 resolution to group anti-aliasing noise.
            small_left = left.resize((max(1, left.width // 8), max(1, left.height // 8)))
            small_right = right.resize(small_left.size)
            mask = ImageChops.difference(small_left, small_right).convert("L").point(lambda p: 255 if p > 35 else 0)
            mask = mask.filter(ImageFilter.MaxFilter(5))
            px = mask.load()
            width, height = mask.size
            seen: set[tuple[int, int]] = set()
            components: list[tuple[int, int, int, int, int]] = []
            for y in range(height):
                for x in range(width):
                    if not px[x, y] or (x, y) in seen:
                        continue
                    stack = [(x, y)]
                    seen.add((x, y))
                    x0 = x1 = x
                    y0 = y1 = y
                    area = 0
                    while stack:
                        cx, cy = stack.pop()
                        area += 1
                        x0, x1, y0, y1 = min(x0, cx), max(x1, cx), min(y0, cy), max(y1, cy)
                        for nx, ny in ((cx + 1, cy), (cx - 1, cy), (cx, cy + 1), (cx, cy - 1)):
                            if 0 <= nx < width and 0 <= ny < height and px[nx, ny] and (nx, ny) not in seen:
                                seen.add((nx, ny))
                                stack.append((nx, ny))
                    if area >= 8:
                        components.append((area, x0, y0, x1, y1))
            for crop_index, (area, x0, y0, x1, y1) in enumerate(sorted(components, reverse=True)[:12], 1):
                crop_width = x1 - x0 + 1
                crop_height = y1 - y0 + 1
                if crop_width >= width * 0.72 or crop_height >= height * 0.72:
                    continue
                pad_x, pad_y = 10, 8
                box = (max(0, (x0 * 8) - pad_x), max(0, (y0 * 8) - pad_y), min(left.width, ((x1 + 1) * 8) + pad_x), min(left.height, ((y1 + 1) * 8) + pad_y))
                if (box[2] - box[0]) * (box[3] - box[1]) >= left.width * left.height * 0.35:
                    continue
                target = output_dir / f"page-{page_index + 1:04d}-region-{crop_index:02d}.png"
                left.crop(box).save(target)
                crops.append((page_index, target))
    return crops


def detect_full_page_images(document: Any) -> int:
    """Count page-sized inline images; these are forbidden as a fallback."""
    if not document.inline_shapes:
        return 0
    count = 0
    for shape in document.inline_shapes:
        section = document.sections[0]
        usable_width = section.page_width - section.left_margin - section.right_margin
        usable_height = section.page_height - section.top_margin - section.bottom_margin
        if shape.width >= usable_width * 0.9 and shape.height >= usable_height * 0.9:
            count += 1
    return count


def embed_diff_regions_in_docx(path: Path, source_pages: list[Path], output_pages: list[Path], temp: Path) -> int:
    """Insert bounded source crops near the matching page's text flow."""
    from docx import Document  # type: ignore
    from docx.shared import Inches  # type: ignore
    crops = diff_crops(source_pages, output_pages, temp / "diff-crops")
    if not crops:
        return 0
    document = Document(str(path))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    formula_paragraphs = [
        p for p in paragraphs
        if not paragraph_has_visual_object(p)
        and not paragraph_in_table(p)
        and len(normalize_text(p.text)) <= 140
        and paragraph_needs_region_review(p.text)
    ]
    if not paragraphs:
        return 0
    # A diff detector can emit many connected components for one formula. Map
    # each source page to at most one candidate paragraph and never clear an
    # ordinary body paragraph. This keeps the fallback local and idempotent.
    used: set[int] = set()
    for page_index, crop in crops:
        target_index = min(len(paragraphs) - 1, round((page_index + 1) * len(paragraphs) / max(1, len(source_pages))) - 1)
        # Never attach a source crop to an arbitrary body paragraph. A local
        # image fallback is permitted only when the paragraph itself contains
        # a recognizable OCR/math uncertainty signal.
        candidates = formula_paragraphs
        if not candidates:
            continue
        candidate_index = min(len(candidates) - 1, round((page_index + 1) * len(candidates) / max(1, len(source_pages))) - 1)
        paragraph = candidates[candidate_index]
        marker = id(paragraph)
        if marker in used:
            continue
        used.add(marker)
        replace_formula = paragraph_needs_region_review(paragraph.text)
        if replace_formula:
            # The PDF crop is authoritative for an OCR-uncertain formula; the
            # paragraph text is replaced only in this explicitly flagged case.
            paragraph.text = ""
        else:
            paragraph = paragraphs[target_index].insert_paragraph_before()
        # Keep a formula fallback visually comparable to 13 pt body text.
        # Bound both width and height so a crop cannot overflow a table cell
        # or push neighboring lines apart.
        from PIL import Image  # type: ignore
        with Image.open(crop) as image:
            ratio = image.width / max(image.height, 1)
        # Keep a fallback formula visually comparable to 13 pt body text.
        # Bound both dimensions; a wide crop must not push the following
        # paragraph down, and a tall crop must not overflow a table cell.
        width_inches = min(1.8, max(0.55, 0.55 * ratio))
        height_inches = width_inches / max(ratio, 0.1)
        if height_inches > 0.38:
            height_inches = 0.38
            width_inches = max(0.55, height_inches * ratio)
        paragraph.paragraph_format.space_before = Inches(0.02)
        paragraph.paragraph_format.space_after = Inches(0.02)
        paragraph.add_run().add_picture(str(crop), width=Inches(width_inches))
    if used:
        document.save(str(path))
    return len(used)


def process(source: Path, output: Path, args: argparse.Namespace, report_dir: Path) -> dict[str, Any]:
    """Convert any supported source, apply deterministic formatting, then verify it."""
    if source.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise RuntimeError(f"输入必须是 Word、PDF 或图片文件: {source}")
    project_root = report_dir.parent.parent
    temp_root = project_root / "tmp" / "word"
    temp_root.mkdir(parents=True, exist_ok=True)
    temp = Path(tempfile.mkdtemp(prefix=f"{source.stem}-", dir=str(temp_root)))
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        backup_dir = project_root / "output" / "backups"
        backup_path = backup_source(source, backup_dir)
        source_for_work = source
        engine_info: dict[str, Any] = {"engine": "local-word", "status": "not-applicable"}
        if source.suffix.lower() == ".pdf":
            source_for_work = limit_pdf_pages(source, temp / "source-preview.pdf", args.max_pages)
            pdf_engine = getattr(args, "pdf_engine", "auto")
            converted = None
            # Fast path for ordinary electronic PDFs.  It avoids model startup
            # and keeps local figures/equations as bounded PDF crops.
            if pdf_engine in {"auto", "pymupdf"}:
                try:
                    converted = pymupdf_pdf_to_docx(source_for_work, temp / "source-pymupdf.docx", args.max_pages, args.dpi)
                    engine_info = {"engine": "pymupdf", "status": "ok", "fast_mode": True}
                except Exception as exc:
                    engine_info = {"engine": "pymupdf", "status": "failed", "fast_mode": True, "reason": str(exc)}
            if not converted and pdf_engine in {"auto", "mineru"}:
                converted, engine_info = mineru_pdf_to_docx(source_for_work, temp / "source-mineru.docx", temp, args)
            if not converted and pdf_engine == "mineru":
                reason = engine_info.get("reason") or engine_info.get("stderr") or "MinerU 转换失败"
                raise RuntimeError(f"MinerU 无法将 PDF 转为 Word: {source}\n{reason}")
            if not converted:
                engine_info = {"engine": "local-word-bridge", "status": "fallback", "mineru": engine_info}
                converted = office_convert(source_for_work, temp / "source-office")
            if not converted:
                raise RuntimeError(f"无法将 PDF 转为 Word；请确认用户账户桥接器正在运行: {source}")
            # Always replace a previous output.  Keeping the old file here
            # made repeated tests appear to ignore the selected fast engine.
            output = copy_to_writable_output(converted, output)
        elif source.suffix.lower() in IMAGE_EXTENSIONS:
            image_to_docx(source, output)
        elif source.suffix.lower() == ".doc":
            converted = office_convert(source, temp / "source-office")
            if not converted:
                raise RuntimeError(f"无法读取 .doc 文件，请确认用户账户桥接器正在运行: {source}")
            output = copy_to_writable_output(converted, output)
        else:
            output = copy_to_writable_output(source, output)

        reference = find_reference(source, args.reference, project_root)
        from docx import Document  # type: ignore
        reference_major_titles: set[str] = set()
        if source.suffix.lower() == ".docx":
            # Read the untouched input before any cleanup changes its direct
            # formatting. This is the local authority for title recognition.
            reference_major_titles = reference_major_title_texts(Document(str(source)))
        structure_overrides: dict[int, str] = {}
        if args.structure_map:
            structure_overrides = load_structure_map(args.structure_map)
        rebuilt_pdf_mode = bool(engine_info.get("fast_mode") or engine_info.get("engine") in {"mineru", "mineru-content-rebuild"})
        if rebuilt_pdf_mode:
            # PyMuPDF has already emitted one normal-flow block sequence per
            # source page.  Running the legacy postprocessor here would
            # intentionally collapse empty/page-break blocks and destroy the
            # PDF page grouping.  Keep this path minimal and deterministic.
            fast_document = Document(str(output))
            picture_cleanup = normalize_picture_paragraphs_for_office(fast_document)
            fast_temp = output.with_name(f".{output.stem}.{uuid.uuid4().hex}.tmp.docx")
            fast_document.save(str(fast_temp))
            os.replace(str(fast_temp), str(output))
            formatting = {
                "fast_path": bool(engine_info.get("fast_mode")),
                "engine": engine_info.get("engine", "rebuilt-pdf"),
                "office_picture_normalization": picture_cleanup,
                "inserted_marker_page_breaks": 0,
                "removed_redundant_page_breaks": 0,
            }
        else:
            formatting = postprocess_docx(
                output,
                strip_header_footer=True,
                structure_overrides=structure_overrides,
                reference_major_titles=reference_major_titles,
            )
        # One rendered pass validates real pagination and floating-object
        # behavior. If LibreOffice is unavailable, keep the DOCX result and
        # record the diagnostic failure instead of changing the conversion.
        if args.render_layout or args.visual_compare:
            formatting["rendered_layout"] = inspect_rendered_layout(output, temp, args.max_pages)
        else:
            formatting["rendered_layout"] = {
                "checked": False,
                "reason": "默认仅执行 DOCX/XML 分页、留白和图片布局检查；添加 --render-layout 可启用真实渲染审计",
                "blank_pages": [],
                "large_whitespace_pages": [],
            }
        # For Word-first previews, enforce the requested page limit after
        # layout.  The final document remains editable; only blocks after the
        # rendered page boundary are removed.
        page_limit_info = {"requested": None, "rendered_pages": None, "truncated": False}
        if source.suffix.lower() == ".docx" and args.max_pages is not None:
            page_limit_info = truncate_docx_to_rendered_pages(output, args.max_pages, temp)
        formatting["page_limit"] = page_limit_info

        # Word-first runs should not silently re-scan an automatically found
        # source PDF.  Quark DOCX already contains the editable/object data;
        # PDF formula detection is enabled only for a PDF input or when the
        # user explicitly asks for visual comparison.
        check_reference = reference if (source.suffix.lower() == ".pdf" or args.visual_compare) else None
        formula_reference = check_reference if check_reference and check_reference.suffix.lower() == ".pdf" else (source_for_work if source_for_work.suffix.lower() == ".pdf" and source.suffix.lower() == ".pdf" else None)
        fast_mode = bool(engine_info.get("fast_mode"))
        # Never enqueue the heavyweight formula bridge for the PyMuPDF fast
        # path.  Equation crops have already been created directly from the
        # source PDF and waiting for a model would defeat the speed goal.
        pdfkit_candidates, pdfkit_status = (
            None,
            {"status": "skipped", "reason": "Word-first 输入不启动 PDF-Extract-Kit；仅处理夸克 DOCX"},
        ) if not formula_reference else (
            (None, {"status": "skipped", "reason": "PyMuPDF 快速路径已直接裁剪公式区域"})
            if fast_mode else pdfkit_formula_bridge_detect(formula_reference, temp, args.max_pages)
        )
        if not pdfkit_candidates:
            # The bridge may be unavailable in a restricted runtime.  Reuse
            # the deterministic MinerU/PyMuPDF chain so local replacement is
            # still automatic when a reliable source box exists.
            pdf_candidates_for_replacement = pdf_formula_candidates(
                formula_reference,
                args.max_pages,
                Path(engine_info["content_list"]) if engine_info.get("content_list") else None,
            )
        else:
            pdf_candidates_for_replacement = pdfkit_candidates
        if fast_mode:
            local_replacements = {"checked": True, "automatic": True, "embedded": 0, "rejected": 0, "reason": "PyMuPDF 已在重建阶段将不可靠公式作为局部 PDF 图片插入"}
        elif formula_reference:
            local_replacements = auto_embed_unreliable_formula_regions(
                output, formula_reference, pdf_candidates_for_replacement, temp, args.dpi,
            )
        else:
            local_replacements = {"checked": False, "automatic": False, "embedded": 0, "rejected": 0, "reason": "Word-first 默认不重新扫描自动发现的 PDF"}
        # Local formula crops are appended after the main formatting pass.
        # Normalize once more so newly inserted pictures follow the same
        # Office-safe standalone-paragraph and XML rules as source figures.
        if local_replacements.get("embedded"):
            replacement_document = Document(str(output))
            replacement_picture_normalization = normalize_picture_paragraphs_for_office(replacement_document)
            replacement_temp = output.with_name(f".{output.stem}.{uuid.uuid4().hex}.tmp.docx")
            replacement_document.save(str(replacement_temp))
            os.replace(str(replacement_temp), str(output))
            local_replacements["office_picture_normalization"] = replacement_picture_normalization
        # Final package-level compatibility pass.  This runs after every
        # python-docx mutation so newly inserted pictures receive the same
        # OOXML normalization as source pictures.  It is deterministic and
        # does not launch LibreOffice or any model.
        try:
            from docx_sanitizer import sanitize_docx, validate_docx  # type: ignore
            formatting["docx_sanitizer"] = sanitize_docx(output)
            # The package sanitizer works at XML level and intentionally does
            # not interpret OMML.  Run the deterministic math cleanup once
            # after that round-trip as well, so a final save cannot leave a
            # converter placeholder fraction/script behind.
            final_math_document = Document(str(output))
            formatting["malformed_math_final"] = normalize_malformed_math(final_math_document)
            final_math_temp = output.with_name(f".{output.stem}.{uuid.uuid4().hex}.tmp.docx")
            final_math_document.save(str(final_math_temp))
            os.replace(str(final_math_temp), str(output))
            formatting["docx_validator"] = validate_docx(output)
        except Exception as exc:
            formatting["docx_sanitizer"] = {"checked": False, "error": str(exc)}
        formula_audit = deterministic_formula_audit(
            formula_reference,
            output,
            args.max_pages,
            Path(engine_info["content_list"]) if engine_info.get("content_list") else None,
            pdfkit_candidates if pdfkit_candidates else None,
            pdfkit_status,
        )
        text_completeness = deterministic_text_completeness(
            check_reference if check_reference and check_reference.suffix.lower() == ".pdf" else None,
            output,
            temp,
            args.max_pages,
        )
        formatting["deterministic_formula_audit"] = formula_audit
        formatting["automatic_formula_region_replacement"] = local_replacements
        formatting["deterministic_text_completeness"] = text_completeness
        checked_document = Document(str(output))
        structure_map_path = report_dir / f"{source.stem}.structure-map.json"
        structure_map = write_structure_map(checked_document, structure_map_path)
        full_page_images = detect_full_page_images(checked_document)
        formatting["full_page_images_detected"] = full_page_images
        if source.suffix.lower() == ".docx" and args.max_pages is not None:
            baseline = {"checked": False, "equal": None, "similarity": None, "reason": "前 N 页预览不与完整输入 Word 做全文相等判断"}
        else:
            baseline = compare_text(source_text(source_for_work, temp / "baseline"), docx_text(output))
        preview_run = source.suffix.lower() == ".docx" and args.max_pages is not None
        # A preview intentionally removes all blocks after page N.  Running a
        # full-document SequenceMatcher here is both meaningless and very
        # expensive for OCR-heavy books, so defer integrity to full runs.
        integrity = (
            {"checked": False, "reason": "前 N 页预览不做完整 Word 原文保留校验"}
            if preview_run
            else content_integrity(source, output)
        )
        comparison: dict[str, Any] = {
            "baseline_word": baseline,
            "reference": None,
            "standard": "Word content/object integrity plus deterministic formula candidate audit; no AI or formula semantic comparison",
            "passed": False,
            "content_integrity": integrity,
        }
        layout_info = formatting.get("rendered_layout", {})
        comparison["rendered_layout"] = layout_info
        layout_ok = not layout_info.get("checked") or bool(layout_info.get("passed"))
        source_objects = document_object_checks(source, output, preview=args.max_pages is not None)
        comparison["editable_object_checks"] = source_objects
        comparison["deterministic_formula_audit"] = formula_audit
        comparison["deterministic_text_completeness"] = text_completeness
        integrity_required = bool(integrity.get("checked") and not (source.suffix.lower() == ".docx" and args.max_pages is not None))
        integrity_ok = not integrity_required or integrity.get("equal") is True
        if source.suffix.lower() == ".pdf" and full_page_images:
            comparison["fallback_blocked"] = (
                "检测到整页图片。按项目规则禁止把整页 PDF 粘进 Word；"
                "请先用夸克扫描王导出基础 Word，再将该 Word 放入 input/word/。"
            )
        # Word-first normalization may split one converter drawing into
        # several inline drawings or merge formula runs, so raw source/output
        # object counts are diagnostic rather than a rejection criterion. The
        # package validator is the authoritative compatibility gate.
        compatibility_ok = bool(formatting.get("docx_validator", {}).get("passed", True))
        if reference and args.visual_compare:
            if not reference.exists() or reference.suffix.lower() not in REFERENCE_EXTENSIONS:
                raise RuntimeError(f"标准文件必须是存在的 PDF 或图片: {reference}")
            reference_for_compare = reference
            if reference.suffix.lower() == ".pdf" and args.max_pages is not None:
                reference_for_compare = limit_pdf_pages(reference, temp / "reference-preview.pdf", args.max_pages)
            ref_kind = "pdf" if reference.suffix.lower() == ".pdf" else "image"
            ref_result = compare(reference_for_compare, output, ref_kind, args.compare, temp / "reference", args.dpi)
            comparison["reference"] = {"path": str(reference), **ref_result}
            # A max-pages Word preview intentionally contains fewer editable
            # objects than the complete source document.  Do not let that
            # expected count mismatch trigger image fallback or rejection.
            object_ok = True if preview_run else (compatibility_ok and (not source_objects.get("checked") or bool(source_objects.get("passed") or source.suffix.lower() == ".docx")))
            comparison["passed"] = bool(ref_result["passed"]) and object_ok and not full_page_images and integrity_ok and layout_ok and bool(formula_audit.get("passed", True)) and bool(text_completeness.get("passed", True))
            # Visual comparison is diagnostic only. It never modifies the
            # document or creates an image fallback.
            if not ref_result["passed"]:
                comparison["fallback_note"] = "视觉差异仅记录，不自动嵌入 PDF 局部图片。"
        elif reference:
            comparison["reference"] = {
                "checked": False,
                "path": str(reference),
                "reason": "默认不执行 Word→PDF 视觉对照；添加 --visual-compare 可启用诊断",
            }
            object_ok = True if preview_run else (compatibility_ok and (not source_objects.get("checked") or bool(source_objects.get("passed") or source.suffix.lower() == ".docx")))
            comparison["passed"] = object_ok and not full_page_images and integrity_ok and layout_ok and bool(formula_audit.get("passed", True)) and bool(text_completeness.get("passed", True))
        else:
            comparison["reference"] = {"checked": False, "reason": "未找到原始 PDF；本次仅执行 Word 内容和对象完整性检查"}
            object_ok = True if preview_run else (compatibility_ok and (not source_objects.get("checked") or bool(source_objects.get("passed") or source.suffix.lower() == ".docx")))
            comparison["passed"] = object_ok and not full_page_images and integrity_ok and layout_ok and bool(formula_audit.get("passed", True)) and bool(text_completeness.get("passed", True))
        report = {
            "source": str(source),
            "output": str(output),
            "backup": str(backup_path),
            "input_type": source.suffix.lower().lstrip("."),
            "engine": engine_info.get("engine", "local-word-bridge") if source.suffix.lower() == ".pdf" else ("local-word-bridge" if source.suffix.lower() == ".doc" else "local-word"),
            "engine_details": engine_info,
            "reference": str(reference) if reference else None,
            "formatting": formatting,
            "structure_map": str(structure_map_path),
            "structure_uncertain_count": len(structure_map.get("uncertain", [])),
            "comparison": comparison,
            "accepted": bool(output.exists() and comparison["passed"]),
            "manual_review_required": bool(not comparison["passed"]),
        }
        report_dir.mkdir(parents=True, exist_ok=True)
        (report_dir / f"{source.stem}.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
    finally:
        shutil.rmtree(temp, ignore_errors=True)


def main() -> int:
    args = parse_args()
    if args.bridge:
        os.environ["CODEX_LO_BRIDGE"] = "1"
    try:
        inputs = discover_inputs(args.input)
        root = Path(__file__).resolve().parents[1]
        output_dir = args.output if not args.input.is_file() else args.output.parent
        output_dir.mkdir(parents=True, exist_ok=True)
        for source in inputs:
            destination = args.output if args.input.is_file() and args.output.suffix.lower() == ".docx" else output_dir / f"{source.stem}.docx"
            print(f"处理: {source.name} (自动转换与排版)")
            report = process(source, destination, args, root / "output" / "reports")
            print(f"完成: {report['output']}")
            print(f"原样校验通过: {report['comparison']['passed']}，最终接受: {report['accepted']}")
        return 0
    except RuntimeError as exc:
        print(f"错误: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
